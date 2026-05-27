#!/usr/bin/env python3
"""Build a side-by-side task-review workbook for current vs proposed noise.

The proposed profile is intentionally scoped to this coverage-extension probe:
it exercises the candidate clean/light/medium/hard policy and the new
coordinate-preserving noise families before they are wired into production
TRACE generation.
"""

from __future__ import annotations

import argparse
import io
import json
import re
import shutil
from pathlib import Path
from typing import Any, Dict, List, Mapping, Sequence, Tuple

import numpy as np
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from PIL import Image, ImageEnhance, ImageFilter, ImageOps

from trace.core.seed import hash64, spawn_rng
from trace.core.taxonomy import resolve_task_taxonomy
from trace.tasks import create_task
from trace.tasks.base import TaskOutput


STRESS5_TASKS: Sequence[Tuple[str, str]] = (
    ("task_pages__command_matrix__command_intent_target_label", "pages_command"),
    ("task_charts__scatter_readout__series_point_lookup_value", "charts_scatter"),
    ("task_games__sudoku__marked_cell_candidate_count", "games_sudoku"),
    ("task_puzzles__nonogram__nonogram_line_completion_label", "puzzles_nonogram"),
    ("task_geometry__graph_paper__line_slope_value", "geometry_slope"),
)

DOMAIN20_STRESS_TASKS: Sequence[Tuple[str, str]] = (
    ("task_charts__heatmap__condition_run_extremum_label", "charts_heatmap"),
    ("task_charts__curve_panels__curve_intersection_count", "charts_scientific"),
    ("task_charts__table__column_summary_value", "charts_table"),
    ("task_charts__bar_3d__condition_count", "charts_3d_bar"),
    ("task_games__chess__check_attacker_count", "games_chess"),
    ("task_games__minesweeper__forced_cell_count", "games_minesweeper"),
    ("task_games__sudoku__marked_cell_candidate_count", "games_sudoku"),
    ("task_geometry__circle_theorem__multi_step_angle_value", "geometry_circle_angle"),
    ("task_geometry__graph_paper__line_slope_value", "geometry_slope"),
    ("task_graph__node_link__edge_color_count", "graph_edge_color"),
    ("task_graph__node_link__mst_weight", "graph_mst"),
    ("task_icons__pattern_grid__color_pattern_violation_index", "icons_color_grid"),
    ("task_icons__overlap_grid__occlusion_order_count", "icons_occlusion"),
    ("task_illustrations__difference_pair__object_difference_count", "illustrations_difference"),
    ("task_pages__command_matrix__command_intent_target_label", "pages_command"),
    ("task_pages__schema__relationship_count", "pages_schema"),
    ("task_physics__wave_interference__interference_point_choice", "physics_waves"),
    ("task_puzzles__nonogram__nonogram_line_completion_label", "puzzles_nonogram"),
    ("task_puzzles__word_search__search_location_label", "puzzles_word_search"),
    ("task_three_d__object_scene__occlusion_order_label", "three_d_occlusion"),
)

TASK_PRESETS: Mapping[str, Sequence[Tuple[str, str]]] = {
    "stress5": STRESS5_TASKS,
    "domain20": DOMAIN20_STRESS_TASKS,
}

CLEAN_PARAMS: Dict[str, Any] = {"visual": {"noise": {"apply_prob": 0.0}}}
PREVIEW_MAX_SIDE = 384
DEFAULT_PROFILE_SHARES: Sequence[Tuple[str, float]] = (
    ("clean", 0.25),
    ("light", 0.25),
    ("medium", 0.25),
    ("hard", 0.25),
)

EDIT_BUCKETS: Mapping[str, Sequence[str]] = {
    "acuity": ("blur", "downsample", "directional_blur", "edge_soften", "unsharp_mask"),
    "compression": ("jpeg", "posterize_quantization"),
    "stochastic_texture": (
        "gaussian_noise",
        "noise",
        "salt_pepper_noise",
        "poisson_noise",
        "speckle_noise",
        "dust_speckle",
    ),
    "illumination": ("brightness_contrast", "exposure_shift", "gamma_shift", "low_contrast_fade", "uneven_illumination", "vignette"),
    "surface_texture": ("screen_or_paper_texture", "scanline_texture", "subpixel_display_texture", "neutral_moire_texture", "ink_bleed", "local_contrast_jitter"),
}

EDIT_RANGES: Mapping[str, Mapping[str, Mapping[str, Tuple[float, float]]]] = {
    "light": {
        "blur": {"radius": (0.05, 0.18)},
        "directional_blur": {"length": (3.0, 5.0), "amount": (0.16, 0.32)},
        "downsample": {"scale": (0.96, 0.99)},
        "edge_soften": {"amount": (0.10, 0.22)},
        "unsharp_mask": {"radius": (0.35, 0.75), "percent": (50.0, 90.0), "threshold": (2.0, 5.0)},
        "jpeg": {"quality": (88.0, 97.0)},
        "posterize_quantization": {"levels": (56.0, 96.0), "amount": (0.08, 0.18)},
        "noise": {"alpha": (0.004, 0.015)},
        "gaussian_noise": {"sigma": (2.0, 6.0)},
        "poisson_noise": {"peak": (900.0, 1500.0)},
        "salt_pepper_noise": {"amount": (0.0008, 0.0025)},
        "speckle_noise": {"sigma": (0.006, 0.016)},
        "dust_speckle": {"amount": (0.0008, 0.0025), "alpha": (0.25, 0.45)},
        "brightness_contrast": {"delta": (0.0, 0.04)},
        "exposure_shift": {"delta": (0.015, 0.045)},
        "gamma_shift": {"gamma": (0.94, 1.06)},
        "low_contrast_fade": {"contrast_drop": (0.04, 0.10), "fade_alpha": (0.02, 0.06)},
        "uneven_illumination": {"strength": (0.025, 0.075)},
        "screen_or_paper_texture": {"alpha": (0.012, 0.028), "grain_sigma": (4.0, 8.0)},
        "scanline_texture": {"alpha": (0.008, 0.018), "period": (3.0, 6.0)},
        "subpixel_display_texture": {"alpha": (0.003, 0.008), "period": (3.0, 4.0)},
        "neutral_moire_texture": {"alpha": (0.003, 0.008), "period": (14.0, 28.0)},
        "ink_bleed": {"amount": (0.04, 0.10)},
        "local_contrast_jitter": {"strength": (0.015, 0.04), "grid_size": (8.0, 14.0)},
        "vignette": {"strength": (0.04, 0.10)},
    },
    "medium": {
        "blur": {"radius": (0.18, 0.35)},
        "directional_blur": {"length": (5.0, 9.0), "amount": (0.28, 0.50)},
        "downsample": {"scale": (0.90, 0.96)},
        "edge_soften": {"amount": (0.22, 0.42)},
        "unsharp_mask": {"radius": (0.75, 1.25), "percent": (90.0, 150.0), "threshold": (1.0, 4.0)},
        "jpeg": {"quality": (76.0, 88.0)},
        "posterize_quantization": {"levels": (32.0, 56.0), "amount": (0.18, 0.34)},
        "noise": {"alpha": (0.015, 0.04)},
        "gaussian_noise": {"sigma": (6.0, 14.0)},
        "poisson_noise": {"peak": (350.0, 900.0)},
        "salt_pepper_noise": {"amount": (0.0025, 0.006)},
        "speckle_noise": {"sigma": (0.016, 0.04)},
        "dust_speckle": {"amount": (0.0025, 0.006), "alpha": (0.35, 0.60)},
        "brightness_contrast": {"delta": (0.04, 0.08)},
        "exposure_shift": {"delta": (0.045, 0.085)},
        "gamma_shift": {"gamma": (0.88, 1.14)},
        "low_contrast_fade": {"contrast_drop": (0.10, 0.22), "fade_alpha": (0.06, 0.14)},
        "uneven_illumination": {"strength": (0.075, 0.16)},
        "screen_or_paper_texture": {"alpha": (0.028, 0.055), "grain_sigma": (8.0, 12.0)},
        "scanline_texture": {"alpha": (0.018, 0.035), "period": (3.0, 7.0)},
        "subpixel_display_texture": {"alpha": (0.008, 0.014), "period": (3.0, 5.0)},
        "neutral_moire_texture": {"alpha": (0.008, 0.018), "period": (10.0, 24.0)},
        "ink_bleed": {"amount": (0.10, 0.20)},
        "local_contrast_jitter": {"strength": (0.04, 0.085), "grid_size": (6.0, 12.0)},
        "vignette": {"strength": (0.10, 0.22)},
    },
    "hard": {
        "blur": {"radius": (0.35, 0.55)},
        "directional_blur": {"length": (9.0, 13.0), "amount": (0.45, 0.70)},
        "downsample": {"scale": (0.84, 0.90)},
        "edge_soften": {"amount": (0.42, 0.62)},
        "unsharp_mask": {"radius": (1.25, 1.8), "percent": (150.0, 220.0), "threshold": (0.0, 3.0)},
        "jpeg": {"quality": (62.0, 76.0)},
        "posterize_quantization": {"levels": (20.0, 32.0), "amount": (0.34, 0.52)},
        "noise": {"alpha": (0.04, 0.07)},
        "gaussian_noise": {"sigma": (14.0, 22.0)},
        "poisson_noise": {"peak": (150.0, 350.0)},
        "salt_pepper_noise": {"amount": (0.006, 0.015)},
        "speckle_noise": {"sigma": (0.04, 0.075)},
        "dust_speckle": {"amount": (0.006, 0.012), "alpha": (0.45, 0.70)},
        "brightness_contrast": {"delta": (0.08, 0.12)},
        "exposure_shift": {"delta": (0.085, 0.13)},
        "gamma_shift": {"gamma": (0.78, 1.24)},
        "low_contrast_fade": {"contrast_drop": (0.22, 0.34), "fade_alpha": (0.14, 0.24)},
        "uneven_illumination": {"strength": (0.16, 0.25)},
        "screen_or_paper_texture": {"alpha": (0.055, 0.08), "grain_sigma": (12.0, 18.0)},
        "scanline_texture": {"alpha": (0.035, 0.06), "period": (3.0, 8.0)},
        "subpixel_display_texture": {"alpha": (0.014, 0.024), "period": (3.0, 5.0)},
        "neutral_moire_texture": {"alpha": (0.018, 0.032), "period": (8.0, 20.0)},
        "ink_bleed": {"amount": (0.20, 0.32)},
        "local_contrast_jitter": {"strength": (0.085, 0.14), "grid_size": (4.0, 10.0)},
        "vignette": {"strength": (0.22, 0.35)},
    },
}


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out-dir", default="plans/coverage-extension")
    parser.add_argument("--sample-count", type=int, default=20)
    parser.add_argument("--seed", type=int, default=20260522)
    parser.add_argument("--max-attempts", type=int, default=120)
    parser.add_argument("--preset", choices=sorted(TASK_PRESETS), default="stress5")
    parser.add_argument(
        "--profile-shares",
        default=",".join(f"{profile}={share}" for profile, share in DEFAULT_PROFILE_SHARES),
        help="Comma-separated clean/light/medium/hard shares, e.g. clean=0.4,light=0.3,medium=0.2,hard=0.1",
    )
    parser.add_argument(
        "--output-stem",
        default="",
        help="Output stem for workbook/manifest/image folder. Defaults to noise_profile_task_review for stress5, otherwise noise_profile_<preset>_task_review.",
    )
    return parser.parse_args()


def _sanitize_sheet_title(raw: str) -> str:
    title = re.sub(r"[\\\\/*?:\\[\\]]+", "_", str(raw).strip())
    return (title or "task")[:31]


def _build_preview(source: Image.Image) -> Image.Image:
    preview = source.convert("RGB")
    border_px = 1
    inner_max = max(1, int(PREVIEW_MAX_SIDE) - 2 * border_px)
    preview.thumbnail((inner_max, inner_max), Image.Resampling.LANCZOS)
    return ImageOps.expand(preview, border=border_px, fill=(0, 0, 0))


def _odd_int(value: float, *, minimum: int = 3) -> int:
    out = max(int(minimum), int(round(float(value))))
    if out % 2 == 0:
        out += 1
    return out


def _parse_profile_shares(raw: str) -> Sequence[Tuple[str, float]]:
    allowed = {profile for profile, _ in DEFAULT_PROFILE_SHARES}
    values: Dict[str, float] = {}
    for part in str(raw).split(","):
        text = part.strip()
        if not text:
            continue
        if "=" not in text:
            raise ValueError(f"profile share must use name=value syntax: {text!r}")
        name, value = text.split("=", 1)
        profile = name.strip()
        if profile not in allowed:
            raise ValueError(f"unknown profile {profile!r}; expected one of {sorted(allowed)}")
        values[profile] = float(value)
    missing = [profile for profile, _ in DEFAULT_PROFILE_SHARES if profile not in values]
    if missing:
        raise ValueError(f"missing profile shares: {missing}")
    total = sum(values.values())
    if total <= 0:
        raise ValueError("--profile-shares total must be positive")
    return [(profile, float(values[profile]) / float(total)) for profile, _ in DEFAULT_PROFILE_SHARES]


def _format_profile_shares(profile_shares: Sequence[Tuple[str, float]]) -> str:
    return "/".join(str(int(round(float(share) * 100))) for _profile, share in profile_shares)


def _profile_schedule(*, count: int, seed: int, task_id: str, profile_shares: Sequence[Tuple[str, float]]) -> List[str]:
    """Return an exact per-sheet profile mix, shuffled deterministically."""
    raw_counts: List[Tuple[str, int, float]] = []
    assigned = 0
    for profile, share in profile_shares:
        exact = float(count) * float(share)
        whole = int(exact)
        assigned += int(whole)
        raw_counts.append((str(profile), int(whole), exact - float(whole)))

    remaining = max(0, int(count) - int(assigned))
    ranked_remainders = sorted(range(len(raw_counts)), key=lambda i: raw_counts[i][2], reverse=True)
    counts = {profile: whole for profile, whole, _ in raw_counts}
    for index in ranked_remainders[:remaining]:
        profile = raw_counts[index][0]
        counts[profile] = int(counts.get(profile, 0)) + 1

    schedule: List[str] = []
    for profile, _share in profile_shares:
        schedule.extend([str(profile)] * int(counts.get(str(profile), 0)))
    rng = spawn_rng(seed, f"coverage_extension.profile_schedule.{task_id}")
    rng.shuffle(schedule)
    return schedule


def _sample_params(profile: str, edit_type: str, rng: Any) -> Dict[str, float]:
    ranges = EDIT_RANGES[str(profile)][str(edit_type)]
    params: Dict[str, float] = {}
    for key, (lo, hi) in ranges.items():
        value = float(rng.uniform(float(lo), float(hi)))
        if key == "quality":
            value = float(int(round(value)))
        if key in {"grid_size", "levels", "period", "percent", "threshold"}:
            value = float(max(1, int(round(value))))
        if str(edit_type) == "directional_blur" and key == "length":
            value = float(_odd_int(value, minimum=3))
        params[str(key)] = value
    if str(edit_type) == "directional_blur":
        params["angle_degrees"] = float(rng.choice([0, 90]))
    if str(edit_type) in {"scanline_texture", "subpixel_display_texture", "uneven_illumination"}:
        params["axis_code"] = float(rng.choice([0, 1]))
    if str(edit_type) == "neutral_moire_texture":
        params["axis_code"] = float(rng.choice([0, 1, 2]))
        params["phase"] = float(rng.uniform(0.0, 6.283185307179586))
    if str(edit_type) in {"exposure_shift", "uneven_illumination"}:
        params["polarity"] = float(rng.choice([-1, 1]))
    if str(edit_type) == "exposure_shift":
        params["factor"] = float(1.0 + (params["polarity"] * float(params.get("delta", 0.03))))
    return params


def _sample_updated_edits(profile: str, instance_seed: int, task_id: str) -> List[Tuple[str, Dict[str, float]]]:
    if profile == "clean":
        return []
    rng = spawn_rng(instance_seed, f"coverage_extension.updated_edits.{task_id}")
    bucket_names = list(EDIT_BUCKETS.keys())
    if profile == "light":
        edit_types = [edit_type for bucket in EDIT_BUCKETS.values() for edit_type in bucket]
        edit_type = str(rng.choice(edit_types))
        return [(edit_type, _sample_params(profile, edit_type, rng))]
    if profile == "medium":
        chosen_buckets = rng.sample(bucket_names, k=2)
    else:
        chosen_buckets = rng.sample(bucket_names, k=int(rng.choice([2, 3])))
    edits: List[Tuple[str, Dict[str, float]]] = []
    for bucket in chosen_buckets:
        edit_type = str(rng.choice(list(EDIT_BUCKETS[bucket])))
        edits.append((edit_type, _sample_params(profile, edit_type, rng)))
    return edits


def _apply_noise_blend(image: Image.Image, *, alpha: float, seed: int, namespace: str) -> Image.Image:
    rng = spawn_rng(seed, namespace)
    base = image.convert("RGB")
    width, height = base.size
    noise = Image.new("RGB", (width, height))
    pixels = noise.load()
    for y in range(height):
        for x in range(width):
            value = int(rng.randint(0, 255))
            pixels[x, y] = (value, value, value)
    return Image.blend(base, noise, max(0.0, min(1.0, float(alpha))))


def _apply_gaussian_noise(image: Image.Image, *, sigma: float, seed: int, namespace: str) -> Image.Image:
    np_rng = np.random.default_rng(hash64(seed, namespace))
    arr = np.asarray(image.convert("RGB"), dtype=np.float32)
    noisy = arr + np_rng.normal(0.0, float(sigma), size=arr.shape)
    return Image.fromarray(np.clip(noisy, 0, 255).astype(np.uint8))


def _apply_poisson_noise(image: Image.Image, *, peak: float, seed: int, namespace: str) -> Image.Image:
    np_rng = np.random.default_rng(hash64(seed, namespace))
    peak_value = max(1.0, float(peak))
    arr = np.asarray(image.convert("RGB"), dtype=np.float32) / 255.0
    noisy = np_rng.poisson(np.clip(arr, 0.0, 1.0) * peak_value) / peak_value
    return Image.fromarray(np.clip(noisy * 255.0, 0, 255).astype(np.uint8))


def _apply_speckle_noise(image: Image.Image, *, sigma: float, seed: int, namespace: str) -> Image.Image:
    np_rng = np.random.default_rng(hash64(seed, namespace))
    arr = np.asarray(image.convert("RGB"), dtype=np.float32)
    multiplier = 1.0 + np_rng.normal(0.0, float(sigma), size=arr.shape[:2] + (1,))
    return Image.fromarray(np.clip(arr * multiplier, 0, 255).astype(np.uint8))


def _apply_dust_speckle(image: Image.Image, *, amount: float, alpha: float, seed: int, namespace: str) -> Image.Image:
    np_rng = np.random.default_rng(hash64(seed, namespace))
    arr = np.asarray(image.convert("RGB"), dtype=np.float32).copy()
    height, width = arr.shape[:2]
    mask = np_rng.random((height, width)) < max(0.0, min(1.0, float(amount)))
    light_dust = np_rng.random((height, width)) < 0.55
    target = np.zeros_like(arr)
    target[light_dust] = 255.0
    blend = max(0.0, min(1.0, float(alpha)))
    arr[mask] = (arr[mask] * (1.0 - blend)) + (target[mask] * blend)
    return Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))


def _apply_brightness_contrast(image: Image.Image, *, delta: float, seed: int, namespace: str) -> Image.Image:
    rng = spawn_rng(seed, namespace)
    brightness = 1.0 + rng.choice([-1.0, 1.0]) * float(rng.uniform(0.0, float(delta)))
    contrast = 1.0 + rng.choice([-1.0, 1.0]) * float(rng.uniform(0.0, float(delta)))
    out = ImageEnhance.Brightness(image.convert("RGB")).enhance(max(0.5, min(1.5, brightness)))
    return ImageEnhance.Contrast(out).enhance(max(0.5, min(1.5, contrast)))


def _apply_gamma_shift(image: Image.Image, *, gamma: float) -> Image.Image:
    arr = np.asarray(image.convert("RGB"), dtype=np.float32)
    luma = (0.299 * arr[:, :, 0]) + (0.587 * arr[:, :, 1]) + (0.114 * arr[:, :, 2])
    gamma_value = max(0.55, min(1.65, float(gamma)))
    adjusted_luma = (np.clip(luma / 255.0, 0.0, 1.0) ** gamma_value) * 255.0
    delta = (adjusted_luma - luma)[:, :, None]
    return Image.fromarray(np.clip(arr + delta, 0, 255).astype(np.uint8))


def _apply_exposure_shift(image: Image.Image, *, factor: float) -> Image.Image:
    arr = np.asarray(image.convert("RGB"), dtype=np.float32)
    return Image.fromarray(np.clip(arr * max(0.5, min(1.5, float(factor))), 0, 255).astype(np.uint8))


def _apply_uneven_illumination(image: Image.Image, *, strength: float, axis_code: float, polarity: float) -> Image.Image:
    arr = np.asarray(image.convert("RGB"), dtype=np.float32)
    height, width = arr.shape[:2]
    if int(round(float(axis_code))) == 1:
        ramp = np.linspace(-1.0, 1.0, height, dtype=np.float32).reshape(height, 1, 1)
    else:
        ramp = np.linspace(-1.0, 1.0, width, dtype=np.float32).reshape(1, width, 1)
    multiplier = 1.0 + (np.sign(float(polarity) or 1.0) * max(0.0, min(1.0, float(strength))) * ramp)
    return Image.fromarray(np.clip(arr * multiplier, 0, 255).astype(np.uint8))


def _apply_directional_blur(image: Image.Image, *, length: float, amount: float, angle_degrees: float) -> Image.Image:
    base = image.convert("RGB")
    kernel_size = max(3, int(round(float(length))))
    if kernel_size % 2 == 0:
        kernel_size += 1
    pad = kernel_size // 2
    arr = np.asarray(base, dtype=np.float32)
    if int(round(float(angle_degrees))) % 180 == 90:
        padded = np.pad(arr, ((pad, pad), (0, 0), (0, 0)), mode="edge")
        blurred = sum(padded[offset : offset + arr.shape[0], :, :] for offset in range(kernel_size)) / float(kernel_size)
    else:
        padded = np.pad(arr, ((0, 0), (pad, pad), (0, 0)), mode="edge")
        blurred = sum(padded[:, offset : offset + arr.shape[1], :] for offset in range(kernel_size)) / float(kernel_size)
    alpha = max(0.0, min(1.0, float(amount)))
    out = (arr * (1.0 - alpha)) + (blurred * alpha)
    return Image.fromarray(np.clip(out, 0, 255).astype(np.uint8))


def _apply_edge_soften(image: Image.Image, *, amount: float) -> Image.Image:
    base = image.convert("RGB")
    softened = base.filter(ImageFilter.SMOOTH_MORE)
    return Image.blend(base, softened, max(0.0, min(1.0, float(amount))))


def _apply_unsharp_mask(image: Image.Image, *, radius: float, percent: float, threshold: float) -> Image.Image:
    return image.convert("RGB").filter(
        ImageFilter.UnsharpMask(
            radius=max(0.1, float(radius)),
            percent=max(1, int(round(float(percent)))),
            threshold=max(0, int(round(float(threshold)))),
        )
    )


def _apply_ink_bleed(image: Image.Image, *, amount: float) -> Image.Image:
    base = image.convert("RGB")
    arr = np.asarray(base, dtype=np.float32)
    gray = np.asarray(base.convert("L"), dtype=np.float32)
    spread = np.asarray(base.convert("L").filter(ImageFilter.MinFilter(3)), dtype=np.float32)
    dark_delta = np.clip((gray - spread) / 255.0, 0.0, 1.0)
    factor = 1.0 - (max(0.0, min(1.0, float(amount))) * dark_delta[:, :, None])
    return Image.fromarray(np.clip(arr * factor, 0, 255).astype(np.uint8))


def _apply_salt_pepper_noise(image: Image.Image, *, amount: float, seed: int, namespace: str) -> Image.Image:
    base = image.convert("RGB")
    arr = np.asarray(base, dtype=np.uint8).copy()
    np_rng = np.random.default_rng(hash64(seed, namespace))
    height, width = arr.shape[:2]
    mask = np_rng.random((height, width)) < max(0.0, min(1.0, float(amount)))
    salt = np_rng.random((height, width)) < 0.5
    arr[mask & salt] = 255
    arr[mask & ~salt] = 0
    return Image.fromarray(arr)


def _apply_posterize_quantization(image: Image.Image, *, levels: float, amount: float) -> Image.Image:
    arr = np.asarray(image.convert("RGB"), dtype=np.float32)
    level_count = max(2.0, float(levels))
    luma = (0.299 * arr[:, :, 0]) + (0.587 * arr[:, :, 1]) + (0.114 * arr[:, :, 2])
    step = 255.0 / max(1.0, level_count - 1.0)
    quantized_luma = np.round(luma / step) * step
    delta = (quantized_luma - luma)[:, :, None] * max(0.0, min(1.0, float(amount)))
    return Image.fromarray(np.clip(arr + delta, 0, 255).astype(np.uint8))


def _apply_low_contrast_fade(image: Image.Image, *, contrast_drop: float, fade_alpha: float) -> Image.Image:
    base = image.convert("RGB")
    contrast = max(0.25, min(1.0, 1.0 - float(contrast_drop)))
    low_contrast = ImageEnhance.Contrast(base).enhance(contrast)
    paper = Image.new("RGB", base.size, (246, 246, 238))
    return Image.blend(low_contrast, paper, max(0.0, min(1.0, float(fade_alpha))))


def _apply_scanline_texture(image: Image.Image, *, alpha: float, period: float, axis_code: float) -> Image.Image:
    arr = np.asarray(image.convert("RGB"), dtype=np.float32)
    height, width = arr.shape[:2]
    line_period = max(2, int(round(float(period))))
    if int(round(float(axis_code))) == 1:
        coords = np.arange(width).reshape(1, width)
    else:
        coords = np.arange(height).reshape(height, 1)
    mask = (coords % line_period) == 0
    multiplier = np.ones((height, width), dtype=np.float32)
    multiplier[mask.repeat(height, axis=0) if mask.shape[0] == 1 else mask.repeat(width, axis=1)] -= max(0.0, min(0.25, float(alpha)))
    return Image.fromarray(np.clip(arr * multiplier[:, :, None], 0, 255).astype(np.uint8))


def _apply_subpixel_display_texture(image: Image.Image, *, alpha: float, period: float, axis_code: float) -> Image.Image:
    arr = np.asarray(image.convert("RGB"), dtype=np.float32)
    height, width = arr.shape[:2]
    line_period = max(3, int(round(float(period))))
    if int(round(float(axis_code))) == 1:
        coords = np.arange(width, dtype=np.float32).reshape(1, width)
        pattern = ((coords % line_period) / max(1.0, float(line_period - 1)) - 0.5).repeat(height, axis=0)
    else:
        coords = np.arange(height, dtype=np.float32).reshape(height, 1)
        pattern = ((coords % line_period) / max(1.0, float(line_period - 1)) - 0.5).repeat(width, axis=1)
    multiplier = 1.0 + (pattern * 2.0 * max(0.0, min(0.05, float(alpha))))
    return Image.fromarray(np.clip(arr * multiplier[:, :, None], 0, 255).astype(np.uint8))


def _apply_neutral_moire_texture(
    image: Image.Image,
    *,
    alpha: float,
    period: float,
    axis_code: float,
    phase: float,
) -> Image.Image:
    arr = np.asarray(image.convert("RGB"), dtype=np.float32)
    height, width = arr.shape[:2]
    yy = np.arange(height, dtype=np.float32).reshape(height, 1)
    xx = np.arange(width, dtype=np.float32).reshape(1, width)
    wave_period = max(4.0, float(period))
    if int(round(float(axis_code))) == 1:
        pattern = np.sin(((yy / wave_period) * 6.283185307179586) + float(phase))
    elif int(round(float(axis_code))) == 2:
        pattern = np.sin((((xx + yy) / wave_period) * 6.283185307179586) + float(phase))
    else:
        pattern = np.sin(((xx / wave_period) * 6.283185307179586) + float(phase))
    multiplier = 1.0 + (pattern[:, :, None] * max(0.0, min(0.05, float(alpha))))
    return Image.fromarray(np.clip(arr * multiplier, 0, 255).astype(np.uint8))


def _apply_screen_or_paper_texture(
    image: Image.Image,
    *,
    alpha: float,
    grain_sigma: float,
    seed: int,
    namespace: str,
) -> Image.Image:
    base = image.convert("RGB")
    width, height = base.size
    np_rng = np.random.default_rng(hash64(seed, namespace))
    arr = np.asarray(base, dtype=np.float32)
    grain = np_rng.normal(0.0, float(grain_sigma), size=(height, width, 1))
    y = np.arange(height, dtype=np.float32).reshape(height, 1, 1)
    stripe_period = float(np_rng.integers(4, 9))
    stripes = np.sin((y / stripe_period) * np.pi * 2.0) * float(grain_sigma) * 0.45
    textured = arr + (grain + stripes) * max(0.0, min(1.0, float(alpha))) * 4.0
    return Image.fromarray(np.clip(textured, 0, 255).astype(np.uint8))


def _apply_local_contrast_jitter(
    image: Image.Image,
    *,
    strength: float,
    grid_size: float,
    seed: int,
    namespace: str,
) -> Image.Image:
    base = image.convert("RGB")
    arr = np.asarray(base, dtype=np.float32)
    height, width = arr.shape[:2]
    cells = max(2, int(round(float(grid_size))))
    np_rng = np.random.default_rng(hash64(seed, namespace))
    low_res = 1.0 + np_rng.uniform(-float(strength), float(strength), size=(cells, cells)).astype(np.float32)
    factor_img = Image.fromarray(np.clip(low_res * 127.5, 0, 255).astype(np.uint8))
    factor_img = factor_img.resize((width, height), resample=Image.Resampling.BILINEAR)
    factor = (np.asarray(factor_img, dtype=np.float32) / 127.5)[:, :, None]
    out = ((arr - 127.5) * factor) + 127.5
    return Image.fromarray(np.clip(out, 0, 255).astype(np.uint8))


def _apply_vignette(image: Image.Image, *, strength: float) -> Image.Image:
    base = image.convert("RGB")
    arr = np.asarray(base, dtype=np.float32)
    height, width = arr.shape[:2]
    yy = np.linspace(-1.0, 1.0, height, dtype=np.float32).reshape(height, 1)
    xx = np.linspace(-1.0, 1.0, width, dtype=np.float32).reshape(1, width)
    distance = np.clip(np.sqrt((xx * xx) + (yy * yy)) / np.sqrt(2.0), 0.0, 1.0)
    mask = 1.0 - (max(0.0, min(1.0, float(strength))) * (distance ** 1.8))
    out = arr * mask[:, :, None]
    return Image.fromarray(np.clip(out, 0, 255).astype(np.uint8))


def _apply_one_edit(image: Image.Image, edit_type: str, params: Mapping[str, float], *, seed: int, edit_index: int) -> Image.Image:
    base = image.convert("RGB")
    namespace = f"coverage_extension.apply.{edit_type}.{edit_index}"
    if edit_type == "blur":
        return base.filter(ImageFilter.GaussianBlur(radius=float(params.get("radius", 0.2))))
    if edit_type == "directional_blur":
        return _apply_directional_blur(
            base,
            length=float(params.get("length", 3.0)),
            amount=float(params.get("amount", 0.25)),
            angle_degrees=float(params.get("angle_degrees", 0.0)),
        )
    if edit_type == "downsample":
        scale = max(0.05, min(1.0, float(params.get("scale", 0.95))))
        width, height = base.size
        low = base.resize(
            (max(1, int(round(width * scale))), max(1, int(round(height * scale)))),
            resample=Image.Resampling.BILINEAR,
        )
        return low.resize((width, height), resample=Image.Resampling.NEAREST)
    if edit_type == "edge_soften":
        return _apply_edge_soften(base, amount=float(params.get("amount", 0.18)))
    if edit_type == "unsharp_mask":
        return _apply_unsharp_mask(
            base,
            radius=float(params.get("radius", 0.8)),
            percent=float(params.get("percent", 100.0)),
            threshold=float(params.get("threshold", 3.0)),
        )
    if edit_type == "jpeg":
        quality = max(5, min(95, int(round(float(params.get("quality", 85.0))))))
        buffer = io.BytesIO()
        base.save(buffer, format="JPEG", quality=quality)
        buffer.seek(0)
        return Image.open(buffer).convert("RGB")
    if edit_type == "posterize_quantization":
        return _apply_posterize_quantization(
            base,
            levels=float(params.get("levels", 64.0)),
            amount=float(params.get("amount", 0.16)),
        )
    if edit_type == "noise":
        return _apply_noise_blend(base, alpha=float(params.get("alpha", 0.02)), seed=seed, namespace=namespace)
    if edit_type == "gaussian_noise":
        return _apply_gaussian_noise(base, sigma=float(params.get("sigma", 6.0)), seed=seed, namespace=namespace)
    if edit_type == "poisson_noise":
        return _apply_poisson_noise(base, peak=float(params.get("peak", 900.0)), seed=seed, namespace=namespace)
    if edit_type == "salt_pepper_noise":
        return _apply_salt_pepper_noise(base, amount=float(params.get("amount", 0.004)), seed=seed, namespace=namespace)
    if edit_type == "speckle_noise":
        return _apply_speckle_noise(base, sigma=float(params.get("sigma", 0.016)), seed=seed, namespace=namespace)
    if edit_type == "dust_speckle":
        return _apply_dust_speckle(
            base,
            amount=float(params.get("amount", 0.002)),
            alpha=float(params.get("alpha", 0.35)),
            seed=seed,
            namespace=namespace,
        )
    if edit_type == "brightness_contrast":
        return _apply_brightness_contrast(base, delta=float(params.get("delta", 0.04)), seed=seed, namespace=namespace)
    if edit_type == "exposure_shift":
        return _apply_exposure_shift(base, factor=float(params.get("factor", 1.0)))
    if edit_type == "gamma_shift":
        return _apply_gamma_shift(base, gamma=float(params.get("gamma", 1.0)))
    if edit_type == "uneven_illumination":
        return _apply_uneven_illumination(
            base,
            strength=float(params.get("strength", 0.06)),
            axis_code=float(params.get("axis_code", 0.0)),
            polarity=float(params.get("polarity", 1.0)),
        )
    if edit_type == "low_contrast_fade":
        return _apply_low_contrast_fade(
            base,
            contrast_drop=float(params.get("contrast_drop", 0.08)),
            fade_alpha=float(params.get("fade_alpha", 0.04)),
        )
    if edit_type == "scanline_texture":
        return _apply_scanline_texture(
            base,
            alpha=float(params.get("alpha", 0.012)),
            period=float(params.get("period", 4.0)),
            axis_code=float(params.get("axis_code", 0.0)),
        )
    if edit_type == "subpixel_display_texture":
        return _apply_subpixel_display_texture(
            base,
            alpha=float(params.get("alpha", 0.006)),
            period=float(params.get("period", 3.0)),
            axis_code=float(params.get("axis_code", 0.0)),
        )
    if edit_type == "neutral_moire_texture":
        return _apply_neutral_moire_texture(
            base,
            alpha=float(params.get("alpha", 0.008)),
            period=float(params.get("period", 16.0)),
            axis_code=float(params.get("axis_code", 0.0)),
            phase=float(params.get("phase", 0.0)),
        )
    if edit_type == "screen_or_paper_texture":
        return _apply_screen_or_paper_texture(
            base,
            alpha=float(params.get("alpha", 0.02)),
            grain_sigma=float(params.get("grain_sigma", 6.0)),
            seed=seed,
            namespace=namespace,
        )
    if edit_type == "ink_bleed":
        return _apply_ink_bleed(base, amount=float(params.get("amount", 0.08)))
    if edit_type == "local_contrast_jitter":
        return _apply_local_contrast_jitter(
            base,
            strength=float(params.get("strength", 0.04)),
            grid_size=float(params.get("grid_size", 8.0)),
            seed=seed,
            namespace=namespace,
        )
    if edit_type == "vignette":
        return _apply_vignette(base, strength=float(params.get("strength", 0.1)))
    return base


def _apply_updated_profile(image: Image.Image, *, profile: str, edits: Sequence[Tuple[str, Mapping[str, float]]], seed: int) -> Image.Image:
    out = image.convert("RGB")
    for index, (edit_type, params) in enumerate(edits):
        out = _apply_one_edit(out, str(edit_type), params, seed=seed, edit_index=index)
    return out


def _serialize_edits(edits: Sequence[Tuple[str, Mapping[str, float]]]) -> List[Dict[str, Any]]:
    return [
        {
            "type": str(edit_type),
            "params": {str(key): round(float(value), 6) for key, value in dict(params).items()},
        }
        for edit_type, params in edits
    ]


def _format_params(params: Mapping[str, Any]) -> str:
    parts: List[str] = []
    for key, value in sorted(dict(params).items()):
        try:
            rendered = f"{float(value):.4g}"
        except Exception:
            rendered = str(value)
        parts.append(f"{key}={rendered}")
    return ", ".join(parts)


def _format_edits(edits: Any) -> str:
    if not isinstance(edits, Sequence) or isinstance(edits, (str, bytes)) or not edits:
        return "none"
    parts: List[str] = []
    for edit in edits:
        if not isinstance(edit, Mapping):
            parts.append(str(edit))
            continue
        edit_type = str(edit.get("type", "unknown"))
        params = edit.get("params", {})
        params_text = _format_params(params) if isinstance(params, Mapping) else str(params)
        parts.append(f"{edit_type}({params_text})" if params_text else edit_type)
    return "\n".join(parts)


def _format_current_noise(meta: Mapping[str, Any]) -> str:
    applied = bool(meta.get("applied", False))
    apply_prob = meta.get("apply_prob", "")
    prefix = f"applied={str(applied).lower()}"
    if apply_prob != "":
        try:
            prefix = f"{prefix}; apply_prob={float(apply_prob):.3g}"
        except Exception:
            prefix = f"{prefix}; apply_prob={apply_prob}"
    edits_text = _format_edits(meta.get("edits", []))
    return f"{prefix}\nedits:\n{edits_text}"


def _format_updated_noise(profile: str, edits: Sequence[Mapping[str, Any]]) -> str:
    return f"profile={profile}\nedits:\n{_format_edits(edits)}"


def _post_image_noise_meta(output: TaskOutput) -> Dict[str, Any]:
    payload = output.trace_payload if isinstance(output.trace_payload, Mapping) else {}
    render_spec = payload.get("render_spec")
    if isinstance(render_spec, Mapping):
        value = render_spec.get("post_image_noise")
        if isinstance(value, Mapping):
            return dict(value)
        style = render_spec.get("style")
        if isinstance(style, Mapping):
            value = style.get("post_image_noise_meta") or style.get("post_image_noise")
            if isinstance(value, Mapping):
                return dict(value)
    value = payload.get("post_image_noise") or payload.get("post_image_noise_meta")
    if isinstance(value, Mapping):
        return dict(value)
    return {}


def _sample_seed(base_seed: int, task_id: str, index: int) -> int:
    return hash64(base_seed, f"coverage_extension.noise_review.{task_id}", index)


def _generate_pair(task_id: str, sample_seed: int, max_attempts: int) -> Tuple[int, TaskOutput, TaskOutput]:
    task = create_task(str(task_id))
    candidate_seeds = [int(sample_seed)]
    candidate_seeds.extend(hash64(sample_seed, f"coverage_extension.retry.{task_id}", i) for i in range(1, 33))
    last_error: Exception | None = None
    for seed in candidate_seeds:
        try:
            current_output = task.generate(int(seed), params={}, max_attempts=max_attempts)
            clean_output = task.generate(int(seed), params=CLEAN_PARAMS, max_attempts=max_attempts)
        except Exception as exc:
            last_error = exc
            continue
        if current_output.answer_gt.to_dict() != clean_output.answer_gt.to_dict():
            last_error = RuntimeError("current and clean generations produced different answers")
            continue
        return int(seed), current_output, clean_output
    raise RuntimeError(f"{task_id} failed after {len(candidate_seeds)} deterministic seed attempts") from last_error


def _add_image(sheet: Any, image: Image.Image, anchor: str, buffers: List[io.BytesIO]) -> None:
    preview = _build_preview(image)
    buffer = io.BytesIO()
    preview.save(buffer, format="PNG")
    buffer.seek(0)
    buffers.append(buffer)
    sheet.add_image(XLImage(buffer), anchor)


def _write_workbook(
    rows_by_task: Mapping[str, Sequence[Mapping[str, Any]]],
    workbook_path: Path,
    *,
    task_specs: Sequence[Tuple[str, str]],
    profile_shares: Sequence[Tuple[str, float]],
) -> None:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    image_buffers: List[io.BytesIO] = []
    bold = Font(bold=True)
    top = Alignment(vertical="top")
    labels_by_task = {task_id: label for task_id, label in task_specs}
    profile_share_text = _format_profile_shares(profile_shares)

    for task_id, rows in rows_by_task.items():
        sheet_title = _sanitize_sheet_title(labels_by_task[str(task_id)])
        sheet = workbook.create_sheet(sheet_title)
        sheet.append(
            [
                "current noise profile",
                f"updated profile: clean/light(easy)/medium/hard = {profile_share_text}",
                "original noise applied",
                "noise applied in updated image",
                "current image path",
                "updated image path",
            ]
        )
        for col in range(1, 7):
            cell = sheet.cell(row=1, column=col)
            cell.font = bold
            cell.alignment = top
            if col <= 2:
                width = 56
            elif col <= 4:
                width = 42
            else:
                width = 64
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = "A2"

        for row_index, row in enumerate(rows, start=2):
            sheet.row_dimensions[row_index].height = PREVIEW_MAX_SIDE * 0.75
            for col in range(1, 7):
                sheet.cell(row=row_index, column=col).alignment = Alignment(wrap_text=True, vertical="top")
            _add_image(sheet, Image.open(row["current_image_path"]), f"A{row_index}", image_buffers)
            _add_image(sheet, Image.open(row["updated_image_path"]), f"B{row_index}", image_buffers)
            sheet.cell(row=row_index, column=3).value = _format_current_noise(row.get("current_post_image_noise", {}))
            sheet.cell(row=row_index, column=4).value = _format_updated_noise(
                str(row.get("updated_profile", "")),
                list(row.get("updated_edits", [])),
            )
            current_path_text = _relative_path_text(str(row["current_image_path"]), base_dir=Path.cwd().resolve())
            updated_path_text = _relative_path_text(str(row["updated_image_path"]), base_dir=Path.cwd().resolve())
            sheet.cell(row=row_index, column=5).value = current_path_text
            sheet.cell(row=row_index, column=6).value = updated_path_text

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(workbook_path)


def _relative_path_text(path_text: str, *, base_dir: Path) -> str:
    path = Path(path_text)
    try:
        return path.relative_to(base_dir).as_posix()
    except ValueError:
        return str(path)


def main() -> int:
    args = _parse_args()
    out_dir = Path(args.out_dir).resolve()
    task_specs = list(TASK_PRESETS[str(args.preset)])
    profile_shares = _parse_profile_shares(str(args.profile_shares))
    output_stem = str(args.output_stem).strip()
    if not output_stem:
        output_stem = "noise_profile_task_review" if str(args.preset) == "stress5" else f"noise_profile_{args.preset}_task_review"
    image_root = out_dir / f"{output_stem}_images"
    workbook_path = out_dir / f"{output_stem}.xlsx"
    manifest_path = out_dir / f"{output_stem}_manifest.json"

    if int(args.sample_count) <= 0:
        raise ValueError("--sample-count must be positive")
    if int(args.max_attempts) <= 0:
        raise ValueError("--max-attempts must be positive")

    shutil.rmtree(image_root, ignore_errors=True)
    image_root.mkdir(parents=True, exist_ok=True)

    rows_by_task: Dict[str, List[Dict[str, Any]]] = {}
    manifest: Dict[str, Any] = {
        "workbook": str(workbook_path.relative_to(out_dir.parent).as_posix()),
        "image_root": str(image_root.relative_to(out_dir.parent).as_posix()),
        "seed": int(args.seed),
        "sample_count_per_task": int(args.sample_count),
        "task_preset": str(args.preset),
        "updated_profile_shares": {profile: share for profile, share in profile_shares},
        "updated_edit_types": sorted({edit_type for edits in EDIT_BUCKETS.values() for edit_type in edits}),
        "tasks": {},
    }

    for task_id, sheet_label in task_specs:
        task = create_task(str(task_id))
        taxonomy = resolve_task_taxonomy(
            str(task_id),
            source_domain=str(getattr(task, "domain", "")),
            source_task_group=str(getattr(task, "task_group", "")),
        )
        task_image_dir = image_root / str(task_id)
        task_image_dir.mkdir(parents=True, exist_ok=True)
        task_rows: List[Dict[str, Any]] = []
        profile_schedule = _profile_schedule(
            count=int(args.sample_count),
            seed=int(args.seed),
            task_id=str(task_id),
            profile_shares=profile_shares,
        )

        for index in range(int(args.sample_count)):
            requested_seed = _sample_seed(int(args.seed), str(task_id), index)
            final_seed, current_output, clean_output = _generate_pair(str(task_id), requested_seed, int(args.max_attempts))
            profile = str(profile_schedule[index])
            edits = _sample_updated_edits(profile, final_seed, str(task_id))
            updated_image = _apply_updated_profile(clean_output.image, profile=profile, edits=edits, seed=final_seed)

            current_path = task_image_dir / f"{index:04d}_current.png"
            updated_path = task_image_dir / f"{index:04d}_updated.png"
            current_output.image.convert("RGB").save(current_path, format="PNG")
            updated_image.convert("RGB").save(updated_path, format="PNG")

            task_rows.append(
                {
                    "index": int(index),
                    "requested_seed": int(requested_seed),
                    "instance_seed": int(final_seed),
                    "task": str(task_id),
                    "domain": str(taxonomy.domain),
                    "scene_id": str(taxonomy.scene_id),
                    "query_id": str(getattr(current_output, "query_id", "")),
                    "query_variant": str(getattr(current_output, "query_variant", "")),
                    "answer_gt": current_output.answer_gt.to_dict(),
                    "current_post_image_noise": _post_image_noise_meta(current_output),
                    "updated_profile": str(profile),
                    "updated_edits": _serialize_edits(edits),
                    "current_image_path": str(current_path),
                    "updated_image_path": str(updated_path),
                }
            )

        rows_by_task[str(task_id)] = task_rows
        manifest["tasks"][str(task_id)] = {
            "sheet": str(sheet_label),
            "domain": str(taxonomy.domain),
            "scene_id": str(taxonomy.scene_id),
            "rows": [
                {
                    key: (
                        _relative_path_text(str(value), base_dir=out_dir.parent)
                        if key.endswith("_image_path")
                        else value
                    )
                    for key, value in row.items()
                    if key not in {"current_image_path", "updated_image_path"} or isinstance(value, str)
                }
                for row in task_rows
            ],
        }
        print(f"[done] {task_id}: {len(task_rows)} rows")

    _write_workbook(rows_by_task, workbook_path, task_specs=task_specs, profile_shares=profile_shares)
    manifest_path.write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(f"[done] workbook: {workbook_path}")
    print(f"[done] manifest: {manifest_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
