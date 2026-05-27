#!/usr/bin/env python3
"""Analyze current accepted task solve rates against heuristic random baselines."""

from __future__ import annotations

import argparse
import json
import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook


_LOW_CONFIDENCE = {"low", "unresolved"}
_CONFIDENCE_RANK = {"high": 3, "medium": 2, "low": 1, "unresolved": 0}
_ANSWER_SUPPORT_KEY_RE = re.compile(r"(?:^|_)(?:target_)?answer(?:_|$)|answer_support|answer_range")
_ANSWER_PROB_KEY_RE = re.compile(r"(?:^|_)(?:target_)?answer(?:_|$).*probabilities$")
_CANDIDATE_WORDS = (
    "candidate",
    "option",
    "choice",
    "label",
    "row",
    "column",
    "landmark",
    "stage",
    "node",
    "field",
    "item",
    "name",
    "control",
    "target",
)
_CANDIDATE_COUNT_KEYS = {
    "option_count",
    "num_options",
    "choice_count",
    "choices_count",
    "candidate_count",
    "candidate_label_count",
    "label_count",
}
_STRING_FIELD_KEYS = {
    "label",
    "name",
    "text",
    "row_label",
    "column_header",
    "landmark_label",
    "stage_label",
    "node_label",
    "field_label",
    "item_label",
    "control_label",
    "target_label",
    "title",
}
_NUMERIC_RANGE_HINT_WORDS = {
    "answer",
    "target",
    "value",
    "count",
    "distance",
    "length",
    "drop",
    "size",
    "sum",
    "total",
    "weight",
    "resistance",
    "force",
    "extension",
    "rank",
}
_NUMERIC_RANGE_EXCLUDE_WORDS = {
    "bbox",
    "coord",
    "coordinate",
    "row_count",
    "column_count",
    "numeric_column_count",
    "total_column_count",
    "canvas",
    "pixel",
    "color",
}


@dataclass(frozen=True)
class ProgressTaskRow:
    domain: str
    task_id: str
    status: str
    best_config: str
    hard: float | None
    easy: float | None
    band: float | None
    artifacts: str
    notes: str


@dataclass(frozen=True)
class SolveStats:
    mean_solve_rate: float | None
    hard_frac: float | None
    easy_frac: float | None
    band_frac: float | None
    source: str
    by_variant: dict[str, float]


@dataclass(frozen=True)
class ReviewSample:
    task_id: str
    query_id: str
    answer_type: str
    answer_value: Any
    payload: dict[str, Any]
    path: Path | None = None


@dataclass(frozen=True)
class BaselineResult:
    random_baseline: float | None
    method: str
    confidence: str
    support: str
    sample_count: int
    answer_type: str
    empirical_unique_count: int
    empirical_uniform_baseline: float | None
    empirical_prior_baseline: float | None
    notes: str = ""


@dataclass(frozen=True)
class TaskAnalysis:
    row: ProgressTaskRow
    solve: SolveStats
    baseline: BaselineResult
    variant_baselines: dict[str, BaselineResult]
    review_sample_count: int


def _parse_float(value: str) -> float | None:
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except Exception:
        return None


def _first_float(record: Mapping[str, Any], keys: Sequence[str]) -> float | None:
    for key in keys:
        if key in record:
            value = _parse_float(str(record.get(key, "")))
            if value is not None:
                return value
    return None


def _split_markdown_row(line: str) -> list[str]:
    """Split a Markdown table row while preserving pipes inside backticks."""
    text = str(line).strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    cells: list[str] = []
    current: list[str] = []
    in_code = False
    index = 0
    while index < len(text):
        char = text[index]
        if char == "`":
            in_code = not in_code
            current.append(char)
        elif char == "|" and not in_code:
            cells.append("".join(current).strip())
            current = []
        else:
            current.append(char)
        index += 1
    cells.append("".join(current).strip())
    return cells


def parse_progress_summary(path: Path) -> list[ProgressTaskRow]:
    """Return task rows from PROGRESS_SUMMARY.md."""
    rows: list[ProgressTaskRow] = []
    domain = ""
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if line.startswith("## "):
            domain = line[3:].strip()
            continue
        if not line.startswith("| `task_"):
            continue
        cells = _split_markdown_row(line)
        if len(cells) < 8:
            continue
        task_match = re.search(r"`([^`]+)`", cells[0])
        if not task_match:
            continue
        rows.append(
            ProgressTaskRow(
                domain=domain,
                task_id=str(task_match.group(1)),
                status=str(cells[1]).strip(),
                best_config=str(cells[2]).strip(),
                hard=_parse_float(cells[3]),
                easy=_parse_float(cells[4]),
                band=_parse_float(cells[5]),
                artifacts=str(cells[6]).strip(),
                notes=str(cells[7]).strip(),
            )
        )
    return rows


def _domain_slug(domain: str) -> str:
    return str(domain).strip().lower().replace(" ", "_")


def _extract_backticks(text: str) -> list[str]:
    return [str(match) for match in re.findall(r"`([^`]+)`", str(text))]


def _resolve_path(repo_root: Path, raw: str) -> Path:
    path = Path(str(raw).strip())
    if path.is_absolute():
        return path
    return repo_root / path


def _extract_review_label(row: ProgressTaskRow) -> str:
    text = f"{row.artifacts} {row.best_config} {row.notes}"
    labels = re.findall(r"\b(v\d+(?:_[A-Za-z0-9]+)*)\b", text)
    return str(labels[0]) if labels else ""


def _candidate_probe_dirs(repo_root: Path, row: ProgressTaskRow) -> list[Path]:
    text = f"{row.artifacts} {row.notes}"
    paths: list[Path] = []
    for raw in re.findall(r"(?:output|probe output)\s*:\s*`([^`]+)`", text):
        path = _resolve_path(repo_root, raw)
        if path.is_dir():
            paths.append(path)
    if paths:
        return sorted(dict.fromkeys(paths))

    probe_root = repo_root / "rlvr" / "outputs" / "curriculum_probe"
    for token in _extract_backticks(text):
        if row.task_id not in token or "probe" not in token:
            continue
        if token.endswith((".xlsx", ".parquet", ".json", ".jsonl")):
            continue
        for path in sorted(probe_root.glob(f"{token}*")):
            if path.is_dir() and (
                (path / "calibration_stats.json").exists()
                or (path / "per_instance.jsonl").exists()
                or (path / "summary.json").exists()
            ):
                paths.append(path)
    review_label = _extract_review_label(row)
    if review_label:
        for path in sorted(probe_root.glob(f"{row.task_id}_probe_100_{review_label}*")):
            if path.is_dir() and (
                (path / "calibration_stats.json").exists()
                or (path / "per_instance.jsonl").exists()
                or (path / "summary.json").exists()
            ):
                paths.append(path)
    return sorted(dict.fromkeys(paths))


def _candidate_solve_workbooks(repo_root: Path, row: ProgressTaskRow) -> list[Path]:
    text = f"{row.artifacts} {row.notes}"
    paths: list[Path] = []
    for raw in re.findall(r"(?:solve(?: workbook)?|distributions?)\s*:\s*`([^`]+\.xlsx)`", text):
        path = _resolve_path(repo_root, raw)
        if path.exists():
            paths.append(path)
    for token in _extract_backticks(text):
        if token.endswith(".xlsx") and "solve_rate_distribution" in token:
            path = _resolve_path(repo_root, token)
            if path.exists():
                paths.append(path)
    if paths:
        return sorted(dict.fromkeys(paths))

    review_dir = repo_root / "plans" / "task-reviews" / _domain_slug(row.domain) / row.task_id
    candidates = sorted(review_dir.glob("*solve_rate_distribution.xlsx"))
    if not candidates:
        return []
    review_label = _extract_review_label(row)
    if review_label:
        matching = [path for path in candidates if f"_{review_label}" in path.name]
        if matching:
            return matching
    return sorted(candidates, key=lambda path: path.stat().st_mtime, reverse=True)


def _record_from_workbook_sheet(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        return []
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    records: list[dict[str, Any]] = []
    for raw in rows[1:]:
        if not raw or all(value is None for value in raw):
            continue
        record = {headers[index]: raw[index] if index < len(raw) else None for index in range(len(headers))}
        records.append(record)
    return records


def _solve_stats_from_json(path: Path) -> SolveStats:
    data = json.loads(path.read_text(encoding="utf-8"))
    overall = data.get("overall")
    if not isinstance(overall, Mapping):
        overall = data.get("summary")
    if not isinstance(overall, Mapping):
        overall = data
    by_variant: dict[str, float] = {}
    for variant_key in ("by_query_id", "by_variant", "by_query_id"):
        raw_by_variant = data.get(variant_key) or {}
        if isinstance(raw_by_variant, Mapping):
            variant_records = [(name, record) for name, record in raw_by_variant.items()]
        elif isinstance(raw_by_variant, list):
            variant_records = [
                (
                    record.get(
                        "group",
                        record.get("name", record.get("query_id", record.get("query_id", record.get("variant", "")))),
                    ),
                    record,
                )
                for record in raw_by_variant
                if isinstance(record, Mapping)
            ]
        else:
            variant_records = []
        for name, record in variant_records:
            if isinstance(record, Mapping):
                mean = _first_float(record, ("mean_solve_rate", "mean_solve"))
                if mean is not None and str(name).strip():
                    by_variant[str(name)] = float(mean)
    return SolveStats(
        mean_solve_rate=_first_float(overall, ("mean_solve_rate", "mean_solve")),
        hard_frac=_first_float(overall, ("hard_frac", "hard")),
        easy_frac=_first_float(overall, ("easy_frac", "easy")),
        band_frac=_first_float(overall, ("band_frac", "band")),
        source=str(path),
        by_variant=by_variant,
    )


def _solve_stats_from_workbook(path: Path) -> SolveStats:
    summary_records = _record_from_workbook_sheet(path, "summary")
    if not summary_records:
        summary_records = _record_from_workbook_sheet(path, "overall")
    overall = summary_records[0] if summary_records else {}
    by_variant: dict[str, float] = {}
    for sheet_name in ("by_query_id", "by_variant", "by_query_id"):
        for record in _record_from_workbook_sheet(path, sheet_name):
            group = record.get(
                "group",
                record.get("name", record.get("query_id", record.get("query_id", record.get("variant", "")))),
            )
            mean = _first_float(record, ("mean_solve_rate", "mean_solve"))
            if str(group).strip() and mean is not None:
                by_variant[str(group)] = float(mean)
    return SolveStats(
        mean_solve_rate=_first_float(overall, ("mean_solve_rate", "mean_solve")),
        hard_frac=_first_float(overall, ("hard_frac", "hard")),
        easy_frac=_first_float(overall, ("easy_frac", "easy")),
        band_frac=_first_float(overall, ("band_frac", "band")),
        source=str(path),
        by_variant=by_variant,
    )


def _solve_stats_from_per_instance(path: Path, row: ProgressTaskRow, *, easy_threshold: int = 58) -> SolveStats:
    prompt_count = 0
    rollout_count = 0
    solved_count = 0
    hard_count = 0
    easy_count = 0
    band_count = 0
    by_variant_rows: dict[str, list[tuple[int, int]]] = defaultdict(list)
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            text = line.strip()
            if not text:
                continue
            record = json.loads(text)
            rollouts = int(record.get("rollout_count") or 0)
            solved = int(record.get("positive_rollout_count") or record.get("solved_rollouts") or 0)
            prompt_count += 1
            rollout_count += int(rollouts)
            solved_count += int(solved)
            if solved == 0:
                hard_count += 1
            elif solved >= int(easy_threshold):
                easy_count += 1
            else:
                band_count += 1
            variant = str(record.get("query_id", "") or "").strip()
            if variant:
                by_variant_rows[variant].append((solved, rollouts))

    by_variant: dict[str, float] = {}
    for variant, values in by_variant_rows.items():
        variant_solved = sum(item[0] for item in values)
        variant_rollouts = sum(item[1] for item in values)
        if variant_rollouts:
            by_variant[str(variant)] = float(variant_solved) / float(variant_rollouts)

    return SolveStats(
        mean_solve_rate=(float(solved_count) / float(rollout_count)) if rollout_count else None,
        hard_frac=(float(hard_count) / float(prompt_count)) if prompt_count else row.hard,
        easy_frac=(float(easy_count) / float(prompt_count)) if prompt_count else row.easy,
        band_frac=(float(band_count) / float(prompt_count)) if prompt_count else row.band,
        source=str(path),
        by_variant=by_variant,
    )


def _solve_stats_from_probe_summary(path: Path, row: ProgressTaskRow) -> SolveStats:
    data = json.loads(path.read_text(encoding="utf-8"))
    overall = data.get("overall") if isinstance(data, Mapping) else None
    if not isinstance(overall, Mapping):
        overall = data if isinstance(data, Mapping) else {}
    mean = _first_float(overall, ("mean_solve_rate", "positive_rollout_rate", "perfect_rollout_rate"))
    hard = _first_float(overall, ("hard_frac", "zero_solve_rate"))
    return SolveStats(
        mean_solve_rate=mean,
        hard_frac=hard if hard is not None else row.hard,
        easy_frac=row.easy,
        band_frac=row.band,
        source=str(path),
        by_variant={},
    )


def _solve_stats_from_progress(row: ProgressTaskRow) -> SolveStats:
    mean_match = re.search(r"mean solve [`']?([0-9.]+)", row.notes)
    mean = _parse_float(mean_match.group(1)) if mean_match else None
    return SolveStats(
        mean_solve_rate=mean,
        hard_frac=row.hard,
        easy_frac=row.easy,
        band_frac=row.band,
        source="plans/PROGRESS_SUMMARY.md",
        by_variant={},
    )


def load_solve_stats(repo_root: Path, row: ProgressTaskRow) -> SolveStats:
    """Load retained solve stats for one progress-summary row."""
    for probe_dir in _candidate_probe_dirs(repo_root, row):
        stats_path = probe_dir / "calibration_stats.json"
        if stats_path.exists():
            return _solve_stats_from_json(stats_path)
        per_instance_path = probe_dir / "per_instance.jsonl"
        if per_instance_path.exists():
            return _solve_stats_from_per_instance(per_instance_path, row)
        summary_path = probe_dir / "summary.json"
        if summary_path.exists():
            return _solve_stats_from_probe_summary(summary_path, row)
    for workbook_path in _candidate_solve_workbooks(repo_root, row):
        try:
            return _solve_stats_from_workbook(workbook_path)
        except Exception:
            continue
    return _solve_stats_from_progress(row)


def _read_manifest_counts(task_dir: Path) -> dict[str, int]:
    manifest_path = task_dir / "manifest.json"
    if not manifest_path.exists():
        return {}
    try:
        data = json.loads(manifest_path.read_text(encoding="utf-8"))
    except Exception:
        return {}
    variants = data.get("variants")
    if not isinstance(variants, Mapping):
        return {}
    counts: dict[str, int] = {}
    for variant, count in variants.items():
        try:
            counts[str(variant)] = int(count)
        except Exception:
            continue
    return counts


def _trace_payload(payload: Mapping[str, Any]) -> Mapping[str, Any]:
    trace_payload = payload.get("trace_payload")
    if isinstance(trace_payload, Mapping):
        return trace_payload
    return payload


def _query_spec(payload: Mapping[str, Any]) -> Mapping[str, Any]:
    trace = _trace_payload(payload)
    query = trace.get("query_spec") if isinstance(trace, Mapping) else None
    return query if isinstance(query, Mapping) else {}


def _query_params(payload: Mapping[str, Any]) -> Mapping[str, Any]:
    query = _query_spec(payload)
    params = query.get("params")
    return params if isinstance(params, Mapping) else {}


def _execution_trace(payload: Mapping[str, Any]) -> Mapping[str, Any]:
    trace = _trace_payload(payload)
    execution = trace.get("execution_trace") if isinstance(trace, Mapping) else None
    return execution if isinstance(execution, Mapping) else {}


def _query_id_from_payload(payload: Mapping[str, Any], fallback: str) -> str:
    query = _query_spec(payload)
    execution = _execution_trace(payload)
    for value in (
        payload.get("query_id"),
        query.get("query_id"),
        _query_params(payload).get("query_id"),
        execution.get("query_id"),
        fallback,
    ):
        if str(value or "").strip():
            return str(value).strip()
    return "default"


def _load_review_payload(path: Path) -> dict[str, Any] | None:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None
    return data if isinstance(data, dict) else None


def load_review_samples(repo_root: Path, row: ProgressTaskRow) -> list[ReviewSample]:
    """Load review JSON rows for one current task, respecting manifest query-id counts when present."""
    task_dir = repo_root / "plans" / "task-reviews" / _domain_slug(row.domain) / row.task_id
    data_dir = task_dir / "data"
    if not data_dir.exists():
        return []
    manifest_counts = _read_manifest_counts(task_dir)
    files_by_variant: dict[str, list[Path]] = defaultdict(list)
    for path in sorted(data_dir.rglob("*.json")):
        files_by_variant[path.parent.name].append(path)

    selected_paths: list[Path] = []
    if manifest_counts:
        for variant, count in sorted(manifest_counts.items()):
            paths = sorted(files_by_variant.get(str(variant), []))
            selected_paths.extend(paths[: max(0, int(count))])
        expected = sum(max(0, int(count)) for count in manifest_counts.values())
        if expected and len(selected_paths) < max(1, int(expected * 0.5)):
            selected_paths = []
    else:
        for paths in files_by_variant.values():
            selected_paths.extend(sorted(paths))
    if not selected_paths:
        for paths in files_by_variant.values():
            selected_paths.extend(sorted(paths))

    samples: list[ReviewSample] = []
    for path in sorted(dict.fromkeys(selected_paths)):
        payload = _load_review_payload(path)
        if payload is None:
            continue
        answer = payload.get("answer_gt")
        if not isinstance(answer, Mapping):
            trace = _trace_payload(payload)
            answer = trace.get("answer_gt") if isinstance(trace, Mapping) else None
        if not isinstance(answer, Mapping):
            continue
        query_id = _query_id_from_payload(payload, fallback=path.parent.name)
        samples.append(
            ReviewSample(
                task_id=row.task_id,
                query_id=query_id,
                answer_type=str(answer.get("type", "")).strip(),
                answer_value=answer.get("value"),
                payload=payload,
                path=path,
            )
        )
    return samples


def canonical_answer(answer_type: str, value: Any) -> str:
    kind = str(answer_type)
    if kind == "integer":
        return str(int(value))
    if kind == "number":
        numeric = float(value)
        if not math.isfinite(numeric):
            raise ValueError("numeric answer must be finite")
        return format(numeric, ".12g")
    return str(value)


def _as_int(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return int(value)
    if isinstance(value, float) and float(value).is_integer():
        return int(value)
    if isinstance(value, str):
        text = value.strip()
        if re.fullmatch(r"-?\d+", text):
            return int(text)
    return None


def _as_number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        numeric = float(value)
        return numeric if math.isfinite(numeric) else None
    if isinstance(value, str):
        try:
            numeric = float(value.strip())
        except Exception:
            return None
        return numeric if math.isfinite(numeric) else None
    return None


def _answer_in_values(answer_type: str, answer_value: Any, values: Sequence[Any]) -> bool:
    if answer_type == "integer":
        answer_int = _as_int(answer_value)
        return answer_int is not None and any(_as_int(value) == answer_int for value in values)
    if answer_type == "number":
        answer_num = _as_number(answer_value)
        return answer_num is not None and any(_as_number(value) == answer_num for value in values)
    answer_text = str(answer_value)
    return any(str(value) == answer_text for value in values)


def _unique_scalar_values(values: Sequence[Any]) -> list[Any]:
    seen: set[str] = set()
    out: list[Any] = []
    for value in values:
        if isinstance(value, (dict, list)):
            continue
        key = str(value)
        if key in seen:
            continue
        seen.add(key)
        out.append(value)
    return out


def _iter_named_values(value: Any, path: str = "", depth: int = 0) -> Iterable[tuple[str, Any]]:
    if depth > 8:
        return
    if isinstance(value, Mapping):
        for key, child in value.items():
            key_text = str(key)
            child_path = f"{path}.{key_text}" if path else key_text
            yield child_path, child
            yield from _iter_named_values(child, child_path, depth + 1)
    elif isinstance(value, list):
        for index, child in enumerate(value):
            child_path = f"{path}[{index}]"
            yield child_path, child
            yield from _iter_named_values(child, child_path, depth + 1)


def _sample_sources(sample: ReviewSample) -> dict[str, Any]:
    payload = sample.payload
    trace = _trace_payload(payload)
    sources: dict[str, Any] = {
        "query_params": _query_params(payload),
        "query_spec": _query_spec(payload),
        "execution_trace": _execution_trace(payload),
    }
    if isinstance(trace, Mapping):
        for key in ("scene_ir", "witness_symbolic", "render_spec", "render_map"):
            value = trace.get(key)
            if isinstance(value, (Mapping, list)):
                sources[key] = value
    return sources


def _support_count_from_sequence(
    *,
    answer_type: str,
    answer_value: Any,
    values: Sequence[Any],
) -> int | None:
    unique = _unique_scalar_values(values)
    if len(unique) < 2:
        return None
    if not _answer_in_values(answer_type, answer_value, unique):
        return None
    return len(unique)


def _integer_range_count(low: Any, high: Any, answer_value: Any) -> int | None:
    low_int = _as_int(low)
    high_int = _as_int(high)
    answer_int = _as_int(answer_value)
    if low_int is None or high_int is None or answer_int is None:
        return None
    if high_int < low_int:
        low_int, high_int = high_int, low_int
    if not (low_int <= answer_int <= high_int):
        return None
    return int(high_int - low_int + 1)


def _find_numeric_answer_support(sample: ReviewSample) -> tuple[int, str, str] | None:
    sources = _sample_sources(sample)
    for source_name, source in sources.items():
        if not isinstance(source, Mapping):
            continue
        for path, value in _iter_named_values(source):
            key = path.split(".")[-1]
            key_no_index = re.sub(r"\[\d+\]$", "", key)
            if isinstance(value, list) and (
                str(key_no_index) in {"target_answer_support", "answer_support", "valid_answer_support"}
                or _ANSWER_SUPPORT_KEY_RE.search(path)
            ):
                count = _support_count_from_sequence(
                    answer_type=sample.answer_type,
                    answer_value=sample.answer_value,
                    values=value,
                )
                if count is not None:
                    return count, f"{source_name}:{path}", "high"
            if isinstance(value, Mapping) and (
                str(key_no_index) in {"target_answer_probabilities", "answer_probabilities"}
                or _ANSWER_PROB_KEY_RE.search(path)
            ):
                keys = list(value.keys())
                count = _support_count_from_sequence(
                    answer_type=sample.answer_type,
                    answer_value=sample.answer_value,
                    values=keys,
                )
                if count is not None:
                    return count, f"{source_name}:{path}", "high"
            lowered_path = path.lower()
            range_key_has_answer_semantics = (
                str(key_no_index).endswith("_range")
                and any(word in lowered_path for word in _NUMERIC_RANGE_HINT_WORDS)
                and not any(word in lowered_path for word in _NUMERIC_RANGE_EXCLUDE_WORDS)
            )
            if isinstance(value, list) and len(value) == 2 and (
                "answer" in lowered_path or "target_answer" in lowered_path or range_key_has_answer_semantics
            ):
                count = _integer_range_count(value[0], value[1], sample.answer_value)
                if count is not None:
                    return count, f"{source_name}:{path}", "high"

        direct_pairs = (
            ("answer_min", "answer_max"),
            ("target_answer_min", "target_answer_max"),
            ("min_answer", "max_answer"),
        )
        for low_key, high_key in direct_pairs:
            if low_key in source and high_key in source:
                count = _integer_range_count(source[low_key], source[high_key], sample.answer_value)
                if count is not None:
                    return count, f"{source_name}:{low_key}/{high_key}", "high"
    return None


def _find_option_candidate_count(sample: ReviewSample) -> tuple[int, str, str] | None:
    sources = _sample_sources(sample)
    for source_name, source in sources.items():
        for path, value in _iter_named_values(source):
            key = re.sub(r"\[\d+\]$", "", path.split(".")[-1])
            if key in _CANDIDATE_COUNT_KEYS:
                count = _as_int(value)
                if count is not None and count >= 2:
                    return int(count), f"{source_name}:{path}", "high"
            if isinstance(value, Mapping) and path.endswith("_probabilities"):
                lowered = path.lower()
                if any(word in lowered for word in ("winner", "candidate", "option", "choice", "label")):
                    count = _support_count_from_sequence(
                        answer_type=sample.answer_type,
                        answer_value=sample.answer_value,
                        values=list(value.keys()),
                    )
                    if count is not None:
                        return count, f"{source_name}:{path}", "high"
            if isinstance(value, list):
                lowered = path.lower()
                if any(word in lowered for word in ("candidate", "option", "choice", "label_pool", "labels")):
                    count = _support_count_from_sequence(
                        answer_type=sample.answer_type,
                        answer_value=sample.answer_value,
                        values=value,
                    )
                    if count is not None:
                        return count, f"{source_name}:{path}", "high"
    return None


def _candidate_count_from_dict_list(
    *,
    answer_type: str,
    answer_value: Any,
    values: Sequence[Any],
) -> int | None:
    candidates: list[Any] = []
    for item in values:
        if not isinstance(item, Mapping):
            continue
        for key in _STRING_FIELD_KEYS:
            if key in item and item[key] not in (None, ""):
                candidates.append(item[key])
    if not candidates:
        return None
    return _support_count_from_sequence(answer_type=answer_type, answer_value=answer_value, values=candidates)


def _find_string_candidate_count(sample: ReviewSample) -> tuple[int, str, str] | None:
    sources = _sample_sources(sample)
    for source_name, source in sources.items():
        for path, value in _iter_named_values(source):
            lowered = path.lower()
            if not any(word in lowered for word in _CANDIDATE_WORDS):
                continue
            if isinstance(value, list):
                count = _support_count_from_sequence(
                    answer_type=sample.answer_type,
                    answer_value=sample.answer_value,
                    values=value,
                )
                if count is not None:
                    confidence = "high" if any(word in lowered for word in ("candidate", "option", "row_labels", "labels")) else "medium"
                    return count, f"{source_name}:{path}", confidence
                dict_count = _candidate_count_from_dict_list(
                    answer_type=sample.answer_type,
                    answer_value=sample.answer_value,
                    values=value,
                )
                if dict_count is not None:
                    return dict_count, f"{source_name}:{path}", "medium"
            if isinstance(value, Mapping):
                keys_count = _support_count_from_sequence(
                    answer_type=sample.answer_type,
                    answer_value=sample.answer_value,
                    values=list(value.keys()),
                )
                if keys_count is not None:
                    return keys_count, f"{source_name}:{path}", "medium"
    return None


def _infer_instance_baseline(sample: ReviewSample) -> tuple[float, int, str, str] | None:
    answer_type = str(sample.answer_type)
    if answer_type.startswith("option"):
        option_count = _find_option_candidate_count(sample)
        if option_count is not None:
            count, source, confidence = option_count
            return 1.0 / float(count), int(count), f"mcq:{source}", confidence
    if answer_type in {"integer", "number"}:
        support = _find_numeric_answer_support(sample)
        if support is not None:
            count, source, confidence = support
            return 1.0 / float(count), int(count), f"numeric_support:{source}", confidence
        return None
    candidate = _find_string_candidate_count(sample)
    if candidate is not None:
        count, source, confidence = candidate
        return 1.0 / float(count), int(count), f"candidate_set:{source}", confidence
    return None


def _empirical_baselines(samples: Sequence[ReviewSample]) -> tuple[int, float | None, float | None, Counter[str]]:
    labels = Counter(canonical_answer(sample.answer_type, sample.answer_value) for sample in samples)
    sample_count = sum(labels.values())
    unique_count = len(labels)
    uniform = (1.0 / float(unique_count)) if unique_count else None
    prior = None
    if sample_count:
        prior = sum((float(count) / float(sample_count)) ** 2.0 for count in labels.values())
    return unique_count, uniform, prior, labels


def _dominant_answer_type(samples: Sequence[ReviewSample]) -> str:
    counts = Counter(str(sample.answer_type) for sample in samples)
    if not counts:
        return ""
    return sorted(counts.items(), key=lambda item: (-int(item[1]), str(item[0])))[0][0]


def _confidence_min(values: Iterable[str]) -> str:
    ranked = sorted((str(value) for value in values), key=lambda value: _CONFIDENCE_RANK.get(value, -1))
    return ranked[0] if ranked else "unresolved"


def estimate_random_baseline(samples: Sequence[ReviewSample], *, split_by_variant: bool = True) -> BaselineResult:
    """Estimate one random baseline from review samples."""
    sample_list = list(samples)
    if not sample_list:
        return BaselineResult(
            random_baseline=None,
            method="missing_review_samples",
            confidence="unresolved",
            support="",
            sample_count=0,
            answer_type="",
            empirical_unique_count=0,
            empirical_uniform_baseline=None,
            empirical_prior_baseline=None,
            notes="no task-review JSON samples found",
        )

    variants = sorted({sample.query_id for sample in sample_list})
    if split_by_variant and len(variants) > 1:
        child_results = {
            variant: estimate_random_baseline(
                [sample for sample in sample_list if sample.query_id == variant],
                split_by_variant=False,
            )
            for variant in variants
        }
        weighted_sum = 0.0
        weight = 0
        missing = False
        for variant, result in child_results.items():
            if result.random_baseline is None:
                missing = True
                continue
            weighted_sum += float(result.random_baseline) * float(result.sample_count)
            weight += int(result.sample_count)
        unique_count, empirical_uniform, empirical_prior, _ = _empirical_baselines(sample_list)
        methods = sorted({result.method for result in child_results.values()})
        supports = sorted({result.support for result in child_results.values() if result.support})
        return BaselineResult(
            random_baseline=(weighted_sum / float(weight)) if weight else None,
            method="weighted_by_query_id",
            confidence="unresolved" if missing else _confidence_min(result.confidence for result in child_results.values()),
            support="; ".join(supports[:4]) + ("; ..." if len(supports) > 4 else ""),
            sample_count=len(sample_list),
            answer_type="mixed" if len({sample.answer_type for sample in sample_list}) > 1 else _dominant_answer_type(sample_list),
            empirical_unique_count=unique_count,
            empirical_uniform_baseline=empirical_uniform,
            empirical_prior_baseline=empirical_prior,
            notes=", ".join(methods[:4]) + ("..." if len(methods) > 4 else ""),
        )

    explicit: list[tuple[float, int, str, str]] = []
    for sample in sample_list:
        inferred = _infer_instance_baseline(sample)
        if inferred is not None:
            explicit.append(inferred)

    unique_count, empirical_uniform, empirical_prior, labels = _empirical_baselines(sample_list)
    answer_type = _dominant_answer_type(sample_list)
    coverage = float(len(explicit)) / float(len(sample_list)) if sample_list else 0.0
    if coverage >= 0.8 and explicit:
        baseline = sum(item[0] for item in explicit) / float(len(explicit))
        counts = [int(item[1]) for item in explicit]
        sources = Counter(item[2].split(":", 1)[0] for item in explicit)
        method = sorted(sources.items(), key=lambda item: (-int(item[1]), str(item[0])))[0][0]
        support = str(counts[0]) if min(counts) == max(counts) else f"{min(counts)}..{max(counts)}"
        confidence = _confidence_min(item[3] for item in explicit)
        notes = "" if coverage >= 0.999 else f"explicit coverage {coverage:.2f}"
        return BaselineResult(
            random_baseline=baseline,
            method=method,
            confidence=confidence,
            support=support,
            sample_count=len(sample_list),
            answer_type=answer_type,
            empirical_unique_count=unique_count,
            empirical_uniform_baseline=empirical_uniform,
            empirical_prior_baseline=empirical_prior,
            notes=notes,
        )

    if answer_type == "integer":
        values = [_as_int(sample.answer_value) for sample in sample_list]
        int_values = [int(value) for value in values if value is not None]
        if int_values:
            low = min(int_values)
            high = max(int_values)
            baseline = 1.0 / float(high - low + 1)
            return BaselineResult(
                random_baseline=baseline,
                method="empirical_integer_range",
                confidence="low",
                support=f"{low}..{high}",
                sample_count=len(sample_list),
                answer_type=answer_type,
                empirical_unique_count=unique_count,
                empirical_uniform_baseline=empirical_uniform,
                empirical_prior_baseline=empirical_prior,
                notes="no explicit answer support found",
            )

    return BaselineResult(
        random_baseline=empirical_uniform,
        method="empirical_unique_answers",
        confidence="low",
        support=str(unique_count),
        sample_count=len(sample_list),
        answer_type=answer_type,
        empirical_unique_count=unique_count,
        empirical_uniform_baseline=empirical_uniform,
        empirical_prior_baseline=empirical_prior,
        notes="no explicit candidate set found",
    )


def analyze_task(repo_root: Path, row: ProgressTaskRow) -> TaskAnalysis:
    samples = load_review_samples(repo_root, row)
    by_variant: dict[str, list[ReviewSample]] = defaultdict(list)
    for sample in samples:
        by_variant[str(sample.query_id)].append(sample)
    variant_baselines = {
        variant: estimate_random_baseline(variant_samples, split_by_variant=False)
        for variant, variant_samples in sorted(by_variant.items())
    }
    return TaskAnalysis(
        row=row,
        solve=load_solve_stats(repo_root, row),
        baseline=estimate_random_baseline(samples, split_by_variant=True),
        variant_baselines=variant_baselines,
        review_sample_count=len(samples),
    )


def _fmt_float(value: float | None, digits: int = 3) -> str:
    if value is None:
        return ""
    return f"{float(value):.{digits}f}"


def _fmt_ratio(mean: float | None, baseline: float | None) -> str:
    if mean is None or baseline is None or baseline <= 0:
        return ""
    return f"{float(mean) / float(baseline):.1f}x"


def _escape_md(value: Any) -> str:
    return str(value).replace("|", "\\|")


def _source_label(path_text: str, repo_root: Path) -> str:
    try:
        path = Path(path_text)
        if path.is_absolute():
            return path.relative_to(repo_root).as_posix()
    except Exception:
        pass
    return str(path_text)


def render_markdown(analyses: Sequence[TaskAnalysis], *, repo_root: Path, progress_path: Path) -> str:
    analyses_by_domain: dict[str, list[TaskAnalysis]] = defaultdict(list)
    for analysis in analyses:
        analyses_by_domain[analysis.row.domain].append(analysis)

    unresolved = [analysis for analysis in analyses if analysis.baseline.confidence in _LOW_CONFIDENCE]
    solved = [analysis for analysis in analyses if analysis.solve.mean_solve_rate is not None]
    baseline_values = [analysis.baseline.random_baseline for analysis in analyses if analysis.baseline.random_baseline is not None]

    lines: list[str] = [
        "# Current Task Random Baseline Analysis",
        "",
        f"Generated: {date.today().isoformat()}",
        "",
        f"Source: `{progress_path.as_posix()}` current active non-counterfactual rows (`accepted` plus `partial`).",
        "",
        "Random baselines are heuristic. MCQ-style tasks use option count; integer tasks prefer explicit answer support and otherwise use empirical generated answer range; label/string tasks prefer visible candidate sets and otherwise use empirical observed answers.",
        "",
        "## Summary",
        "",
        f"- Tasks analyzed: `{len(analyses)}`",
        f"- Tasks with solve-rate artifacts: `{len(solved)}`",
        f"- Tasks with random baseline estimate: `{len(baseline_values)}`",
        f"- Low-confidence baseline rows: `{len(unresolved)}`",
        "",
        "| Domain | Tasks | Mean solve | Mean random baseline | Low-confidence baselines |",
        "|---|---:|---:|---:|---:|",
    ]

    for domain in sorted(analyses_by_domain):
        domain_rows = analyses_by_domain[domain]
        means = [item.solve.mean_solve_rate for item in domain_rows if item.solve.mean_solve_rate is not None]
        baselines = [item.baseline.random_baseline for item in domain_rows if item.baseline.random_baseline is not None]
        low = sum(1 for item in domain_rows if item.baseline.confidence in _LOW_CONFIDENCE)
        mean_solve = sum(means) / float(len(means)) if means else None
        mean_baseline = sum(baselines) / float(len(baselines)) if baselines else None
        lines.append(
            f"| {_escape_md(domain)} | {len(domain_rows)} | {_fmt_float(mean_solve)} | {_fmt_float(mean_baseline)} | {low} |"
        )

    for domain in sorted(analyses_by_domain):
        lines.extend(["", f"## {domain}", ""])
        lines.append("| Task | Mean solve | Random baseline | Solve/random | Method | Conf. | Support | Review n |")
        lines.append("|---|---:|---:|---:|---|---|---:|---:|")
        for analysis in sorted(analyses_by_domain[domain], key=lambda item: item.row.task_id):
            baseline = analysis.baseline
            lines.append(
                "| "
                f"`{analysis.row.task_id}` | "
                f"{_fmt_float(analysis.solve.mean_solve_rate)} | "
                f"{_fmt_float(baseline.random_baseline)} | "
                f"{_fmt_ratio(analysis.solve.mean_solve_rate, baseline.random_baseline)} | "
                f"{_escape_md(baseline.method)} | "
                f"{baseline.confidence} | "
                f"{_escape_md(baseline.support)} | "
                f"{baseline.sample_count} |"
            )

    variant_rows: list[tuple[str, str, str, BaselineResult, float | None]] = []
    for analysis in analyses:
        if len(analysis.variant_baselines) <= 1:
            continue
        baseline_values_for_task = {
            round(float(result.random_baseline), 8)
            for result in analysis.variant_baselines.values()
            if result.random_baseline is not None
        }
        methods = {result.method for result in analysis.variant_baselines.values()}
        answer_types = {result.answer_type for result in analysis.variant_baselines.values()}
        if len(baseline_values_for_task) <= 1 and len(methods) <= 1 and len(answer_types) <= 1:
            continue
        for variant, result in analysis.variant_baselines.items():
            variant_rows.append(
                (
                    analysis.row.domain,
                    analysis.row.task_id,
                    variant,
                    result,
                    analysis.solve.by_variant.get(str(variant)),
                )
            )

    if variant_rows:
        lines.extend(["", "## Variant Baselines", ""])
        lines.append("| Domain | Task | Variant | Variant mean | Random baseline | Method | Conf. | Support | n |")
        lines.append("|---|---|---|---:|---:|---|---|---:|---:|")
        for domain, task_id, variant, result, variant_mean in sorted(variant_rows):
            lines.append(
                "| "
                f"{_escape_md(domain)} | `{task_id}` | `{_escape_md(variant)}` | "
                f"{_fmt_float(variant_mean)} | "
                f"{_fmt_float(result.random_baseline)} | "
                f"{_escape_md(result.method)} | "
                f"{result.confidence} | "
                f"{_escape_md(result.support)} | "
                f"{result.sample_count} |"
            )

    if unresolved:
        lines.extend(["", "## Low-Confidence Notes", ""])
        lines.append("| Domain | Task | Method | Empirical unique | Empirical prior | Note |")
        lines.append("|---|---|---|---:|---:|---|")
        for analysis in sorted(unresolved, key=lambda item: (item.row.domain, item.row.task_id)):
            baseline = analysis.baseline
            lines.append(
                "| "
                f"{_escape_md(analysis.row.domain)} | `{analysis.row.task_id}` | "
                f"{_escape_md(baseline.method)} | "
                f"{baseline.empirical_unique_count} | "
                f"{_fmt_float(baseline.empirical_prior_baseline)} | "
                f"{_escape_md(baseline.notes)} |"
            )

    lines.extend(["", "## Artifact Sources", ""])
    lines.append("| Domain | Task | Solve source |")
    lines.append("|---|---|---|")
    for analysis in sorted(analyses, key=lambda item: (item.row.domain, item.row.task_id)):
        lines.append(
            f"| {_escape_md(analysis.row.domain)} | `{analysis.row.task_id}` | `{_source_label(analysis.solve.source, repo_root)}` |"
        )

    return "\n".join(lines) + "\n"


def _parse_cli() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--progress-summary", default="plans/PROGRESS_SUMMARY.md", type=Path)
    parser.add_argument("--out", default="plans/current_task_random_baseline_analysis.md", type=Path)
    parser.add_argument("--include-status", action="append", default=["accepted", "partial"])
    parser.add_argument("--exclude-domain", action="append", default=["Counterfactual"])
    return parser.parse_args()


def main() -> int:
    args = _parse_cli()
    repo_root = Path.cwd().resolve()
    progress_path = args.progress_summary if args.progress_summary.is_absolute() else repo_root / args.progress_summary
    out_path = args.out if args.out.is_absolute() else repo_root / args.out
    include_statuses = {str(status).strip() for status in args.include_status if str(status).strip()}
    exclude_domains = {str(domain).strip().lower() for domain in args.exclude_domain if str(domain).strip()}

    rows = [
        row
        for row in parse_progress_summary(progress_path)
        if row.status in include_statuses and row.domain.strip().lower() not in exclude_domains
    ]
    analyses = [analyze_task(repo_root, row) for row in rows]
    markdown = render_markdown(analyses, repo_root=repo_root, progress_path=args.progress_summary)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(markdown, encoding="utf-8")
    print(f"wrote {out_path}")
    print(f"tasks analyzed: {len(analyses)}")
    low_conf = sum(1 for analysis in analyses if analysis.baseline.confidence in _LOW_CONFIDENCE)
    print(f"low-confidence baselines: {low_conf}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
