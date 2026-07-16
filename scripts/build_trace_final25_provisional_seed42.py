#!/usr/bin/env python3
"""Assemble the historical 7B base-vs-TRACE seed-42 Final25 snapshot."""

from __future__ import annotations

from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[1]
RESULTS_ROOT = REPO_ROOT / "results"

CORE_SOURCE = RESULTS_ROOT / "qwen25vl3b_7b_answer_grpo_selected25_temp06_3seed_mean_std.xlsx"
EVOCHART_SOURCE = (
    RESULTS_ROOT
    / "trace_evochart_temp06_seed42_qwen25vl7b_base_vs_answer_directjudge_20260715T001847Z_results.xlsx"
)
MMSTAR_SOURCE = (
    RESULTS_ROOT
    / "trace_mmstar_temp06_seed42_qwen25vl7b_base_vs_answer_20260715T013619Z_qwen3_extracted.xlsx"
)
MME_SOURCE = (
    RESULTS_ROOT
    / "trace_mme_reasoning_temp06_seed42_qwen25vl7b_base_vs_answer_20260714T230633Z"
    / "mme_reasoning_results.xlsx"
)

OUTPUT_STEM = RESULTS_ROOT / "trace_final25_provisional_seed42_qwen25vl7b_base_vs_jiswyznz"

BASE_SLUG = "qwen25vl7b-base"
TRACE_SLUG = "trace-qwen25vl7b-easyr1-all1000-answer-nokl-step500"

CATEGORIES = OrderedDict(
    [
        (
            "Charts & Tables",
            ["ChartMuseum", "ChartQAPro", "CharXivReason", "TableVQABench", "EvoChart"],
        ),
        ("Visual Math", ["MathVision", "MathVista", "MathVerse", "WeMath"]),
        ("Science & General Reasoning", ["PhyX mini MC", "Physics", "MMMU-ProVis", "MMStar"]),
        (
            "Spatial & Grounding",
            ["ScreenSpot", "SpatialVizBench COT", "CV-Bench 3D", "ERQA"],
        ),
        (
            "Perception & Counting",
            ["Blink", "CountBenchQA", "CountQA", "TreeBench"],
        ),
        ("Puzzles & Logic", ["PuzzleVQA", "VisualPuzzles", "LogicVista", "MME-Reasoning"]),
    ]
)

ROW_COUNTS = {
    "ChartMuseum": 1000,
    "ChartQAPro": 1948,
    "CharXivReason": 1000,
    "TableVQABench": 1500,
    "EvoChart": 1250,
    "MathVision": 3040,
    "MathVista": 1000,
    "MathVerse": 788,
    "WeMath": 1740,
    "PhyX mini MC": 1000,
    "Physics": 1297,
    "MMMU-ProVis": 1730,
    "MMStar": 1500,
    "ScreenSpot": 1272,
    "SpatialVizBench COT": 1180,
    "CV-Bench 3D": 1200,
    "ERQA": 400,
    "Blink": 1901,
    "CountBenchQA": 487,
    "CountQA": 1528,
    "TreeBench": 405,
    "PuzzleVQA": 2000,
    "VisualPuzzles": 1168,
    "LogicVista": 447,
    "MME-Reasoning": 1188,
}


def rows_as_dicts(path: Path, sheet: str) -> list[dict[str, object]]:
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    rows = ws.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    return [dict(zip(headers, row)) for row in rows]


def assemble() -> list[dict[str, object]]:
    core = {
        str(row["Task"]): row
        for row in rows_as_dicts(CORE_SOURCE, "seed_values")
        if row["seed"] == 42 and row["Task"] != "Average"
    }
    values: dict[str, tuple[float, float, str]] = {}
    for benchmark in ROW_COUNTS:
        if benchmark in core:
            values[benchmark] = (
                float(core[benchmark]["7B Base"]),
                float(core[benchmark]["7B Answer GRPO 500"]),
                CORE_SOURCE.name + ":seed_values",
            )

    evo_rows = rows_as_dicts(EVOCHART_SOURCE, "summary")
    evo = {str(row["model"]): row for row in evo_rows}
    values["EvoChart"] = (
        float(evo[BASE_SLUG]["score"]),
        float(evo[TRACE_SLUG]["score"]),
        EVOCHART_SOURCE.name + ":summary/direct Qwen3 judge",
    )

    mmstar_rows = rows_as_dicts(MMSTAR_SOURCE, "summary")
    mmstar = {str(row["slug"]): row for row in mmstar_rows if row["slug"]}
    values["MMStar"] = (
        float(mmstar[BASE_SLUG]["judge_extracted_accuracy"]),
        float(mmstar[TRACE_SLUG]["judge_extracted_accuracy"]),
        MMSTAR_SOURCE.name + ":summary/judge_extracted_accuracy",
    )

    mme_rows = rows_as_dicts(MME_SOURCE, "Sheet1")
    mme = {str(row["model"]): row for row in mme_rows}
    values["MME-Reasoning"] = (
        float(mme[BASE_SLUG]["accuracy"]),
        float(mme[TRACE_SLUG]["accuracy"]),
        str(MME_SOURCE.relative_to(RESULTS_ROOT)) + ":Sheet1",
    )

    missing = set(ROW_COUNTS) - set(values)
    if missing:
        raise RuntimeError(f"Missing Final25 values: {sorted(missing)}")

    result = []
    for category, benchmarks in CATEGORIES.items():
        for benchmark in benchmarks:
            base, trace, source = values[benchmark]
            result.append(
                {
                    "category": category,
                    "benchmark": benchmark,
                    "rows": ROW_COUNTS[benchmark],
                    "base": base,
                    "trace": trace,
                    "delta": trace - base,
                    "source": source,
                }
            )
    if len(result) != 25 or sum(int(row["rows"]) for row in result) != 31969:
        raise RuntimeError("Final25 benchmark or row-count contract mismatch")
    return result


def mean(rows: list[dict[str, object]], key: str) -> float:
    return sum(float(row[key]) for row in rows) / len(rows)


def summary_rows(details: list[dict[str, object]]) -> list[dict[str, object]]:
    summaries = []
    for category in CATEGORIES:
        rows = [row for row in details if row["category"] == category]
        summaries.append(
            {
                "category": category,
                "benchmarks": len(rows),
                "base": mean(rows, "base"),
                "trace": mean(rows, "trace"),
                "delta": mean(rows, "delta"),
            }
        )
    summaries.append(
        {
            "category": "Overall Final25 macro average",
            "benchmarks": len(details),
            "base": mean(details, "base"),
            "trace": mean(details, "trace"),
            "delta": mean(details, "delta"),
        }
    )
    return summaries


def markdown_table(headers: list[str], rows: list[list[str]], aligns: list[str]) -> str:
    output = ["| " + " | ".join(headers) + " |"]
    output.append("|" + "|".join(f":{'-' * 3}:" if a == "center" else (f"{'-' * 4}:" if a == "right" else f":{'-' * 4}") for a in aligns) + "|")
    output.extend("| " + " | ".join(row) + " |" for row in rows)
    return "\n".join(output)


def write_markdown(details: list[dict[str, object]], summaries: list[dict[str, object]]) -> None:
    lines = [
        "# TRACE Final25 Provisional Seed-42: Qwen2.5-VL-7B Base vs Previous TRACE",
        "",
        "> **Provisional historical assembly, not a validated Final25 campaign result.** "
        "These scores were stitched from earlier runs and were not regenerated or fully rescored "
        "under the frozen Final25 verifier.",
        "",
        "- TRACE checkpoint: previous step-500 model associated with W&B run `jiswyznz`.",
        "- Generation: seed 42, temperature 0.6, top-p 1.0, top-k -1, maximum 4096 tokens "
        "with benchmark-specific caps where applicable.",
        "- Aggregate: unweighted macro average of the 25 primary benchmark scores.",
        "- Coverage: 25 benchmarks and 31,969 rows per model.",
        "- MMStar uses Qwen3-32B extraction followed by exact option matching, the closest available "
        "historical artifact to the frozen contract.",
        "",
        "## Summary",
        "",
        markdown_table(
            ["Category", "Benchmarks", "7B Base", "Previous TRACE", "Delta"],
            [
                [
                    str(row["category"]),
                    str(row["benchmarks"]),
                    f'{float(row["base"]):.2f}',
                    f'{float(row["trace"]):.2f}',
                    f'{float(row["delta"]):+.2f}',
                ]
                for row in summaries
            ],
            ["left", "right", "right", "right", "right"],
        ),
        "",
        "## Detailed Results",
        "",
    ]
    for category in CATEGORIES:
        rows = [row for row in details if row["category"] == category]
        lines.extend(
            [
                f"### {category}",
                "",
                markdown_table(
                    ["Benchmark", "Rows", "7B Base", "Previous TRACE", "Delta"],
                    [
                        [
                            str(row["benchmark"]),
                            f'{int(row["rows"]):,}',
                            f'{float(row["base"]):.2f}',
                            f'{float(row["trace"]):.2f}',
                            f'{float(row["delta"]):+.2f}',
                        ]
                        for row in rows
                    ],
                    ["left", "right", "right", "right", "right"],
                ),
                "",
            ]
        )
    lines.extend(
        [
            "## Limitations",
            "",
            "- The 22 Selected25-derived rows predate the final scoring audit. Historical BLINK, "
            "CV-Bench 3D, ERQA, PhyX, and TreeBench artifacts have known contract issues.",
            "- EvoChart and MMStar only have seed-42 coverage in these historical results.",
            "- MMStar had four unparsed base and two unparsed TRACE extraction responses in the old run.",
            "- MME-Reasoning should be rescored with the current strict deterministic judge wrapper.",
            "- This table must not be substituted for the clean three-seed Final25 comparison.",
            "",
            "## Sources",
            "",
            f"- `{CORE_SOURCE.relative_to(REPO_ROOT)}`",
            f"- `{EVOCHART_SOURCE.relative_to(REPO_ROOT)}`",
            f"- `{MMSTAR_SOURCE.relative_to(REPO_ROOT)}`",
            f"- `{MME_SOURCE.relative_to(REPO_ROOT)}`",
            "",
        ]
    )
    OUTPUT_STEM.with_suffix(".md").write_text("\n".join(lines))


def style_sheet(ws, widths: dict[int, int]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width


def write_excel(details: list[dict[str, object]], summaries: list[dict[str, object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(["category", "benchmarks", "7B base", "previous TRACE", "delta"])
    for row in summaries:
        ws.append([row["category"], row["benchmarks"], row["base"], row["trace"], row["delta"]])
    style_sheet(ws, {1: 48, 2: 12, 3: 14, 4: 18, 5: 12})
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = "0.00"

    ws = wb.create_sheet("detailed_results")
    ws.append(["category", "benchmark", "rows", "7B base", "previous TRACE", "delta", "source"])
    for row in details:
        ws.append(
            [
                row["category"],
                row["benchmark"],
                row["rows"],
                row["base"],
                row["trace"],
                row["delta"],
                row["source"],
            ]
        )
    style_sheet(ws, {1: 48, 2: 25, 3: 10, 4: 14, 5: 18, 6: 12, 7: 90})
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = "0.00"

    ws = wb.create_sheet("metadata")
    ws.append(["key", "value"])
    metadata = [
        ("status", "PROVISIONAL historical assembly; not a validated Final25 campaign"),
        ("generation_seed", 42),
        ("trace_checkpoint", "previous step-500 checkpoint associated with W&B jiswyznz"),
        ("benchmark_count", 25),
        ("rows_per_model", 31969),
        ("aggregation", "unweighted macro average of 25 primary benchmark scores"),
        ("core_source", str(CORE_SOURCE.relative_to(REPO_ROOT))),
        ("evochart_source", str(EVOCHART_SOURCE.relative_to(REPO_ROOT))),
        ("mmstar_source", str(MMSTAR_SOURCE.relative_to(REPO_ROOT))),
        ("mme_reasoning_source", str(MME_SOURCE.relative_to(REPO_ROOT))),
        ("mmstar_metric", "judge_extracted_accuracy"),
        ("limitations", "See the accompanying Markdown report; clean three-seed rerun required"),
    ]
    for row in metadata:
        ws.append(row)
    style_sheet(ws, {1: 25, 2: 110})
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUTPUT_STEM.with_suffix(".xlsx"))


def main() -> None:
    details = assemble()
    summaries = summary_rows(details)
    write_markdown(details, summaries)
    write_excel(details, summaries)
    print(OUTPUT_STEM.with_suffix(".md"))
    print(OUTPUT_STEM.with_suffix(".xlsx"))


if __name__ == "__main__":
    main()
