#!/usr/bin/env python3
"""Export one task-review workbook from an existing RLVR parquet + TRACE dataset root.

This is intended for calibration review sets under plans/task-reviews/, where we
want the same workbook-style artifact as task-reviews/ but for one exact probe
set that already exists on disk.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import shutil
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Sequence, Tuple

import pandas as pd
import zstandard as zstd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
from PIL import ImageOps as PILImageOps

from trace.core.review_overlays import render_evidence_overlay, resolve_overlay_evidence
from trace.core.taxonomy import resolve_task_query_id, resolve_task_taxonomy


_PREVIEW_MAX_SIDE = 384
_EXCEL_HEADERS: List[str] = [
    "image",
    "evidence_image",
    "task",
    "query_variant",
    "scene_id",
    "query_id",
    "prompt_answer",
    "ground_truth_answer",
    "prompt_answer_and_evidence",
    "ground_truth_answer_and_evidence",
    "answer_type",
    "evidence_type",
    "instance_seed",
    "image_path",
    "data_path",
]

_EXCEL_COLUMN_WIDTHS: Dict[str, float] = {
    "A": 54,
    "B": 54,
    "C": 28,
    "D": 24,
    "E": 24,
    "F": 24,
    "G": 34,
    "H": 22,
    "I": 40,
    "J": 24,
    "K": 14,
    "L": 16,
    "M": 24,
    "N": 28,
    "O": 28,
}

_WRAP_COLUMNS = {"C", "D", "E", "F", "G", "H", "I", "J", "M", "N", "O"}


def _parse_cli() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export one workbook for an exact task probe set")
    parser.add_argument("--parquet", required=True, help="RLVR parquet for the exact probe set")
    parser.add_argument("--dataset-root", required=True, help="TRACE dataset root that produced the parquet")
    parser.add_argument("--out-root", default="plans/task-reviews", help="Output root for workbook artifacts")
    parser.add_argument("--task-id", default="", help="Optional task id override")
    parser.add_argument(
        "--review-label",
        default="",
        help="Optional label suffix for the workbook, e.g. level0_200",
    )
    parser.add_argument(
        "--calibration-baseline",
        default="v0",
        help="Calibration artifact baseline label to write into the review manifest.",
    )
    return parser.parse_args()


def _json_cell(value: Any) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    return "" if value is None else str(value)


def _build_preview_image(source: PILImage.Image, *, max_image_side: int = _PREVIEW_MAX_SIDE) -> PILImage.Image:
    preview = source.convert("RGB")
    border_px = 1
    inner_max_side = max(1, int(max_image_side) - (2 * border_px))
    preview.thumbnail((inner_max_side, inner_max_side), PILImage.Resampling.LANCZOS)
    return PILImageOps.expand(preview, border=border_px, fill=(0, 0, 0))


def _sanitize_sheet_title(raw: str) -> str:
    title = re.sub(r"[\\\\/*?:\\[\\]]+", "_", str(raw).strip())
    if not title:
        return "default"
    return title[:31]


def _dedupe_sheet_title(base: str, used: set[str]) -> str:
    candidate = _sanitize_sheet_title(base)
    if candidate not in used:
        used.add(candidate)
        return candidate
    index = 2
    while True:
        suffix = f"_{index}"
        trimmed = candidate[: max(1, 31 - len(suffix))]
        attempt = f"{trimmed}{suffix}"
        if attempt not in used:
            used.add(attempt)
            return attempt
        index += 1


def _review_variant_key(*, query_variant: str, query_id: str) -> str:
    """Return the review grouping key for one exported task instance."""

    query_variant_text = str(query_variant).strip()
    query_id_text = str(query_id).strip()
    if query_variant_text in {"", "default"} and query_id_text:
        return query_id_text
    return query_variant_text


def _ensure_link_or_copy(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists() or destination.is_symlink():
        destination.unlink()
    try:
        relative_target = os.path.relpath(str(source), start=str(destination.parent))
        destination.symlink_to(relative_target)
    except Exception:
        shutil.copy2(source, destination)


def _load_jsonl(path: Path) -> Iterable[Dict[str, Any]]:
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            text = line.strip()
            if not text:
                continue
            yield json.loads(text)


def _load_selected_refs(parquet_path: Path) -> Tuple[str, List[Dict[str, Any]]]:
    df = pd.read_parquet(parquet_path, columns=["task", "trace_ref"])
    if len(df) == 0:
        raise ValueError(f"parquet has no rows: {parquet_path}")
    task_ids = sorted({str(item) for item in df["task"].tolist()})
    if len(task_ids) != 1:
        raise ValueError(f"expected exactly one task in parquet, found: {task_ids}")
    selected_refs: List[Dict[str, Any]] = []
    for value in df["trace_ref"].tolist():
        parsed = value if isinstance(value, Mapping) else json.loads(str(value))
        selected_refs.append(
            {
                "shard_id": str(parsed["shard_id"]),
                "line_index": int(parsed["line_index"]),
                "trace_record_hash": str(parsed.get("trace_record_hash", "")),
            }
        )
    return str(task_ids[0]), selected_refs


def _load_train_instance_index(dataset_root: Path) -> Dict[Tuple[str, int], Dict[str, Any]]:
    index: Dict[Tuple[str, int], Dict[str, Any]] = {}
    for row in _load_jsonl(dataset_root / "train_instances.jsonl"):
        trace_ref = row.get("trace_ref", {})
        if not isinstance(trace_ref, Mapping):
            continue
        key = (str(trace_ref.get("shard_id", "")), int(trace_ref.get("line_index", -1)))
        index[key] = row
    return index


def _load_trace_index(dataset_root: Path, selected_refs: Sequence[Mapping[str, Any]]) -> Dict[Tuple[str, int], Dict[str, Any]]:
    needed_by_shard: Dict[str, set[int]] = defaultdict(set)
    for ref in selected_refs:
        needed_by_shard[str(ref["shard_id"])].add(int(ref["line_index"]))

    out: Dict[Tuple[str, int], Dict[str, Any]] = {}
    for shard_id, line_indices in needed_by_shard.items():
        shard_path = dataset_root / "traces" / str(shard_id)
        with shard_path.open("rb") as handle:
            reader = zstd.ZstdDecompressor().stream_reader(handle)
            for line_index, raw in enumerate(reader.read().decode("utf-8").splitlines()):
                if line_index not in line_indices:
                    continue
                out[(str(shard_id), int(line_index))] = json.loads(raw)
    return out


def _populate_inspection_sheet(
    sheet: Any,
    *,
    rows: Sequence[Mapping[str, Any]],
    out_root: Path,
    image_buffers: List[io.BytesIO],
) -> None:
    sheet.append(list(_EXCEL_HEADERS))

    bold = Font(bold=True)
    for column in range(1, len(_EXCEL_HEADERS) + 1):
        sheet.cell(row=1, column=column).font = bold

    for letter, width in _EXCEL_COLUMN_WIDTHS.items():
        sheet.column_dimensions[str(letter)].width = float(width)

    wrap_top = Alignment(wrap_text=True, vertical="top")

    for row_idx, row in enumerate(rows, start=2):
        image_path = out_root / str(row.get("image_path", ""))
        if image_path.exists():
            with PILImage.open(image_path) as source:
                source_rgb = source.convert("RGB")
                preview = _build_preview_image(source_rgb)
                evidence_preview = _build_preview_image(
                    render_evidence_overlay(
                        source_rgb,
                        evidence_type=str(row.get("overlay_evidence_type", row.get("evidence_type", ""))),
                        evidence_value=row.get("overlay_evidence_value", row.get("answer_evidence")),
                    )
                )
                row_height = max(int(preview.height), int(evidence_preview.height))

                preview_buffer = io.BytesIO()
                preview.save(preview_buffer, format="PNG")
                preview_buffer.seek(0)
                image_buffers.append(preview_buffer)

                evidence_buffer = io.BytesIO()
                evidence_preview.save(evidence_buffer, format="PNG")
                evidence_buffer.seek(0)
                image_buffers.append(evidence_buffer)

            preview_img = XLImage(preview_buffer)
            preview_img.anchor = f"A{row_idx}"
            sheet.add_image(preview_img)

            evidence_img = XLImage(evidence_buffer)
            evidence_img.anchor = f"B{row_idx}"
            sheet.add_image(evidence_img)

            sheet.row_dimensions[row_idx].height = max(60, float(row_height) * 0.75)
        else:
            sheet.row_dimensions[row_idx].height = 60

        values = [
            row.get("task", ""),
            row.get("query_variant", ""),
            row.get("scene_id", ""),
            row.get("query_id", ""),
            row.get("prompt_answer", row.get("prompt_answer_only", "")),
            _json_cell(row.get("ground_truth_answer", row.get("answer_only"))),
            row.get("prompt_answer_and_evidence", row.get("prompt", "")),
            _json_cell(row.get("ground_truth_answer_and_evidence", row.get("answer"))),
            row.get("answer_type", ""),
            row.get("evidence_type", ""),
            int(row.get("instance_seed", 0)),
            row.get("image_path", ""),
            row.get("data_path", ""),
        ]
        for offset, value in enumerate(values, start=3):
            cell = sheet.cell(row=row_idx, column=offset, value=value)
            if get_column_letter(offset) in _WRAP_COLUMNS:
                cell.alignment = wrap_top

    sheet.freeze_panes = "C2"


def _write_inspection_excel(
    rows_by_variant: Mapping[str, Sequence[Mapping[str, Any]]],
    path: Path,
    *,
    out_root: Path,
) -> Dict[str, str]:
    workbook = Workbook()
    image_buffers: List[io.BytesIO] = []
    used_titles: set[str] = set()
    variant_to_sheet: Dict[str, str] = {}

    sorted_variants = sorted(str(variant) for variant in rows_by_variant.keys()) or [""]
    for index, query_variant in enumerate(sorted_variants):
        base_title = str(query_variant).strip() or "default"
        sheet_title = _dedupe_sheet_title(base_title, used_titles)
        if index == 0:
            sheet = workbook.active
            sheet.title = sheet_title
        else:
            sheet = workbook.create_sheet(title=sheet_title)
        variant_to_sheet[str(query_variant)] = str(sheet_title)
        _populate_inspection_sheet(
            sheet,
            rows=list(rows_by_variant.get(str(query_variant), [])),
            out_root=out_root,
            image_buffers=image_buffers,
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return variant_to_sheet


def main() -> int:
    args = _parse_cli()
    parquet_path = Path(str(args.parquet)).resolve()
    dataset_root = Path(str(args.dataset_root)).resolve()
    out_root = Path(str(args.out_root)).resolve()

    inferred_task_id, selected_refs = _load_selected_refs(parquet_path)
    task_id = str(args.task_id).strip() or inferred_task_id
    if task_id != inferred_task_id:
        raise ValueError(f"--task-id {task_id} does not match parquet task {inferred_task_id}")

    taxonomy = resolve_task_taxonomy(str(task_id))
    task_dir = out_root / str(taxonomy.domain) / str(taxonomy.scene_id) / task_id
    images_dir = task_dir / "images"
    data_dir = task_dir / "data"
    shutil.rmtree(images_dir, ignore_errors=True)
    shutil.rmtree(data_dir, ignore_errors=True)
    images_dir.mkdir(parents=True, exist_ok=True)
    data_dir.mkdir(parents=True, exist_ok=True)

    review_label = str(args.review_label).strip()
    workbook_name = f"{task_id}_{review_label}.xlsx" if review_label else f"{task_id}.xlsx"
    workbook_path = task_dir / workbook_name

    train_index = _load_train_instance_index(dataset_root)
    trace_index = _load_trace_index(dataset_root, selected_refs)

    rows_by_variant: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    per_variant_counts: Dict[str, int] = defaultdict(int)

    for ref in selected_refs:
        key = (str(ref["shard_id"]), int(ref["line_index"]))
        instance = train_index.get(key)
        trace_record = trace_index.get(key)
        if instance is None:
            raise KeyError(f"missing train instance for {key}")
        if trace_record is None:
            raise KeyError(f"missing trace record for {key}")

        query_spec = trace_record.get("query_spec", {}) if isinstance(trace_record, Mapping) else {}
        execution_trace = trace_record.get("execution_trace", {}) if isinstance(trace_record, Mapping) else {}
        query_variant = str(query_spec.get("query_variant") or execution_trace.get("query_variant") or "").strip()
        scene_id = str(instance.get("scene_id") or query_spec.get("scene_id") or taxonomy.scene_id)
        query_id = str(
            instance.get("query_id")
            or query_spec.get("query_id")
            or execution_trace.get("query_id")
            or resolve_task_query_id(query_variant=query_variant, trace_payload=trace_record)
        )
        review_variant = _review_variant_key(query_variant=query_variant, query_id=query_id)
        variant_dir = review_variant or "default"
        variant_index = int(per_variant_counts[variant_dir])
        per_variant_counts[variant_dir] += 1

        image_info = list(instance.get("images", []) or [])
        if not image_info:
            raise ValueError(f"instance missing images for {key}")
        image_source = dataset_root / str(image_info[0].get("path", ""))
        image_destination = images_dir / variant_dir / f"{variant_index:04d}.png"
        _ensure_link_or_copy(image_source, image_destination)
        rel_image_path = image_destination.relative_to(out_root).as_posix()

        prompt_variants = dict(instance.get("prompt_variants", {}) or {})
        prompt_answer_only = str(prompt_variants.get("answer_only", instance.get("prompt", "")))
        prompt_answer_and_evidence = str(prompt_variants.get("answer_and_evidence", instance.get("prompt", "")))

        answer_gt = dict(instance.get("answer_gt", {}) or {})
        evidence_gt = dict(instance.get("evidence_gt", {}) or {})
        answer_only_ground_truth = {
            "answer": answer_gt.get("value"),
        }
        canonical_answer = {
            "evidence": evidence_gt.get("value"),
            "answer": answer_gt.get("value"),
        }

        overlay_evidence_type, overlay_evidence_value = resolve_overlay_evidence(
            evidence_type=str(evidence_gt.get("type", "")),
            evidence_value=evidence_gt.get("value"),
            trace_payload={"projected_evidence": trace_record.get("projected_evidence", {})},
        )

        data_payload = {
            "task": task_id,
            "domain": str(taxonomy.domain),
            "scene_id": scene_id,
            "query_variant": query_variant,
            "query_id": query_id,
            "instance_seed": int(instance.get("instance_seed", 0)),
            "instance_id": instance.get("instance_id"),
            "prompt": prompt_answer_and_evidence,
            "prompt_answer": prompt_answer_only,
            "prompt_answer_only": prompt_answer_only,
            "prompt_answer_and_evidence": prompt_answer_and_evidence,
            "prompt_variants": prompt_variants,
            "answer_gt": answer_gt,
            "evidence_gt": evidence_gt,
            "reward_contract": instance.get("reward_contract", {}),
            "image": {
                "path": rel_image_path,
                "source_path": str(image_source),
            },
            "trace_ref": instance.get("trace_ref", {}),
            "query_spec": query_spec,
            "execution_trace": execution_trace,
            "projected_evidence": trace_record.get("projected_evidence", {}),
            "versions": instance.get("versions", {}),
        }
        data_path = data_dir / variant_dir / f"{variant_index:04d}.json"
        data_path.parent.mkdir(parents=True, exist_ok=True)
        data_path.write_text(json.dumps(data_payload, ensure_ascii=False, indent=2), encoding="utf-8")
        rel_data_path = data_path.relative_to(out_root).as_posix()

        rows_by_variant[review_variant].append(
            {
                "task": task_id,
                "query_variant": query_variant,
                "scene_id": scene_id,
                "query_id": query_id,
                "prompt": prompt_answer_and_evidence,
                "prompt_answer": prompt_answer_only,
                "prompt_answer_only": prompt_answer_only,
                "prompt_answer_and_evidence": prompt_answer_and_evidence,
                "ground_truth_answer": answer_only_ground_truth,
                "ground_truth_answer_and_evidence": canonical_answer,
                "answer": canonical_answer,
                "answer_evidence": evidence_gt.get("value"),
                "answer_type": str(answer_gt.get("type", "")),
                "evidence_type": str(evidence_gt.get("type", "")),
                "instance_seed": int(instance.get("instance_seed", 0)),
                "image_path": rel_image_path,
                "data_path": rel_data_path,
                "overlay_evidence_type": str(overlay_evidence_type),
                "overlay_evidence_value": overlay_evidence_value,
            }
        )

    workbook_sheets = _write_inspection_excel(rows_by_variant, workbook_path, out_root=out_root)

    manifest = {
        "task_id": task_id,
        "domain": str(taxonomy.domain),
        "scene_id": str(taxonomy.scene_id),
        "calibration_baseline": str(args.calibration_baseline),
        "review_label": review_label,
        "inspection_count": int(sum(len(rows) for rows in rows_by_variant.values())),
        "variants": {
            str(variant): int(len(rows_by_variant[variant]))
            for variant in sorted(rows_by_variant.keys())
        },
        "source_parquet": str(parquet_path),
        "source_dataset_root": str(dataset_root),
        "workbook": str(workbook_path.relative_to(out_root).as_posix()),
        "workbook_sheets": dict(workbook_sheets),
    }
    (task_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
