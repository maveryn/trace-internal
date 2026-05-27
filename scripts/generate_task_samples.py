#!/usr/bin/env python3
"""Generate task sample images/data plus per-task and per-domain combined review workbooks."""

from __future__ import annotations

import argparse
import io
import json
import math
import re
import shutil
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Mapping, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
from PIL import ImageOps as PILImageOps

from trace.core.evidence_sanitization import sanitize_trace_payload_for_public_evidence
from trace.core.json_io import write_json_file
from trace.core.review_overlays import render_evidence_overlay, resolve_overlay_evidence
from trace.core.seed import hash64
from trace.tasks import TASK_REGISTRY, create_task


_FIELD_LABELS: Dict[str, str] = {
    "domain": "domain",
    "task_group": "task_group",
    "task": "task",
    "sample_index": "sample_index",
    "instance_seed": "instance_seed",
    "query_id": "query_id",
    "image_path": "image_path",
    "data_path": "data_path",
    "prompt_answer": "prompt_answer",
    "ground_truth_answer": "ground_truth_answer",
    "prompt_answer_and_evidence": "prompt_answer_and_evidence",
    "ground_truth_answer_and_evidence": "ground_truth_answer_and_evidence",
    "answer_type": "answer_type",
    "answer_value": "answer_value",
    "evidence_type": "evidence_type",
}

_TASK_SHEET_FIELDS: List[str] = [
    "task",
    "query_id",
    "prompt_answer",
    "ground_truth_answer",
    "prompt_answer_and_evidence",
    "ground_truth_answer_and_evidence",
    "domain",
    "task_group",
    "sample_index",
    "instance_seed",
    "image_path",
    "data_path",
    "answer_type",
    "evidence_type",
]

_COLUMN_WIDTHS_BY_FIELD: Dict[str, float] = {
    "domain": 12,
    "task_group": 16,
    "task": 24,
    "sample_index": 12,
    "instance_seed": 20,
    "query_id": 28,
    "image_path": 22,
    "data_path": 38,
    "prompt_answer": 34,
    "ground_truth_answer": 22,
    "prompt_answer_and_evidence": 40,
    "ground_truth_answer_and_evidence": 24,
    "answer_type": 14,
    "answer_value": 12,
    "evidence_type": 14,
}

_PREVIEW_COLUMN_WIDTH = 56
_INT_FIELDS = {"sample_index", "instance_seed"}
_JSON_FIELDS = {"answer_value", "ground_truth_answer", "ground_truth_answer_and_evidence"}
_WRAP_FIELDS = {
    "task",
    "image_path",
    "data_path",
    "prompt_answer",
    "ground_truth_answer",
    "prompt_answer_and_evidence",
    "ground_truth_answer_and_evidence",
    "answer_value",
}
_TASK_WRAP_MAX_CHARS = 20

_TASK_PREVIEW_MAX_SIDE = 384
_DISTRIBUTION_REVIEW_SAMPLE_COUNT = 500
_DISTRIBUTION_SIGMA_THRESHOLD = 2.0
_DISTRIBUTION_MAX_NUMERIC_BINS = 20


def _parse_json_dict(raw: str, *, arg_name: str) -> Dict[str, Any]:
    """Parse CLI JSON object argument into a dictionary."""
    if not raw:
        return {}
    value = json.loads(raw)
    if not isinstance(value, dict):
        raise ValueError(f"{arg_name} must be a JSON object")
    return dict(value)


def _resolve_task_ids(raw_tasks: str) -> List[str]:
    """Resolve target task ids from CLI input or registry defaults."""
    if not raw_tasks.strip():
        return sorted(TASK_REGISTRY.keys())
    task_ids = [item.strip() for item in raw_tasks.split(",") if item.strip()]
    if not task_ids:
        raise ValueError("--tasks resolved to an empty list")
    unknown = [task_id for task_id in task_ids if task_id not in TASK_REGISTRY]
    if unknown:
        raise ValueError(f"unknown task ids: {', '.join(sorted(unknown))}")
    return sorted(dict.fromkeys(task_ids))


def _json_cell(value: Any) -> str:
    """Serialize nested values for spreadsheet cells."""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    return "" if value is None else str(value)


def _field_value_for_sheet(row: Dict[str, Any], field: str) -> Any:
    """Return one row field value formatted for spreadsheet export."""
    value = row.get(field)
    if field == "task":
        return _wrap_task_label(str(value))
    if field in _INT_FIELDS:
        return int(value)
    if field in _JSON_FIELDS:
        return _json_cell(value)
    return value


def _wrap_task_label(value: str) -> str:
    """Wrap long underscore-separated task ids into multiple lines for readability."""
    text = str(value).strip()
    if not text or "_" not in text:
        return text

    parts = text.split("_")
    lines: List[str] = []
    current = ""
    for idx, part in enumerate(parts):
        suffix = "_" if idx < len(parts) - 1 else ""
        token = f"{part}{suffix}"
        if current and (len(current) + len(token)) > int(_TASK_WRAP_MAX_CHARS):
            lines.append(current)
            current = token
        else:
            current = f"{current}{token}"
    if current:
        lines.append(current)
    return "\n".join(lines)


def _build_preview_image(source: PILImage.Image, *, max_image_side: int) -> PILImage.Image:
    """Resize and border one preview image without mutating source image."""
    preview = source.convert("RGB")
    border_px = 1
    inner_max_side = max(1, int(max_image_side) - (2 * border_px))
    preview.thumbnail((inner_max_side, inner_max_side), PILImage.Resampling.LANCZOS)
    return PILImageOps.expand(preview, border=border_px, fill=(0, 0, 0))


def _configure_sheet_columns(sheet: Any, fields: List[str], *, start_col: int = 1) -> None:
    """Apply configured column widths for a field sequence."""
    for idx, field in enumerate(fields, start=start_col):
        letter = get_column_letter(idx)
        sheet.column_dimensions[letter].width = _COLUMN_WIDTHS_BY_FIELD.get(field, 20)


def _populate_task_review_sheet(
    sheet: Any,
    *,
    rows: List[Dict[str, Any]],
    out_root: Path,
    max_image_side: int,
    image_buffers: List[io.BytesIO],
) -> None:
    """Fill one task-style review sheet with preview image columns and row metadata."""
    headers = ["image", "evidence_image", *[_FIELD_LABELS[field] for field in _TASK_SHEET_FIELDS]]
    sheet.append(headers)

    bold = Font(bold=True)
    for column in range(1, len(headers) + 1):
        sheet.cell(row=1, column=column).font = bold

    sheet.column_dimensions["A"].width = _PREVIEW_COLUMN_WIDTH
    sheet.column_dimensions["B"].width = _PREVIEW_COLUMN_WIDTH
    _configure_sheet_columns(sheet, _TASK_SHEET_FIELDS, start_col=3)

    wrap_top = Alignment(wrap_text=True, vertical="top")
    for row_idx, row in enumerate(rows, start=2):
        preview_height = 0
        image_path = out_root / str(row["image_path"])
        if image_path.exists():
            with PILImage.open(image_path) as source:
                source_rgb = source.convert("RGB")
                preview = _build_preview_image(source_rgb, max_image_side=max_image_side)
                evidence_overlay = render_evidence_overlay(
                    source_rgb,
                    evidence_type=str(row.get("_overlay_evidence_type", row.get("evidence_type", ""))),
                    evidence_value=row.get("_overlay_evidence_value", row.get("answer_evidence")),
                )
                evidence_preview = _build_preview_image(evidence_overlay, max_image_side=max_image_side)
                preview_height = max(int(preview.height), int(evidence_preview.height))

                base_buffer = io.BytesIO()
                preview.save(base_buffer, format="PNG")
                base_buffer.seek(0)
                image_buffers.append(base_buffer)

                evidence_buffer = io.BytesIO()
                evidence_preview.save(evidence_buffer, format="PNG")
                evidence_buffer.seek(0)
                image_buffers.append(evidence_buffer)

            xl_img = XLImage(base_buffer)
            xl_img.anchor = f"A{row_idx}"
            sheet.add_image(xl_img)

            evidence_img = XLImage(evidence_buffer)
            evidence_img.anchor = f"B{row_idx}"
            sheet.add_image(evidence_img)

        for col_idx, field in enumerate(_TASK_SHEET_FIELDS, start=3):
            cell = sheet.cell(row=row_idx, column=col_idx, value=_field_value_for_sheet(row, field))
            if field in _WRAP_FIELDS:
                cell.alignment = wrap_top
        sheet.row_dimensions[row_idx].height = max(60, float(preview_height) * 0.75)

    sheet.freeze_panes = "C2"


def _sanitize_sheet_title(raw: str) -> str:
    """Return an Excel-safe sheet title."""
    title = re.sub(r"[\\/*?:\[\]]+", "_", str(raw).strip())
    if not title:
        return "task"
    return title[:31]


def _dedupe_sheet_title(base: str, used: set[str]) -> str:
    """Return a unique sheet title under Excel's 31-character limit."""
    candidate = _sanitize_sheet_title(base)
    if candidate not in used:
        used.add(candidate)
        return candidate

    suffix_index = 2
    while True:
        suffix = f"_{suffix_index}"
        trimmed = candidate[: max(1, 31 - len(suffix))]
        attempt = f"{trimmed}{suffix}"
        if attempt not in used:
            used.add(attempt)
            return attempt
        suffix_index += 1


def _write_task_excel(
    rows: List[Dict[str, Any]],
    path: Path,
    *,
    out_root: Path,
    max_image_side: int = _TASK_PREVIEW_MAX_SIDE,
) -> None:
    """Write per-task sample review workbook with embedded preview images."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "samples"
    image_buffers: List[io.BytesIO] = []
    _populate_task_review_sheet(
        sheet,
        rows=rows,
        out_root=out_root,
        max_image_side=int(max_image_side),
        image_buffers=image_buffers,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _write_domain_combined_excel(
    rows_by_task: Dict[str, List[Dict[str, Any]]],
    path: Path,
    *,
    out_root: Path,
    max_image_side: int = _TASK_PREVIEW_MAX_SIDE,
) -> None:
    """Write one domain-level combined workbook with one task-format sheet per task."""
    workbook = Workbook()
    used_titles: set[str] = set()
    image_buffers: List[io.BytesIO] = []
    sorted_task_ids = sorted(rows_by_task.keys())
    for idx, task_id in enumerate(sorted_task_ids):
        title = _dedupe_sheet_title(task_id, used_titles)
        if idx == 0:
            sheet = workbook.active
            sheet.title = title
        else:
            sheet = workbook.create_sheet(title=title)
        _populate_task_review_sheet(
            sheet,
            rows=rows_by_task[task_id],
            out_root=out_root,
            max_image_side=int(max_image_side),
            image_buffers=image_buffers,
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _as_float_number(value: Any) -> float | None:
    """Parse plain numeric values when possible."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            return float(text)
        except Exception:
            return None
    return None


def _answer_numeric_for_binning(value: Any) -> float | None:
    """Parse one answer value into a numeric scalar for distribution binning."""
    numeric = _as_float_number(value)
    if numeric is not None:
        return float(numeric)
    if not isinstance(value, str):
        return None
    text = value.strip().replace(" ", "")
    if not text:
        return None
    lowered = text.lower().replace("*", "")
    if lowered in {"π", "pi"}:
        return 1.0
    match = re.fullmatch(r"([+-]?(?:\d+(?:\.\d+)?))(?:π|pi)", lowered, flags=re.IGNORECASE)
    if match is None:
        return None
    try:
        return float(match.group(1))
    except Exception:
        return None


def _format_numeric_label(value: float) -> str:
    """Format one numeric boundary with compact deterministic precision."""
    return f"{float(value):.6g}"


def _normalize_probability_map(raw: Any) -> Dict[str, float]:
    """Normalize one probability/weight map into a stable probability map."""
    if not isinstance(raw, Mapping):
        return {}
    parsed: Dict[str, float] = {}
    for key, value in raw.items():
        numeric = _as_float_number(value)
        if numeric is None or (not math.isfinite(float(numeric))) or float(numeric) <= 0.0:
            continue
        parsed[str(key)] = float(numeric)
    total = float(sum(parsed.values()))
    if total <= 0.0:
        return {}
    return {key: (float(parsed[key]) / total) for key in sorted(parsed.keys())}


def _probability_maps_match(left: Mapping[str, float], right: Mapping[str, float], *, tol: float = 1e-9) -> bool:
    """Return true when two probability maps have equal supports and values within tolerance."""
    left_keys = set(str(key) for key in left.keys())
    right_keys = set(str(key) for key in right.keys())
    if left_keys != right_keys:
        return False
    for key in left_keys:
        if abs(float(left[str(key)]) - float(right[str(key)])) > float(tol):
            return False
    return True


def _extract_query_id_distribution_hints(trace_payload: Mapping[str, Any]) -> Dict[str, Any]:
    """Extract optional query-id probability hints from one task output trace payload."""
    query_spec = trace_payload.get("query_spec", {}) if isinstance(trace_payload, Mapping) else {}
    query_params = query_spec.get("params", {}) if isinstance(query_spec, Mapping) else {}
    execution_trace = trace_payload.get("execution_trace", {}) if isinstance(trace_payload, Mapping) else {}

    variant_probabilities = _normalize_probability_map(
        execution_trace.get("query_id_probabilities")
        if isinstance(execution_trace, Mapping)
        else {}
    )
    if not variant_probabilities:
        variant_probabilities = _normalize_probability_map(
            query_params.get("query_id_probabilities")
            if isinstance(query_params, Mapping)
            else {}
        )

    source_kind = ""
    if isinstance(execution_trace, Mapping) and execution_trace.get("source_kind") is not None:
        source_kind = str(execution_trace.get("source_kind"))

    source_kind_probabilities = _normalize_probability_map(
        execution_trace.get("source_kind_probabilities")
        if isinstance(execution_trace, Mapping)
        else {}
    )
    if not source_kind_probabilities:
        source_kind_probabilities = _normalize_probability_map(
            query_params.get("source_kind_probabilities")
            if isinstance(query_params, Mapping)
            else {}
        )

    option_labels: List[str] = []
    for key in ("answer_option_labels", "option_labels", "options"):
        raw = execution_trace.get(key) if isinstance(execution_trace, Mapping) else None
        if not isinstance(raw, list):
            continue
        labels = [str(item).strip() for item in raw if str(item).strip()]
        if labels:
            option_labels = labels
            break

    return {
        "query_id_probabilities": variant_probabilities,
        "source_kind": str(source_kind),
        "source_kind_probabilities": source_kind_probabilities,
        "answer_option_labels": list(option_labels),
    }


def _resolve_query_id_expected_probabilities(
    rows_by_query_id: Mapping[str, List[Dict[str, Any]]],
) -> Tuple[Dict[str, float], str]:
    """Resolve expected variant probabilities from trace metadata or fallback uniform."""
    observed_variants = sorted(str(key) for key in rows_by_query_id.keys())
    if not observed_variants:
        return {}, "no_variants"

    direct_maps: List[Dict[str, float]] = []
    for rows in rows_by_query_id.values():
        for row in rows:
            parsed = _normalize_probability_map(row.get("_query_id_probabilities"))
            if parsed:
                direct_maps.append(parsed)
    if direct_maps:
        first = direct_maps[0]
        if all(_probability_maps_match(first, item) for item in direct_maps[1:]):
            return dict(first), "query_id_probabilities"

    source_probability_maps: List[Dict[str, float]] = []
    source_to_variant: Dict[str, str] = {}
    source_mapping_valid = True
    for variant, rows in rows_by_query_id.items():
        for row in rows:
            source_kind = str(row.get("_source_kind", "")).strip()
            if source_kind:
                existing = source_to_variant.get(source_kind)
                if existing is None:
                    source_to_variant[source_kind] = str(variant)
                elif str(existing) != str(variant):
                    source_mapping_valid = False
            parsed = _normalize_probability_map(row.get("_source_kind_probabilities"))
            if parsed:
                source_probability_maps.append(parsed)
    if source_mapping_valid and source_probability_maps:
        first_source_map = source_probability_maps[0]
        if all(_probability_maps_match(first_source_map, item) for item in source_probability_maps[1:]):
            aggregated: Dict[str, float] = defaultdict(float)
            for source_kind, probability in first_source_map.items():
                variant = source_to_variant.get(str(source_kind))
                if variant is None:
                    continue
                aggregated[str(variant)] += float(probability)
            normalized_aggregated = _normalize_probability_map(aggregated)
            if normalized_aggregated:
                return normalized_aggregated, "derived_from_source_kind_probabilities"

    uniform_probability = 1.0 / float(len(observed_variants))
    return {variant: float(uniform_probability) for variant in observed_variants}, "uniform_fallback"


def _finalize_binned_check(
    *,
    bins: List[Dict[str, Any]],
    sample_count: int,
) -> Dict[str, Any]:
    """Compute z-score based distribution diagnostics for one set of bins."""
    out = {
        "status": "skipped",
        "reason": "",
        "bin_count": int(len(bins)),
        "sample_count": int(sample_count),
        "sigma_threshold": float(_DISTRIBUTION_SIGMA_THRESHOLD),
        "recommended_sample_count": int(_DISTRIBUTION_REVIEW_SAMPLE_COUNT),
        "bins": list(bins),
        "metrics": {
            "max_abs_delta_count": 0.0,
            "max_abs_delta_ratio": 0.0,
            "max_bin_z": 0.0,
        },
    }
    if not bins:
        out["reason"] = "empty_bins"
        return out
    max_abs_delta_count = 0.0
    max_abs_delta_ratio = 0.0
    max_bin_z = 0.0
    for bin_entry in bins:
        probability = float(bin_entry.get("expected_probability", 0.0))
        count = int(bin_entry.get("count", 0))
        expected_count = float(sample_count) * float(probability)
        delta_count = float(count) - float(expected_count)
        sigma = math.sqrt(float(sample_count) * float(probability) * max(0.0, 1.0 - float(probability)))
        z_score = (abs(float(delta_count)) / float(sigma)) if float(sigma) > 0.0 else 0.0
        bin_entry["expected_count"] = float(expected_count)
        bin_entry["delta_count"] = float(delta_count)
        bin_entry["z_score"] = float(z_score)
        max_abs_delta_count = max(float(max_abs_delta_count), abs(float(delta_count)))
        if int(sample_count) > 0:
            max_abs_delta_ratio = max(float(max_abs_delta_ratio), abs(float(delta_count)) / float(sample_count))
        max_bin_z = max(float(max_bin_z), float(z_score))

    out["metrics"] = {
        "max_abs_delta_count": float(max_abs_delta_count),
        "max_abs_delta_ratio": float(max_abs_delta_ratio),
        "max_bin_z": float(max_bin_z),
    }
    out["status"] = "pass" if float(max_bin_z) <= float(_DISTRIBUTION_SIGMA_THRESHOLD) else "fail"
    return out


def _build_answer_distribution_for_variant(rows: Sequence[Dict[str, Any]]) -> Dict[str, Any]:
    """Build one answer-distribution report for a query-id slice."""
    answer_counter = Counter(_json_cell(row.get("answer_value")) for row in rows)
    report: Dict[str, Any] = {
        "accepted_samples": int(len(rows)),
        "answer_counts": dict(sorted(answer_counter.items())),
        "unique_answers": int(len(answer_counter)),
    }

    numeric_values: List[float] = []
    option_signatures: List[Tuple[str, ...]] = []
    all_numeric = True
    for row in rows:
        numeric = _answer_numeric_for_binning(row.get("answer_value"))
        if numeric is None:
            all_numeric = False
        else:
            numeric_values.append(float(numeric))
        raw_options = row.get("_answer_option_labels")
        if isinstance(raw_options, list):
            labels = tuple(str(item).strip() for item in raw_options if str(item).strip())
            if labels:
                option_signatures.append(labels)

    if all_numeric and numeric_values:
        answer_min = float(min(numeric_values))
        answer_max = float(max(numeric_values))
        answer_span = float(answer_max - answer_min)
        if answer_span <= 0.0:
            bin_count = 1
        else:
            bin_count = max(1, int(min(_DISTRIBUTION_MAX_NUMERIC_BINS, math.floor(answer_span))))
        bin_width = (float(answer_span) / float(bin_count)) if int(bin_count) > 0 else 0.0
        bins: List[Dict[str, Any]] = []
        for idx in range(int(bin_count)):
            lower = float(answer_min + (float(idx) * float(bin_width))) if answer_span > 0.0 else float(answer_min)
            if idx == int(bin_count) - 1:
                upper = float(answer_max)
                label = f"[{_format_numeric_label(lower)}, {_format_numeric_label(upper)}]"
            else:
                upper = float(answer_min + (float(idx + 1) * float(bin_width)))
                label = f"[{_format_numeric_label(lower)}, {_format_numeric_label(upper)})"
            bins.append(
                {
                    "bin_id": int(idx),
                    "label": str(label),
                    "range_min": float(lower),
                    "range_max": float(upper),
                    "count": 0,
                    "expected_probability": (1.0 / float(bin_count)),
                }
            )
        if int(bin_count) == 1:
            bins[0]["count"] = int(len(numeric_values))
        else:
            for value in numeric_values:
                if value >= answer_max:
                    bin_index = int(bin_count) - 1
                else:
                    width = float(bin_width) if float(bin_width) > 0.0 else 1.0
                    ratio = (float(value) - float(answer_min)) / float(width)
                    bin_index = max(0, min(int(bin_count) - 1, int(math.floor(ratio))))
                bins[int(bin_index)]["count"] = int(bins[int(bin_index)]["count"]) + 1

        report["answer_numeric"] = {
            "count": int(len(numeric_values)),
            "min": float(answer_min),
            "max": float(answer_max),
            "mean": float(sum(numeric_values) / float(len(numeric_values))),
            "binning": {
                "kind": "numeric_equal_width",
                "bin_count": int(bin_count),
                "range_min": float(answer_min),
                "range_max": float(answer_max),
                "range_span": float(answer_span),
            },
        }
        report["distribution_check"] = _finalize_binned_check(
            bins=bins,
            sample_count=int(len(numeric_values)),
        )
        return report

    category_labels: List[str] = []
    if option_signatures:
        first_signature = option_signatures[0]
        if all(signature == first_signature for signature in option_signatures[1:]):
            category_labels = [str(label) for label in first_signature]
    if not category_labels:
        category_labels = sorted(answer_counter.keys())

    if not category_labels:
        report["distribution_check"] = _finalize_binned_check(bins=[], sample_count=0)
        return report

    bin_probability = 1.0 / float(len(category_labels))
    bins = [
        {
            "bin_id": int(index),
            "label": str(label),
            "count": int(answer_counter.get(str(label), 0)),
            "expected_probability": float(bin_probability),
        }
        for index, label in enumerate(category_labels)
    ]
    report["distribution_check"] = _finalize_binned_check(
        bins=bins,
        sample_count=int(sum(answer_counter.values())),
    )
    report["answer_categorical"] = {
        "count": int(sum(answer_counter.values())),
        "labels": [str(label) for label in category_labels],
        "binning": {
            "kind": "categorical",
            "bin_count": int(len(category_labels)),
        },
    }
    return report


def _build_variant_distribution_report(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Build query-id and answer-bin distribution reports for one task."""
    rows_by_query_id: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for row in rows:
        rows_by_query_id[str(row.get("query_id", "default"))].append(row)

    expected_query_id_probabilities, expected_query_id_source = _resolve_query_id_expected_probabilities(rows_by_query_id)
    query_id_labels = sorted(expected_query_id_probabilities.keys() or rows_by_query_id.keys())
    variant_counts = {str(variant): int(len(rows_by_query_id.get(str(variant), []))) for variant in query_id_labels}
    query_id_bins = [
        {
            "bin_id": int(index),
            "label": str(variant),
            "count": int(variant_counts.get(str(variant), 0)),
            "expected_probability": float(expected_query_id_probabilities.get(str(variant), 0.0)),
        }
        for index, variant in enumerate(query_id_labels)
    ]
    if query_id_bins and sum(float(item.get("expected_probability", 0.0)) for item in query_id_bins) <= 0.0:
        uniform_probability = 1.0 / float(len(query_id_bins))
        for item in query_id_bins:
            item["expected_probability"] = float(uniform_probability)
        expected_query_id_probabilities = {str(item["label"]): float(uniform_probability) for item in query_id_bins}
        expected_query_id_source = "uniform_fallback"

    variant_distribution_check = _finalize_binned_check(
        bins=query_id_bins,
        sample_count=int(len(rows)),
    )
    variant_distribution_report = {
        "accepted_samples": int(len(rows)),
        "expected_probabilities": dict(sorted(expected_query_id_probabilities.items())),
        "expected_source": str(expected_query_id_source),
        "check": variant_distribution_check,
    }

    per_variant: Dict[str, Any] = {}
    checks_run = 0
    checks_failed = 0
    checks_skipped = 0

    if str(variant_distribution_check.get("status")) in {"pass", "fail"}:
        checks_run += 1
        if str(variant_distribution_check.get("status")) == "fail":
            checks_failed += 1
    else:
        checks_skipped += 1

    for query_id in sorted(rows_by_query_id.keys()):
        variant_report = _build_answer_distribution_for_variant(rows_by_query_id[query_id])
        distribution_check = variant_report.get("distribution_check", {})
        status = str(distribution_check.get("status", "skipped"))
        if status in {"pass", "fail"}:
            checks_run += 1
            if status == "fail":
                checks_failed += 1
        else:
            checks_skipped += 1
        per_variant[str(query_id)] = variant_report

    return {
        "thresholds": {
            "recommended_sample_count": int(_DISTRIBUTION_REVIEW_SAMPLE_COUNT),
            "sigma_threshold": float(_DISTRIBUTION_SIGMA_THRESHOLD),
            "max_numeric_bins": int(_DISTRIBUTION_MAX_NUMERIC_BINS),
        },
        "checks_run": int(checks_run),
        "checks_failed": int(checks_failed),
        "checks_skipped": int(checks_skipped),
        "query_id_distribution": variant_distribution_report,
        "per_query_id": per_variant,
    }


def _generate_samples_for_task(
    *,
    task_id: str,
    out_root: Path,
    count: int,
    base_seed: int,
    max_attempts_per_instance: int,
    image_format: str,
    params: Dict[str, Any],
) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    """Generate image/data sample artifacts for one task id."""
    task = create_task(task_id)
    task_dir = out_root / task.domain / task.task_group / task.task_id
    image_dir = task_dir / "images"
    data_dir = task_dir / "data"
    image_dir.mkdir(parents=True, exist_ok=True)
    data_dir.mkdir(parents=True, exist_ok=True)

    accepted = 0
    attempted_candidates = 0
    max_candidates = max(int(count) * 20, int(count))
    rejections: Dict[str, int] = {}
    rows: List[Dict[str, Any]] = []
    accepted_by_variant: Dict[str, int] = defaultdict(int)

    def _attempt_sample(*, seed_namespace: str, seed_index: int) -> None:
        nonlocal accepted
        instance_seed = hash64(int(base_seed), seed_namespace, int(seed_index))
        task_params = dict(params)

        try:
            output = task.generate(
                int(instance_seed),
                params=task_params,
                max_attempts=int(max_attempts_per_instance),
            )
        except Exception as exc:
            reason = type(exc).__name__
            rejections[reason] = rejections.get(reason, 0) + 1
            return

        image_name = f"{accepted:06d}.{image_format}"
        image_path = image_dir / image_name
        output.image.save(image_path, format=image_format.upper())

        rel_image_path = image_path.relative_to(out_root).as_posix()
        data_path = data_dir / f"{accepted:06d}.json"
        rel_data_path = data_path.relative_to(out_root).as_posix()

        prompt_variants = dict(getattr(output, "prompt_variants", {}) or {})
        prompt_answer = str(prompt_variants.get("answer_only", output.prompt))
        prompt_answer_and_evidence = str(prompt_variants.get("answer_and_evidence", output.prompt))

        trace_payload = sanitize_trace_payload_for_public_evidence(
            output.trace_payload if isinstance(output.trace_payload, Mapping) else {},
            evidence_gt=output.evidence_gt,
        )

        payload = {
            "domain": task.domain,
            "task_group": task.task_group,
            "task": task.task_id,
            "sample_index": int(accepted),
            "instance_seed": int(instance_seed),
            "query_id": str(getattr(output, "query_id", "default")),
            "prompt": prompt_answer_and_evidence,
            "prompt_answer": prompt_answer,
            "prompt_answer_and_evidence": prompt_answer_and_evidence,
            "prompt_variants": prompt_variants,
            "answer_gt": output.answer_gt.to_dict(),
            "evidence_gt": output.evidence_gt.to_dict(),
            "complexity": output.complexity.to_dict(),
            "image": {
                "path": rel_image_path,
                "format": image_format,
            },
            "versions": dict(output.task_versions),
            "trace_payload": trace_payload,
        }
        write_json_file(data_path, payload)

        distribution_hints = _extract_query_id_distribution_hints(trace_payload)
        overlay_evidence_type, overlay_evidence_value = resolve_overlay_evidence(
            evidence_type=str(output.evidence_gt.type),
            evidence_value=output.evidence_gt.value,
            trace_payload=trace_payload,
        )
        canonical_answer = {
            "evidence": output.evidence_gt.value,
            "answer": output.answer_gt.value,
        }
        answer_only_ground_truth = {
            "answer": output.answer_gt.value,
        }
        rows.append(
            {
                "domain": task.domain,
                "task_group": task.task_group,
                "task": task.task_id,
                "sample_index": int(accepted),
                "instance_seed": int(instance_seed),
                "query_id": str(getattr(output, "query_id", "default")),
                "image_path": rel_image_path,
                "data_path": rel_data_path,
                "prompt": prompt_answer_and_evidence,
                "prompt_answer": prompt_answer,
                "prompt_answer_only": prompt_answer,
                "prompt_answer_and_evidence": prompt_answer_and_evidence,
                "ground_truth_answer": answer_only_ground_truth,
                "ground_truth_answer_and_evidence": canonical_answer,
                "answer": canonical_answer,
                "answer_type": output.answer_gt.type,
                "answer_value": output.answer_gt.value,
                "evidence_type": output.evidence_gt.type,
                "answer_evidence": output.evidence_gt.value,
                "_overlay_evidence_type": overlay_evidence_type,
                "_overlay_evidence_value": overlay_evidence_value,
                "_query_id_probabilities": dict(distribution_hints.get("query_id_probabilities", {})),
                "_source_kind": str(distribution_hints.get("source_kind", "")),
                "_source_kind_probabilities": dict(distribution_hints.get("source_kind_probabilities", {})),
                "_answer_option_labels": list(distribution_hints.get("answer_option_labels", [])),
            }
        )
        accepted += 1
        accepted_by_variant[str(getattr(output, "query_id", "default"))] += 1

    seed_index = 0
    while accepted < int(count) and seed_index < int(max_candidates):
        _attempt_sample(
            seed_namespace=f"{task_id}:sample_instance_seed",
            seed_index=seed_index,
        )
        seed_index += 1
    attempted_candidates = int(seed_index)
    requested_samples = int(count)

    distribution_report = _build_variant_distribution_report(rows)

    summary = {
        "domain": task.domain,
        "task_group": task.task_group,
        "task": task.task_id,
        "requested_samples": int(requested_samples),
        "accepted_samples": int(accepted),
        "accepted_samples_by_query_id": dict(sorted((k, int(v)) for k, v in accepted_by_variant.items())),
        "attempted_candidates": int(attempted_candidates),
        "max_candidates": int(max_candidates),
        "query_ids": sorted(accepted_by_variant.keys()),
        "rejections_by_error": dict(sorted(rejections.items())),
        "distribution_checks": {
            "checks_run": int(distribution_report.get("checks_run", 0)),
            "checks_failed": int(distribution_report.get("checks_failed", 0)),
            "checks_skipped": int(distribution_report.get("checks_skipped", 0)),
        },
        "answer_distribution": dict(distribution_report),
    }
    write_json_file(task_dir / "summary.json", summary)
    write_json_file(task_dir / "distribution_report.json", distribution_report)
    return summary, rows


def main() -> int:
    """Parse CLI args and generate task sample artifacts."""
    parser = argparse.ArgumentParser(
        description="Generate TRACE task sample images/data, per-task Excel files, and per-domain combined workbooks"
    )
    parser.add_argument("--out", default="task-reviews/generated_samples", help="Output root directory")
    parser.add_argument("--tasks", default="", help="Comma-separated task ids (default: all registered tasks)")
    parser.add_argument("--count", type=int, default=50, help="Samples per task (default: 50)")
    parser.add_argument(
        "--count-per-variant",
        type=int,
        default=0,
        help="Reserved for future variant-targeted sampling (must be 0 for now)",
    )
    parser.add_argument("--seed", type=int, default=0, help="Base sampling seed")
    parser.add_argument("--max-attempts-per-instance", type=int, default=100, help="Max generation attempts per instance")
    parser.add_argument("--image-format", default="png", choices=["png", "jpeg", "jpg"], help="Saved image format")
    parser.add_argument("--params", default="", help="Global task params JSON object")
    parser.add_argument("--task-params", default="", help="Per-task params JSON object mapping task_id -> params object")
    parser.add_argument(
        "--combined-excel",
        default="combined_samples.xlsx",
        help="Combined Excel filename written under each domain directory",
    )
    parser.add_argument(
        "--distribution-check",
        action="store_true",
        help=(
            "Require distribution checks to run/pass for every selected task "
            f"(requires --count >= {_DISTRIBUTION_REVIEW_SAMPLE_COUNT})"
        ),
    )
    parser.add_argument("--clean", action="store_true", help="Clean only selected task directories before generation")
    parser.add_argument("--clean-all", action="store_true", help="Dangerous: remove entire --out before generation")
    args = parser.parse_args()

    if int(args.count) <= 0:
        raise ValueError("--count must be positive")
    if int(args.count_per_query_id) != 0:
        raise ValueError("--count-per-variant is not supported yet; use --count")
    if args.clean and args.clean_all:
        raise ValueError("use only one of --clean or --clean-all")
    if args.clean and not args.tasks.strip():
        raise ValueError("--clean requires --tasks so we only remove explicit task directories")
    if args.distribution_check and int(args.count) < int(_DISTRIBUTION_REVIEW_SAMPLE_COUNT):
        raise ValueError(
            "--distribution-check requires --count >= "
            f"{int(_DISTRIBUTION_REVIEW_SAMPLE_COUNT)} (received {int(args.count)})"
        )

    task_ids = _resolve_task_ids(args.tasks)
    global_params = _parse_json_dict(args.params, arg_name="--params")
    per_task_params = _parse_json_dict(args.task_params, arg_name="--task-params")
    for task_id, params in per_task_params.items():
        if task_id not in TASK_REGISTRY:
            raise ValueError(f"--task-params includes unknown task_id: {task_id}")
        if not isinstance(params, dict):
            raise ValueError(f"--task-params[{task_id}] must be a JSON object")

    out_root = Path(args.out)
    if args.clean_all and out_root.exists():
        shutil.rmtree(out_root)
    elif args.clean and out_root.exists():
        for task_id in task_ids:
            task = create_task(task_id)
            task_dir = out_root / task.domain / task.task_group / task.task_id
            if task_dir.exists():
                shutil.rmtree(task_dir)

    all_rows: List[Dict[str, Any]] = []
    rows_by_domain_task: Dict[str, Dict[str, List[Dict[str, Any]]]] = defaultdict(dict)
    all_summaries: List[Dict[str, Any]] = []
    shortfall_tasks: List[str] = []
    distribution_check_failures: List[str] = []

    for task_id in task_ids:
        params = dict(global_params)
        params.update(dict(per_task_params.get(task_id, {})))
        summary, rows = _generate_samples_for_task(
            task_id=task_id,
            out_root=out_root,
            count=int(args.count),
            base_seed=int(args.seed),
            max_attempts_per_instance=int(args.max_attempts_per_instance),
            image_format=("jpeg" if args.image_format == "jpg" else args.image_format),
            params=params,
        )
        all_summaries.append(summary)
        all_rows.extend(rows)
        rows_by_domain_task[str(summary["domain"])][str(summary["task"])] = list(rows)

        task_dir = out_root / str(summary["domain"]) / str(summary["task_group"]) / str(summary["task"])
        task_excel_path = task_dir / f"{summary['task']}.xlsx"
        _write_task_excel(rows, task_excel_path, out_root=out_root)

        if int(summary["accepted_samples"]) < int(summary["requested_samples"]):
            shortfall_tasks.append(str(summary["task"]))
        print(
            f"[task] {summary['task']}: accepted={summary['accepted_samples']}/{summary['requested_samples']} "
            f"attempts={summary['attempted_candidates']} "
            f"dist(run/fail/skip)="
            f"{summary.get('distribution_checks', {}).get('checks_run', 0)}/"
            f"{summary.get('distribution_checks', {}).get('checks_failed', 0)}/"
            f"{summary.get('distribution_checks', {}).get('checks_skipped', 0)}"
        )
        if args.distribution_check:
            checks = summary.get("distribution_checks", {})
            failed = int(checks.get("checks_failed", 0))
            skipped = int(checks.get("checks_skipped", 0))
            if failed > 0 or skipped > 0:
                distribution_check_failures.append(
                    f"{summary['task']} (failed={failed}, skipped={skipped})"
                )

    domain_combined_excels: Dict[str, str] = {}
    for domain in sorted(rows_by_domain_task.keys()):
        combined_excel_path = out_root / str(domain) / str(args.combined_excel)
        _write_domain_combined_excel(
            rows_by_domain_task[str(domain)],
            combined_excel_path,
            out_root=out_root,
        )
        domain_combined_excels[str(domain)] = combined_excel_path.relative_to(out_root).as_posix()
        print(f"[done] combined excel ({domain}): {combined_excel_path}")

    summary_payload = {
        "num_tasks": len(task_ids),
        "samples_per_task": int(args.count),
        "total_rows": len(all_rows),
        "domain_combined_excels": dict(sorted(domain_combined_excels.items())),
        "tasks": all_summaries,
    }
    write_json_file(out_root / "summary_all_tasks.json", summary_payload)

    if shortfall_tasks:
        print(
            f"[error] sample shortfall for tasks: {', '.join(sorted(shortfall_tasks))}. "
            "Check task summaries under <out>/<domain>/<task_group>/<task>/summary.json",
            file=sys.stderr,
        )
        return 1
    if args.distribution_check and distribution_check_failures:
        print(
            f"[error] distribution check failed for tasks: {', '.join(sorted(distribution_check_failures))}. "
            "Inspect distribution_report.json under <out>/<domain>/<task_group>/<task>/",
            file=sys.stderr,
        )
        return 1

    print(f"[done] wrote {len(all_rows)} samples to {out_root}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
