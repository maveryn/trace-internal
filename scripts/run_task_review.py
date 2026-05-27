#!/usr/bin/env python3
"""Run task review workflows with distribution reports and manual inspection workbook exports."""

from __future__ import annotations

import argparse
import io
import os
import json
from pathlib import Path
import re
import shutil
from typing import Any, Dict, List, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
from PIL import ImageOps as PILImageOps

from trace.core.answer_distribution import evaluate_answer_distribution
from trace.core.evidence_sanitization import sanitize_trace_payload_for_public_evidence
from trace.core.json_io import write_json_file
from trace.core.review_overlays import render_evidence_overlay, resolve_overlay_evidence
from trace.core.seed import hash64
from trace.core.taxonomy import inject_taxonomy_metadata, resolve_task_query_id, resolve_task_taxonomy
from trace.core.task_review_sampling import collect_query_id_samples, generate_random_samples, resolve_review_query_id
from trace.tasks import TASK_REGISTRY, create_task


_PREVIEW_MAX_SIDE = 384
_SCENE_PREVIEW_ROWS_PER_TASK = 100
_EXCEL_HEADERS: List[str] = [
    "image",
    "evidence_image",
    "task",
    "query_id",
    "scene_id",
    "query_id",
    "prompt_answer",
    "ground_truth_answer",
    "prompt_answer_and_evidence",
    "ground_truth_answer_and_evidence",
    "answer_type",
    "evidence_type",
    "instance_seed",
    "image_path",
    "data_path",
]

_EXCEL_COLUMN_WIDTHS: Dict[str, float] = {
    "A": 54,
    "B": 54,
    "C": 28,
    "D": 24,
    "E": 24,
    "F": 24,
    "G": 34,
    "H": 22,
    "I": 40,
    "J": 24,
    "K": 14,
    "L": 16,
    "M": 24,
    "N": 28,
    "O": 28,
}

_WRAP_COLUMNS = {"C", "D", "E", "F", "G", "H", "I", "J", "M", "N", "O"}

_MODEL_STATS_HEADERS: List[str] = [
    "task",
    "combined_status",
    "model",
    "reasons",
    "hard_frac",
    "easy_frac",
    "band_frac",
    "mean_solve_rate",
    "response_cap_rate",
    "prompt_max",
    "prompt_over_limit_count",
    "rollout_count",
    "prompt_count",
    "max_response_length",
    "solve_workbook",
    "calibration_stats",
    "output_dir",
]
_MODEL_STATS_WIDTHS: Dict[str, float] = {
    "A": 56,
    "B": 20,
    "C": 16,
    "D": 34,
    "E": 12,
    "F": 12,
    "G": 12,
    "H": 16,
    "I": 18,
    "J": 12,
    "K": 22,
    "L": 14,
    "M": 14,
    "N": 20,
    "O": 72,
    "P": 72,
    "Q": 72,
}
_MODEL_IDS: Dict[str, str] = {
    "qwen25vl7b": "Qwen/Qwen2.5-VL-7B-Instruct",
    "qwen3vl8b": "Qwen/Qwen3-VL-8B-Instruct",
    "qwen3vl4b": "Qwen/Qwen3-VL-4B-Instruct",
}
_MODEL_RESPONSE_CAP_THRESHOLDS: Dict[str, float] = {
    "qwen25vl7b": 0.25,
    "qwen3vl8b": 0.25,
    "qwen3vl4b": 0.25,
}
_CURRENT_CALIBRATION_MODEL_SLUGS = {"qwen25vl7b"}
_CURRENT_CALIBRATION_BASELINE = "v0"
_DIFFICULTY_TAIL_THRESHOLD = 0.50


def _infer_task_domain(task_id: str, *, task_obj: Any | None = None) -> str:
    """Infer one task domain for review-artifact routing."""

    if task_obj is not None:
        taxonomy = resolve_task_taxonomy(
            str(task_id),
            source_domain=str(getattr(task_obj, "domain", "")),
            source_task_group=str(getattr(task_obj, "task_group", "")),
        )
        return str(taxonomy.domain)
    taxonomy = resolve_task_taxonomy(str(task_id))
    if taxonomy.domain != "unknown":
        return str(taxonomy.domain)
    if task_obj is not None:
        domain = getattr(task_obj, "domain", None)
        if str(domain or "").strip():
            return str(domain)
    match = re.match(r"^task_([a-z0-9]+)_", str(task_id).strip())
    if match:
        return str(match.group(1))
    return "unknown"


def _resolve_task_review_dir(*, out_root: Path, task_id: str, task_obj: Any | None = None) -> Path:
    """Return the canonical domain-scoped review directory for one task."""

    taxonomy = resolve_task_taxonomy(
        str(task_id),
        source_domain=str(getattr(task_obj, "domain", "")) if task_obj is not None else "",
        source_task_group=str(getattr(task_obj, "task_group", "")) if task_obj is not None else "",
    )
    return Path(out_root) / str(taxonomy.domain) / str(taxonomy.scene_id) / str(task_id)


def _resolve_task_ids(raw_tasks: str) -> List[str]:
    """Resolve selected task ids from CLI input."""
    if not str(raw_tasks).strip():
        return sorted(TASK_REGISTRY.keys())
    task_ids = [item.strip() for item in str(raw_tasks).split(",") if item.strip()]
    if not task_ids:
        raise ValueError("--tasks resolved to an empty list")
    unknown = [task_id for task_id in task_ids if task_id not in TASK_REGISTRY]
    if unknown:
        raise ValueError(f"unknown task ids: {', '.join(sorted(unknown))}")
    return sorted(dict.fromkeys(task_ids))


def _task_review_dir(*, out_root: Path, task_id: str, domain: str) -> Path:
    """Return the domain-scoped review directory for one task."""

    taxonomy = resolve_task_taxonomy(str(task_id), source_domain=str(domain))
    return Path(out_root) / str(taxonomy.domain) / str(taxonomy.scene_id) / str(task_id)


def _parse_cli() -> argparse.Namespace:
    """Parse CLI arguments for task-review workflow execution."""
    parser = argparse.ArgumentParser(description="Run TRACE task review workflow")
    parser.add_argument("--tasks", default="", help="Comma-separated task ids (default: all registered tasks)")
    parser.add_argument(
        "--mode",
        choices=["full", "distribution", "inspection"],
        default="full",
        help="full=random+distribution+inspection, distribution=random+distribution, inspection=excel-only",
    )
    parser.add_argument("--seed", type=int, default=0, help="Base seed")
    parser.add_argument("--out-root", default="task-reviews", help="Output root for task review artifacts")
    parser.add_argument("--random-count", type=int, default=100, help="Random sample count per task")
    parser.add_argument("--count-per-variant", type=int, default=100, help="Target samples per query id")
    parser.add_argument(
        "--inspection-count-per-variant",
        type=int,
        default=25,
        help="Source balanced inspection samples per query; used only with --balanced-inspection-by-query",
    )
    parser.add_argument(
        "--balanced-inspection-by-query",
        action="store_true",
        help="Generate inspection workbooks with a fixed sample count per query id instead of --random-count total samples per task",
    )
    parser.add_argument("--max-attempts-per-instance", type=int, default=200, help="Max attempts per instance generation")
    parser.add_argument(
        "--max-total-samples-per-task",
        type=int,
        default=30000,
        help="Safety cap while collecting per-variant samples",
    )
    parser.add_argument(
        "--workers",
        type=int,
        default=max(1, int(os.cpu_count() or 1)),
        help="Thread workers for sampling (default: all visible CPUs)",
    )
    parser.add_argument(
        "--allow-fail",
        action="store_true",
        help="Exit 0 even when one or more distribution checks fail",
    )
    parser.add_argument(
        "--skip-scene-workbooks",
        action="store_true",
        help="Do not refresh per-scene combined inspection workbooks after inspection export",
    )
    return parser.parse_args()


def _as_float(value: Any) -> float | None:
    """Parse one scalar as float when possible."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            return float(text)
        except Exception:
            return None
    return None


def _normalize_probability_map(raw: Any) -> Dict[str, float]:
    """Normalize one probability/weight map into a sorted probability map."""
    if not isinstance(raw, Mapping):
        return {}
    parsed: Dict[str, float] = {}
    for key, value in raw.items():
        numeric = _as_float(value)
        if numeric is None:
            continue
        if float(numeric) <= 0.0:
            continue
        parsed[str(key)] = float(numeric)
    total = float(sum(parsed.values()))
    if total <= 0.0:
        return {}
    return {key: (float(parsed[key]) / total) for key in sorted(parsed.keys())}


def _extract_sampling_axes(output: Any) -> Dict[str, Dict[str, Any]]:
    """Extract sampling-axis observed values and expected probabilities from one output."""
    trace_payload = getattr(output, "trace_payload", {})
    if not isinstance(trace_payload, Mapping):
        trace_payload = {}
    execution_trace = trace_payload.get("execution_trace", {})
    if not isinstance(execution_trace, Mapping):
        execution_trace = {}
    query_spec = trace_payload.get("query_spec", {})
    if not isinstance(query_spec, Mapping):
        query_spec = {}
    query_params = query_spec.get("params", {})
    if not isinstance(query_params, Mapping):
        query_params = {}

    axes: Dict[str, Dict[str, Any]] = {}

    def _set_axis(axis: str, *, observed: Any, expected: Mapping[str, float], source: str) -> None:
        axis_name = str(axis).strip()
        if not axis_name:
            return
        entry = axes.get(axis_name)
        expected_dict = dict(expected)
        if entry is None:
            axes[axis_name] = {
                "observed": "" if observed is None else str(observed),
                "expected_probabilities": expected_dict,
                "expected_source": str(source),
                "expected_conflict": False,
            }
            return
        if entry.get("observed") in {"", None} and observed is not None:
            entry["observed"] = str(observed)
        existing_expected = entry.get("expected_probabilities", {})
        if existing_expected and existing_expected != expected_dict:
            entry["expected_conflict"] = True
        elif (not existing_expected) and expected_dict:
            entry["expected_probabilities"] = expected_dict
            entry["expected_source"] = str(source)

    for source_name, source in (("execution_trace", execution_trace), ("query_params", query_params)):
        if not isinstance(source, Mapping):
            continue
        for key, value in source.items():
            key_text = str(key)
            if not key_text.endswith("_probabilities"):
                continue
            axis = key_text[: -len("_probabilities")]
            probs = _normalize_probability_map(value)
            if not probs:
                continue
            observed = None
            if str(axis) == "query_id":
                observed = getattr(output, "query_id", "")
            elif str(axis) == "query_id":
                observed = str(getattr(output, "query_id", "") or "")
            else:
                observed = execution_trace.get(str(axis), query_params.get(str(axis), ""))
            _set_axis(str(axis), observed=observed, expected=probs, source=str(source_name))

    if str(getattr(output, "query_id", "") or "").strip() in {"", "default"}:
        query_axis = axes.get("query_id", {})
        if isinstance(query_axis, Mapping) and query_axis.get("expected_probabilities"):
            axes["query_id"] = {
                "observed": str(getattr(output, "query_id", "") or query_axis.get("observed", "")),
                "expected_probabilities": dict(query_axis.get("expected_probabilities", {})),
                "expected_source": str(query_axis.get("expected_source", "")),
                "expected_conflict": bool(query_axis.get("expected_conflict", False)),
            }

    if "query_id" not in axes:
        axes["query_id"] = {
            "observed": str(getattr(output, "query_id", "") or ""),
            "expected_probabilities": {},
            "expected_source": "",
            "expected_conflict": False,
        }
    return axes


def _extract_replay_generation_params(output: Any) -> Dict[str, Any]:
    """Extract explicit generation params that can replay one sampled inspection row."""
    trace_payload = getattr(output, "trace_payload", {})
    if not isinstance(trace_payload, Mapping):
        return {}
    query_spec = trace_payload.get("query_spec", {})
    if not isinstance(query_spec, Mapping):
        return {}
    execution_trace = trace_payload.get("execution_trace", {})
    if not isinstance(execution_trace, Mapping):
        execution_trace = {}
    query_params = query_spec.get("params", {})
    if not isinstance(query_params, Mapping):
        return {}

    replay: Dict[str, Any] = {}
    for key, value in query_params.items():
        key_text = str(key)
        if key_text.endswith("_probabilities") or key_text.endswith("_support"):
            continue
        if isinstance(value, (str, int, float, bool)) or value is None:
            replay[key_text] = value
        elif isinstance(value, (list, tuple)):
            replay[key_text] = list(value)
        elif isinstance(value, Mapping):
            replay[key_text] = dict(value)
    query_id = str(getattr(output, "query_id", "") or "")
    review_query_id = resolve_review_query_id(output)
    if review_query_id:
        if query_id.strip() in {"", "default"}:
            internal_query_id = str(execution_trace.get("internal_query_id", "") or "").strip()
            if internal_query_id and internal_query_id != "default":
                replay["query_id"] = str(internal_query_id)
            else:
                replay["query_id"] = str(review_query_id)
        else:
            replay.setdefault("query_id", query_id)
    return replay


def _random_collector(output: Any, instance_seed: int) -> Dict[str, Any]:
    """Collect task-review fields from one generated output."""
    query_id = str(getattr(output, "query_id", "") or "")
    return {
        "instance_seed": int(instance_seed),
        "review_query_id": resolve_review_query_id(output),
        "scene_id": str(getattr(output, "scene_id", "") or ""),
        "query_id": str(getattr(output, "query_id", "") or resolve_task_query_id(query_id=query_id, trace_payload=getattr(output, "trace_payload", {}))),
        "answer_type": str(output.answer_gt.type),
        "answer_value": output.answer_gt.value,
        "sampling_axes": _extract_sampling_axes(output),
        "generation_params": _extract_replay_generation_params(output),
    }


def _answer_collector(output: Any, instance_seed: int) -> Dict[str, Any]:
    """Collect answer-only fields for distribution checks."""
    query_id = str(getattr(output, "query_id", "") or "")
    return {
        "instance_seed": int(instance_seed),
        "review_query_id": resolve_review_query_id(output),
        "scene_id": str(getattr(output, "scene_id", "") or ""),
        "query_id": str(getattr(output, "query_id", "") or resolve_task_query_id(query_id=query_id, trace_payload=getattr(output, "trace_payload", {}))),
        "answer_type": str(output.answer_gt.type),
        "answer_value": output.answer_gt.value,
        "generation_params": _extract_replay_generation_params(output),
    }


def _has_true_variants(variant_counts: Mapping[str, int], expected_probabilities: Mapping[str, float]) -> bool:
    """Return true when a task has more than one meaningful variant."""
    expected_non_empty = [str(key) for key in expected_probabilities.keys() if str(key).strip()]
    if len(expected_non_empty) > 1:
        return True
    observed_non_empty = [str(key) for key, value in variant_counts.items() if str(key).strip() and int(value) > 0]
    return len(observed_non_empty) > 1


def _build_sampling_axis_reports(rows: Sequence[Mapping[str, Any]]) -> Dict[str, Dict[str, Any]]:
    """Build observed/expected sampling distribution report for all discovered axes."""
    axis_names: set[str] = set()
    for row in rows:
        axes = row.get("sampling_axes", {})
        if isinstance(axes, Mapping):
            axis_names.update(str(axis) for axis in axes.keys())

    reports: Dict[str, Dict[str, Any]] = {}
    for axis in sorted(axis_names):
        observed_counts: Dict[str, int] = {}
        expected_map: Dict[str, float] = {}
        expected_conflict = False
        expected_source = ""

        for row in rows:
            axes = row.get("sampling_axes", {})
            axis_entry = axes.get(axis, {}) if isinstance(axes, Mapping) else {}
            observed = axis_entry.get("observed")
            observed_label = str(observed) if observed is not None else ""
            if str(axis) == "query_id" and not observed_label:
                observed_label = str(row.get("review_query_id", "") or row.get("query_id", "") or "")
            observed_counts[observed_label] = int(observed_counts.get(observed_label, 0) + 1)

            raw_expected = axis_entry.get("expected_probabilities", {}) if isinstance(axis_entry, Mapping) else {}
            normalized_expected = _normalize_probability_map(raw_expected)
            if normalized_expected:
                if not expected_map:
                    expected_map = dict(normalized_expected)
                    expected_source = str(axis_entry.get("expected_source", ""))
                elif expected_map != normalized_expected:
                    expected_conflict = True
            if bool(axis_entry.get("expected_conflict", False)):
                expected_conflict = True

        total = int(sum(observed_counts.values()))
        max_count = int(max(observed_counts.values())) if observed_counts else 0
        reports[str(axis)] = {
            "sample_count": int(total),
            "observed_counts": dict(sorted(observed_counts.items(), key=lambda item: item[0])),
            "max_observed_frequency": (float(max_count) / float(total)) if int(total) > 0 else 0.0,
            "expected_probabilities": dict(expected_map),
            "expected_source": str(expected_source),
            "expected_conflict": bool(expected_conflict),
        }

    return reports


def _build_random_review_report(*, task_id: str, rows: Sequence[Mapping[str, Any]]) -> Dict[str, Any]:
    """Build random-sample review report for one task."""
    answer_rows = [
        {
            "answer_type": str(row.get("answer_type", "")),
            "answer_value": row.get("answer_value"),
        }
        for row in rows
    ]
    answer_report = evaluate_answer_distribution(answer_rows)
    sampling_axes = _build_sampling_axis_reports(rows)
    variant_axis = sampling_axes.get("query_id", {})
    variant_counts = variant_axis.get("observed_counts", {}) if isinstance(variant_axis, Mapping) else {}
    variant_expected = variant_axis.get("expected_probabilities", {}) if isinstance(variant_axis, Mapping) else {}
    has_variants = _has_true_variants(variant_counts, variant_expected)

    return {
        "task_id": str(task_id),
        "sample_count": int(len(rows)),
        "answer_distribution": answer_report,
        "sampling_axes": sampling_axes,
        "query_id_distribution": {
            "has_variants": bool(has_variants),
            "status": "reported" if bool(has_variants) else "skipped_no_variants",
            "observed_counts": dict(variant_counts),
            "expected_probabilities": dict(variant_expected),
            "expected_source": str(variant_axis.get("expected_source", "")) if isinstance(variant_axis, Mapping) else "",
            "expected_conflict": bool(variant_axis.get("expected_conflict", False)) if isinstance(variant_axis, Mapping) else False,
        },
    }


def _evaluate_rows(rows: Sequence[Mapping[str, Any]]) -> Dict[str, Any]:
    """Evaluate standard answer-distribution checks for one row slice."""
    answer_rows = [
        {
            "answer_type": str(row.get("answer_type", "")),
            "answer_value": row.get("answer_value"),
        }
        for row in rows
    ]
    if not answer_rows:
        return {
            "sample_count": 0,
            "unique_answers": 0,
            "max_answer_count": 0,
            "max_answer_frequency": 0.0,
            "checks": {
                "min_unique_answers": {"pass": False},
                "max_answer_frequency": {"pass": False},
                "max_five_bin_frequency": {"pass": None, "observed": None, "status": "not_reported_no_samples"},
            },
            "pass": False,
        }
    return evaluate_answer_distribution(answer_rows)


def _build_distribution_review_report(
    *,
    task_id: str,
    task_group: str,
    domain: str,
    scene_id: str,
    random_rows: Sequence[Mapping[str, Any]],
    random_report: Mapping[str, Any],
    variant_rows: Mapping[str, List[Dict[str, Any]]] | None,
    query_id_collection_meta: Mapping[str, Any] | None,
) -> Dict[str, Any]:
    """Build distribution review report from random and per-variant samples."""
    variant_distribution = random_report.get("query_id_distribution", {})
    has_variants = bool(variant_distribution.get("has_variants", False))

    if not has_variants:
        single_report = _evaluate_rows(random_rows)
        return {
            "task_id": str(task_id),
            "domain": str(domain),
            "task_group": str(task_group),
            "scene_id": str(scene_id),
            "mode": "single_sample",
            "has_query_ids": False,
            "overall": dict(single_report),
            "per_query_id": {"": dict(single_report)},
            "failed_variants": [""] if not bool(single_report.get("pass", False)) else [],
            "incomplete_query_ids": [],
            "pass": bool(single_report.get("pass", False)),
            "query_id_distribution": dict(variant_distribution),
            "sampling_axes": dict(random_report.get("sampling_axes", {})),
        }

    variant_rows = variant_rows or {}
    query_id_collection_meta = query_id_collection_meta or {}
    expected_query_ids = list(query_id_collection_meta.get("expected_query_ids", []))
    if not expected_query_ids:
        expected_query_ids = sorted(variant_rows.keys())
    per_variant: Dict[str, Any] = {}
    combined_rows: List[Dict[str, Any]] = []

    for query_id in expected_query_ids:
        rows = list(variant_rows.get(str(query_id), []))
        combined_rows.extend(rows)
        if rows:
            per_variant[str(query_id)] = _evaluate_rows(rows)

    overall = _evaluate_rows(combined_rows)
    failed_variants = [
        str(query_id)
        for query_id, result in per_variant.items()
        if not bool(result.get("pass", False))
    ]
    incomplete_query_ids = list(query_id_collection_meta.get("incomplete_query_ids", []))
    no_samples_collected = not bool(combined_rows)
    task_pass = bool((not failed_variants) and (not incomplete_query_ids) and (not no_samples_collected))

    return {
        "task_id": str(task_id),
        "domain": str(domain),
        "task_group": str(task_group),
        "scene_id": str(scene_id),
        "mode": "per_variant",
        "has_query_ids": True,
        "overall": overall,
        "per_query_id": per_variant,
        "failed_variants": failed_variants,
        "incomplete_query_ids": incomplete_query_ids,
        "no_samples_collected": bool(no_samples_collected),
        "pass": bool(task_pass),
        "query_id_distribution": dict(variant_distribution),
        "sampling_axes": dict(random_report.get("sampling_axes", {})),
        "collection": {
            "target_count_per_query_id": int(query_id_collection_meta.get("target_count_per_query_id", 0)),
            "total_generated": int(query_id_collection_meta.get("total_generated", 0)),
            "expected_query_ids": list(expected_query_ids),
            "generated_query_id_counts": dict(query_id_collection_meta.get("generated_query_id_counts", {})),
            "collected_query_id_counts": dict(query_id_collection_meta.get("collected_query_id_counts", {})),
            "generation_error_counts": dict(query_id_collection_meta.get("generation_error_counts", {})),
        },
    }


def _build_preview_image(source: PILImage.Image, *, max_image_side: int = _PREVIEW_MAX_SIDE) -> PILImage.Image:
    """Resize and border one preview image for workbook embedding."""
    preview = source.convert("RGB")
    border_px = 1
    inner_max_side = max(1, int(max_image_side) - (2 * border_px))
    preview.thumbnail((inner_max_side, inner_max_side), PILImage.Resampling.LANCZOS)
    return PILImageOps.expand(preview, border=border_px, fill=(0, 0, 0))


def _json_cell(value: Any) -> str:
    """Serialize nested values for workbook cells."""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    return "" if value is None else str(value)


def _sanitize_sheet_title(raw: str) -> str:
    """Return one Excel-safe sheet title."""
    title = re.sub(r"[\\\\/*?:\\[\\]]+", "_", str(raw).strip())
    if not title:
        return "default"
    return title[:31]


def _dedupe_sheet_title(base: str, used: set[str]) -> str:
    """Return one unique Excel-safe sheet title."""
    candidate = _sanitize_sheet_title(base)
    if candidate not in used:
        used.add(candidate)
        return candidate
    index = 2
    while True:
        suffix = f"_{index}"
        trimmed = candidate[: max(1, 31 - len(suffix))]
        attempt = f"{trimmed}{suffix}"
        if attempt not in used:
            used.add(attempt)
            return attempt
        index += 1


def _populate_inspection_sheet(
    sheet: Any,
    *,
    rows: Sequence[Mapping[str, Any]],
    out_root: Path,
    image_buffers: List[io.BytesIO],
) -> None:
    """Populate one inspection sheet with preview images and metadata rows."""
    sheet.append(list(_EXCEL_HEADERS))

    bold = Font(bold=True)
    for column in range(1, len(_EXCEL_HEADERS) + 1):
        sheet.cell(row=1, column=column).font = bold

    for letter, width in _EXCEL_COLUMN_WIDTHS.items():
        sheet.column_dimensions[str(letter)].width = float(width)

    wrap_top = Alignment(wrap_text=True, vertical="top")

    for row_idx, row in enumerate(rows, start=2):
        image_path = out_root / str(row.get("image_path", ""))
        if image_path.exists():
            with PILImage.open(image_path) as source:
                source_rgb = source.convert("RGB")
                preview = _build_preview_image(source_rgb)
                evidence_preview = _build_preview_image(
                    render_evidence_overlay(
                        source_rgb,
                        evidence_type=str(row.get("overlay_evidence_type", row.get("evidence_type", ""))),
                        evidence_value=row.get("overlay_evidence_value", row.get("answer_evidence")),
                    )
                )
                row_height = max(int(preview.height), int(evidence_preview.height))

                preview_buffer = io.BytesIO()
                preview.save(preview_buffer, format="PNG")
                preview_buffer.seek(0)
                image_buffers.append(preview_buffer)

                evidence_buffer = io.BytesIO()
                evidence_preview.save(evidence_buffer, format="PNG")
                evidence_buffer.seek(0)
                image_buffers.append(evidence_buffer)

            preview_img = XLImage(preview_buffer)
            preview_img.anchor = f"A{row_idx}"
            sheet.add_image(preview_img)

            evidence_img = XLImage(evidence_buffer)
            evidence_img.anchor = f"B{row_idx}"
            sheet.add_image(evidence_img)

            sheet.row_dimensions[row_idx].height = max(60, float(row_height) * 0.75)
        else:
            sheet.row_dimensions[row_idx].height = 60

        values = [
            row.get("task", ""),
            row.get("query_id", ""),
            row.get("scene_id", ""),
            row.get("query_id", ""),
            row.get("prompt_answer", row.get("prompt_answer_only", "")),
            _json_cell(row.get("ground_truth_answer", row.get("answer_only"))),
            row.get("prompt_answer_and_evidence", row.get("prompt", "")),
            _json_cell(row.get("ground_truth_answer_and_evidence", row.get("answer"))),
            row.get("answer_type", ""),
            row.get("evidence_type", ""),
            int(row.get("instance_seed", 0)),
            row.get("image_path", ""),
            row.get("data_path", ""),
        ]
        for offset, value in enumerate(values, start=3):
            cell = sheet.cell(row=row_idx, column=offset, value=value)
            if get_column_letter(offset) in _WRAP_COLUMNS:
                cell.alignment = wrap_top

    sheet.freeze_panes = "C2"


def _write_inspection_excel(
    rows_by_query_id: Mapping[str, Sequence[Mapping[str, Any]]],
    path: Path,
    *,
    out_root: Path,
) -> Dict[str, str]:
    """Write one inspection workbook with one sheet per query id."""
    workbook = Workbook()
    image_buffers: List[io.BytesIO] = []
    used_titles: set[str] = set()
    variant_to_sheet: Dict[str, str] = {}

    sorted_variants = sorted(str(variant) for variant in rows_by_query_id.keys())
    if not sorted_variants:
        sorted_variants = [""]

    for index, query_id in enumerate(sorted_variants):
        base_title = str(query_id).strip() or "default"
        sheet_title = _dedupe_sheet_title(base_title, used_titles)
        if int(index) == 0:
            sheet = workbook.active
            sheet.title = sheet_title
        else:
            sheet = workbook.create_sheet(title=sheet_title)
        variant_to_sheet[str(query_id)] = str(sheet_title)
        _populate_inspection_sheet(
            sheet,
            rows=list(rows_by_query_id.get(str(query_id), [])),
            out_root=out_root,
            image_buffers=image_buffers,
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return variant_to_sheet


def _task_sheet_title(task_id: str, *, domain: str) -> str:
    """Return a compact human-readable sheet title for one task id."""

    task_text = str(task_id).strip()
    prefix = f"task_{str(domain).strip()}_"
    if task_text.startswith(prefix):
        return task_text[len(prefix) :]
    if task_text.startswith("task_"):
        return task_text[len("task_") :]
    return task_text


def _scene_status_candidates(out_root: Path) -> List[Path]:
    """Return candidate calibration status files for one review root."""

    root = Path(out_root)
    candidates = [
        root.parent / "calibration_sweep_status.json",
        Path("plans/calibration_sweep_status.json"),
    ]
    unique: List[Path] = []
    seen: set[str] = set()
    for candidate in candidates:
        key = str(candidate.resolve()) if candidate.exists() else str(candidate)
        if key in seen:
            continue
        seen.add(key)
        unique.append(candidate)
    return unique


def _status_reasons_from_stats(stats: Mapping[str, Any], *, model_slug: str) -> tuple[str, List[str]]:
    """Return calibration status/reasons from stats using current gates."""

    overall = stats.get("overall", {})
    if not isinstance(overall, Mapping):
        overall = {}
    prompt_stats = stats.get("prompt_token_stats", {})
    if not isinstance(prompt_stats, Mapping):
        prompt_stats = {}
    response_stats = stats.get("response_token_stats", {})
    if not isinstance(response_stats, Mapping):
        response_stats = {}

    blocked_reasons: List[str] = []
    if int(prompt_stats.get("over_limit_count") or 0) > 0:
        blocked_reasons.append("prompt_over_limit")
    cap_threshold = float(_MODEL_RESPONSE_CAP_THRESHOLDS.get(str(model_slug), 0.25))
    if float(response_stats.get("cap_rate") or 0.0) > cap_threshold:
        blocked_reasons.append("response_cap_rate")
    if not overall:
        blocked_reasons.append("missing_overall_stats")
    if blocked_reasons:
        return "blocked", blocked_reasons

    tuning_reasons: List[str] = []
    if float(overall.get("hard_frac") or 0.0) >= _DIFFICULTY_TAIL_THRESHOLD:
        tuning_reasons.append("hard_frac")
    if float(overall.get("easy_frac") or 0.0) >= _DIFFICULTY_TAIL_THRESHOLD:
        tuning_reasons.append("easy_frac")
    mean = float(overall.get("mean_solve_rate") or 0.0)
    if mean < 0.10 or mean > 0.75:
        tuning_reasons.append("mean_solve_rate")
    if tuning_reasons:
        return "needs_manual_tuning", tuning_reasons
    return "accepted", []


def _model_stats_row_from_record(
    *,
    task_id: str,
    scene_id: str,
    combined_status: str,
    model_slug: str,
    model_record: Mapping[str, Any],
) -> Dict[str, Any] | None:
    """Build one model stats row from a calibration sweep model record."""

    stats = model_record.get("stats", {})
    if not isinstance(stats, Mapping):
        return None
    overall = stats.get("overall", {})
    if not isinstance(overall, Mapping):
        overall = {}
    prompt_stats = stats.get("prompt_token_stats", {})
    if not isinstance(prompt_stats, Mapping):
        prompt_stats = {}
    response_stats = stats.get("response_token_stats", {})
    if not isinstance(response_stats, Mapping):
        response_stats = {}
    reasons = model_record.get("reasons", [])
    if isinstance(reasons, str):
        reason_text = reasons
    elif isinstance(reasons, Sequence):
        reason_text = ", ".join(str(item) for item in reasons)
    else:
        reason_text = ""

    return {
        "task": str(task_id),
        "scene_id": str(scene_id),
        "combined_status": str(combined_status),
        "model": str(model_slug),
        "model_id": str(model_record.get("model_id", _MODEL_IDS.get(str(model_slug), ""))),
        "model_status": str(model_record.get("status", "")),
        "reasons": reason_text,
        "hard_frac": overall.get("hard_frac"),
        "easy_frac": overall.get("easy_frac"),
        "band_frac": overall.get("band_frac"),
        "mean_solve_rate": overall.get("mean_solve_rate"),
        "response_cap_rate": response_stats.get("cap_rate"),
        "prompt_max": prompt_stats.get("max"),
        "prompt_over_limit_count": prompt_stats.get("over_limit_count"),
        "rollout_count": overall.get("rollout_count"),
        "prompt_count": overall.get("prompt_count"),
        "max_response_length": model_record.get("max_response_length"),
        "solve_workbook": str(model_record.get("solve_workbook", "")),
        "calibration_stats": str(model_record.get("calibration_stats", "")),
        "output_dir": str(model_record.get("output_dir", "")),
    }


def _model_stats_row_from_stats_file(
    *,
    stats_path: Path,
    model_slug: str,
    task_id: str,
    scene_id: str,
    combined_status: str,
) -> Dict[str, Any] | None:
    """Build one model stats row from a calibration_stats.json file."""

    try:
        stats = json.loads(Path(stats_path).read_text(encoding="utf-8"))
    except Exception:
        return None
    if not isinstance(stats, Mapping):
        return None
    config = stats.get("config", {})
    if not isinstance(config, Mapping):
        config = {}
    if str(config.get("calibration_baseline", "")).strip() != _CURRENT_CALIBRATION_BASELINE:
        return None
    artifacts = stats.get("artifacts", {})
    if not isinstance(artifacts, Mapping):
        artifacts = {}
    status, reasons = _status_reasons_from_stats(stats, model_slug=str(model_slug))
    record = {
        "model_id": _MODEL_IDS.get(str(model_slug), str(model_slug)),
        "status": status,
        "reasons": reasons,
        "stats": stats,
        "max_response_length": config.get("max_response_length"),
        "solve_workbook": artifacts.get("solve_workbook", ""),
        "calibration_stats": artifacts.get("calibration_stats", str(stats_path)),
        "output_dir": config.get("probe_output_dir", str(stats_path.parent)),
    }
    return _model_stats_row_from_record(
        task_id=str(task_id),
        scene_id=str(scene_id),
        combined_status=str(combined_status or status),
        model_slug=str(model_slug),
        model_record=record,
    )


def _latest_stats_files_by_task_model(
    *,
    domain: str,
    scene_id: str,
    task_ids: Sequence[str],
) -> Dict[tuple[str, str], Path]:
    """Return latest calibration stats files keyed by (task_id, model_slug)."""

    task_set = {str(task_id) for task_id in task_ids}
    probe_root = Path("rlvr/outputs/calibration/current")
    if not probe_root.exists():
        return {}
    selected: Dict[tuple[str, str], Path] = {}
    priorities: Dict[tuple[str, str], tuple[bool, int, float]] = {}
    pattern = f"*/{str(domain)}/{str(scene_id)}/task_*/100x*_seed*/calibration_stats.json"
    for stats_path in probe_root.glob(pattern):
        try:
            rel = stats_path.relative_to(probe_root)
        except ValueError:
            continue
        if len(rel.parts) < 6:
            continue
        model_slug = str(rel.parts[0])
        task_id = str(rel.parts[3])
        if task_id not in task_set:
            continue
        rollout_match = re.match(r"100x(?P<rollouts>\d+)_seed", str(rel.parts[4]))
        if rollout_match is None:
            continue
        rollout_count = int(rollout_match.group("rollouts"))
        key = (task_id, model_slug)
        mtime = float(stats_path.stat().st_mtime)
        priority = (int(rollout_count) == 24, int(rollout_count), mtime)
        if key not in selected or priority >= priorities.get(key, (False, -1, 0.0)):
            selected[key] = stats_path
            priorities[key] = priority
    return selected


def _load_scene_model_stats_rows(
    *,
    out_root: Path,
    domain: str,
    scene_id: str,
    task_ids: Sequence[str],
) -> List[Dict[str, Any]]:
    """Load per-task/model calibration stats for one scene when available."""

    status: Dict[str, Any] = {}
    for candidate in _scene_status_candidates(Path(out_root)):
        if not candidate.exists():
            continue
        try:
            loaded = json.loads(candidate.read_text(encoding="utf-8"))
        except Exception:
            continue
        if isinstance(loaded, Mapping):
            status = dict(loaded)
            break
    status_config = status.get("config", {}) if isinstance(status.get("config"), Mapping) else {}
    if str(status_config.get("calibration_baseline", "")).strip() != _CURRENT_CALIBRATION_BASELINE:
        status = {}
    records = status.get("tasks", {})
    if not isinstance(records, Mapping):
        records = {}

    row_map: Dict[tuple[str, str], Dict[str, Any]] = {}
    combined_status_by_task: Dict[str, str] = {}
    reviewer_override_tasks: set[str] = set()
    reviewer_override_models: set[tuple[str, str]] = set()
    suppress_stats_fallback_tasks: set[str] = set()
    for task_id in sorted(str(item) for item in task_ids):
        record = records.get(task_id, {})
        if not isinstance(record, Mapping):
            continue
        if str(record.get("domain", domain)) != str(domain):
            continue
        if str(record.get("scene_id", scene_id)) != str(scene_id):
            continue

        combined_status_by_task[task_id] = str(record.get("status", ""))
        if str(record.get("status", "")).strip() in {"dry_run", "reviewed_pending_probe"}:
            suppress_stats_fallback_tasks.add(task_id)
        record_has_reviewer_override = bool(record.get("reviewer_override"))
        if record_has_reviewer_override:
            reviewer_override_tasks.add(task_id)
        models = record.get("models", {})
        if not isinstance(models, Mapping):
            continue
        for model_slug in sorted(str(key) for key in models.keys()):
            if model_slug not in _CURRENT_CALIBRATION_MODEL_SLUGS:
                continue
            model_record = models.get(model_slug, {})
            if not isinstance(model_record, Mapping):
                continue
            row = _model_stats_row_from_record(
                task_id=task_id,
                scene_id=str(record.get("scene_id", scene_id)),
                combined_status=str(record.get("status", "")),
                model_slug=model_slug,
                model_record=model_record,
            )
            if row is not None:
                row_map[(task_id, model_slug)] = row
                if record_has_reviewer_override or bool(model_record.get("reviewer_override")):
                    reviewer_override_models.add((task_id, model_slug))

    for (task_id, model_slug), stats_path in _latest_stats_files_by_task_model(
        domain=str(domain),
        scene_id=str(scene_id),
        task_ids=task_ids,
    ).items():
        if model_slug not in _CURRENT_CALIBRATION_MODEL_SLUGS:
            continue
        if task_id in suppress_stats_fallback_tasks:
            continue
        if (task_id, model_slug) in reviewer_override_models:
            continue
        row = _model_stats_row_from_stats_file(
            stats_path=stats_path,
            model_slug=model_slug,
            task_id=task_id,
            scene_id=str(scene_id),
            combined_status=combined_status_by_task.get(task_id, ""),
        )
        if row is not None:
            row_map[(task_id, model_slug)] = row

    statuses_by_task: Dict[str, List[str]] = {}
    for (task_id, _model_slug), row in row_map.items():
        status_text = str(row.get("model_status", "")).strip()
        if status_text:
            statuses_by_task.setdefault(task_id, []).append(status_text)
    combined_by_task: Dict[str, str] = {}
    for task_id, statuses in statuses_by_task.items():
        if any(status == "blocked" for status in statuses):
            combined_by_task[task_id] = "blocked"
        elif statuses and all(status == "accepted" for status in statuses):
            combined_by_task[task_id] = "accepted"
        elif statuses:
            combined_by_task[task_id] = "needs_manual_tuning"
    for (task_id, _model_slug), row in row_map.items():
        if task_id in combined_by_task:
            row["combined_status"] = combined_by_task[task_id]
        if task_id in reviewer_override_tasks:
            row["combined_status"] = combined_status_by_task.get(task_id, row.get("combined_status", ""))

    return [row_map[key] for key in sorted(row_map.keys())]


def _populate_model_stats_sheet(sheet: Any, rows: Sequence[Mapping[str, Any]]) -> None:
    """Populate the scene-level model stats sheet."""

    sheet.append(list(_MODEL_STATS_HEADERS))
    bold = Font(bold=True)
    for column in range(1, len(_MODEL_STATS_HEADERS) + 1):
        sheet.cell(row=1, column=column).font = bold
    for letter, width in _MODEL_STATS_WIDTHS.items():
        sheet.column_dimensions[str(letter)].width = float(width)

    wrap_top = Alignment(wrap_text=True, vertical="top")
    wrap_columns = {"A", "E", "G", "R", "S", "T"}
    for row_idx, row in enumerate(rows, start=2):
        for column_idx, header in enumerate(_MODEL_STATS_HEADERS, start=1):
            value = row.get(header, "")
            cell = sheet.cell(row=row_idx, column=column_idx, value=value)
            if get_column_letter(column_idx) in wrap_columns:
                cell.alignment = wrap_top
    sheet.freeze_panes = "A2"


def _write_scene_inspection_excel(
    rows_by_task: Mapping[str, Sequence[Mapping[str, Any]]],
    path: Path,
    *,
    out_root: Path,
    domain: str,
) -> tuple[Dict[str, str], str, int]:
    """Write one scene-level workbook with one sheet per task."""

    workbook = Workbook()
    image_buffers: List[io.BytesIO] = []
    used_titles: set[str] = set()
    task_to_sheet: Dict[str, str] = {}

    sorted_tasks = sorted(str(task_id) for task_id in rows_by_task.keys())
    if not sorted_tasks:
        sorted_tasks = ["empty"]

    model_stats_rows = _load_scene_model_stats_rows(
        out_root=out_root,
        domain=str(domain),
        scene_id=str(path.parent.name),
        task_ids=sorted_tasks,
    )
    model_stats_sheet = _dedupe_sheet_title("model_stats", used_titles)
    sheet = workbook.active
    sheet.title = model_stats_sheet
    _populate_model_stats_sheet(sheet, rows=model_stats_rows)

    for task_id in sorted_tasks:
        sheet_title = _dedupe_sheet_title(_task_sheet_title(str(task_id), domain=str(domain)), used_titles)
        sheet = workbook.create_sheet(title=sheet_title)
        task_to_sheet[str(task_id)] = str(sheet_title)
        _populate_inspection_sheet(
            sheet,
            rows=list(rows_by_task.get(str(task_id), [])),
            out_root=out_root,
            image_buffers=image_buffers,
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return task_to_sheet, model_stats_sheet, int(len(model_stats_rows))



def _inspection_rows_from_task_dir(*, out_root: Path, task_dir: Path) -> List[Dict[str, Any]]:
    """Load inspection workbook rows from one task's JSON sidecars."""

    data_root = task_dir / "data"
    if not data_root.exists():
        return []

    rows: List[Dict[str, Any]] = []
    for data_path in sorted(data_root.rglob("*.json")):
        try:
            payload = json.loads(data_path.read_text(encoding="utf-8"))
        except Exception:
            continue
        if not isinstance(payload, Mapping):
            continue

        answer_gt = payload.get("answer_gt", {})
        if not isinstance(answer_gt, Mapping):
            answer_gt = {}
        evidence_gt = payload.get("evidence_gt", {})
        if not isinstance(evidence_gt, Mapping):
            evidence_gt = {}
        image_payload = payload.get("image", {})
        if not isinstance(image_payload, Mapping):
            image_payload = {}
        prompt_variants = payload.get("prompt_variants", {})
        if not isinstance(prompt_variants, Mapping):
            prompt_variants = {}
        trace_payload = payload.get("trace_payload", {})
        if not isinstance(trace_payload, Mapping):
            trace_payload = {}

        evidence_type = str(evidence_gt.get("type", ""))
        evidence_value = evidence_gt.get("value")
        overlay_evidence_type, overlay_evidence_value = resolve_overlay_evidence(
            evidence_type=evidence_type,
            evidence_value=evidence_value,
            trace_payload=trace_payload,
        )

        prompt = str(payload.get("prompt", ""))
        prompt_answer = str(prompt_variants.get("answer_only", prompt))
        prompt_answer_and_evidence = str(prompt_variants.get("answer_and_evidence", prompt))
        image_path = str(image_payload.get("path", ""))
        try:
            rel_data_path = data_path.relative_to(out_root).as_posix()
        except ValueError:
            rel_data_path = data_path.as_posix()

        rows.append(
            {
                "task": str(payload.get("task", task_dir.name)),
                "query_id": str(payload.get("query_id", "")),
                "scene_id": str(payload.get("scene_id", task_dir.parent.name)),
                "query_id": str(payload.get("query_id", "")),
                "prompt": prompt,
                "prompt_answer": prompt_answer,
                "prompt_answer_only": prompt_answer,
                "prompt_answer_and_evidence": prompt_answer_and_evidence,
                "ground_truth_answer": {
                    "answer": answer_gt.get("value"),
                },
                "ground_truth_answer_and_evidence": {
                    "evidence": evidence_value,
                    "answer": answer_gt.get("value"),
                },
                "answer": {
                    "evidence": evidence_value,
                    "answer": answer_gt.get("value"),
                },
                "answer_evidence": evidence_value,
                "answer_type": str(answer_gt.get("type", "")),
                "evidence_type": evidence_type,
                "instance_seed": int(payload.get("instance_seed", 0)),
                "image_path": image_path,
                "data_path": str(rel_data_path),
                "overlay_evidence_type": str(overlay_evidence_type),
                "overlay_evidence_value": overlay_evidence_value,
            }
        )

    return sorted(
        rows,
        key=lambda item: (
            str(item.get("query_id", "")),
            int(item.get("instance_seed", 0)),
            str(item.get("data_path", "")),
        ),
    )


def _scene_preview_rows(
    rows: Sequence[Mapping[str, Any]],
    *,
    domain: str,
    scene_id: str,
    task_id: str,
) -> List[Dict[str, Any]]:
    """Return the deterministic 100-row human-preview sample for a scene sheet."""

    resolved = [dict(row) for row in rows]
    if len(resolved) <= _SCENE_PREVIEW_ROWS_PER_TASK:
        return list(resolved)

    ranked = sorted(
        resolved,
        key=lambda row: hash64(
            20260518,
            f"{domain}/{scene_id}/{task_id}/{row.get('data_path', '')}",
            0,
        ),
    )[:_SCENE_PREVIEW_ROWS_PER_TASK]
    return sorted(
        ranked,
        key=lambda item: (
            str(item.get("query_id", "")),
            int(item.get("instance_seed", 0)),
            str(item.get("data_path", "")),
        ),
    )


def build_scene_review_workbook(*, out_root: Path, domain: str, scene_id: str) -> Dict[str, Any]:
    """Build one combined scene workbook with one sheet per task."""

    scene_dir = Path(out_root) / str(domain) / str(scene_id)
    rows_by_task: Dict[str, List[Dict[str, Any]]] = {}
    task_entries: Dict[str, Dict[str, Any]] = {}

    if not scene_dir.exists():
        return {
            "domain": str(domain),
            "scene_id": str(scene_id),
            "task_count": 0,
            "inspection_count": 0,
            "workbook": "",
            "tasks": {},
        }

    for task_dir in sorted(path for path in scene_dir.iterdir() if path.is_dir() and path.name.startswith("task_")):
        task_id = str(task_dir.name)
        if task_id not in TASK_REGISTRY:
            continue
        task = create_task(task_id)
        taxonomy = resolve_task_taxonomy(
            task_id,
            source_domain=str(getattr(task, "domain", "")),
            source_task_group=str(getattr(task, "task_group", "")),
        )
        if str(taxonomy.domain) != str(domain) or str(taxonomy.scene_id) != str(scene_id):
            continue
        all_rows = _inspection_rows_from_task_dir(out_root=Path(out_root), task_dir=task_dir)
        if not all_rows:
            continue
        rows = _scene_preview_rows(
            all_rows,
            domain=str(domain),
            scene_id=str(scene_id),
            task_id=str(task_id),
        )
        rows_by_task[task_id] = rows
        manifest_path = task_dir / "manifest.json"
        task_entries[task_id] = {
            "inspection_count": int(len(rows)),
            "source_inspection_count": int(len(all_rows)),
            "scene_preview_rows_per_task": int(_SCENE_PREVIEW_ROWS_PER_TASK),
            "task_manifest": str(manifest_path.relative_to(out_root).as_posix()) if manifest_path.exists() else "",
            "task_workbook": str((task_dir / f"{task_id}.xlsx").relative_to(out_root).as_posix())
            if (task_dir / f"{task_id}.xlsx").exists()
            else "",
        }

    workbook_path = scene_dir / "scene_review.xlsx"
    if not rows_by_task:
        if workbook_path.exists():
            workbook_path.unlink()
        manifest = {
            "domain": str(domain),
            "scene_id": str(scene_id),
            "task_count": 0,
            "inspection_count": 0,
            "workbook": "",
            "tasks": {},
        }
        write_json_file(scene_dir / "scene_review_manifest.json", manifest)
        return manifest

    task_sheets, model_stats_sheet, model_stats_count = _write_scene_inspection_excel(
        rows_by_task,
        workbook_path,
        out_root=Path(out_root),
        domain=str(domain),
    )
    for task_id, sheet_name in task_sheets.items():
        task_entries.setdefault(str(task_id), {})["sheet"] = str(sheet_name)

    manifest = {
        "domain": str(domain),
        "scene_id": str(scene_id),
        "task_count": int(len(rows_by_task)),
        "inspection_count": int(sum(len(rows) for rows in rows_by_task.values())),
        "workbook": str(workbook_path.relative_to(out_root).as_posix()),
        "model_stats_sheet": str(model_stats_sheet),
        "model_stats_count": int(model_stats_count),
        "tasks": task_entries,
    }
    write_json_file(scene_dir / "scene_review_manifest.json", manifest)
    return manifest


def build_scene_review_workbooks(*, out_root: Path, scene_keys: Sequence[tuple[str, str]] | None = None) -> Dict[str, Any]:
    """Build combined scene workbooks for selected or all scenes under out_root."""

    root = Path(out_root)
    if scene_keys is None:
        discovered: List[tuple[str, str]] = []
        for domain_dir in sorted(path for path in root.iterdir() if path.is_dir()):
            for scene_dir in sorted(path for path in domain_dir.iterdir() if path.is_dir()):
                discovered.append((str(domain_dir.name), str(scene_dir.name)))
        scene_keys = discovered

    scene_manifests: Dict[str, Any] = {}
    for domain, scene_id in sorted({(str(domain), str(scene_id)) for domain, scene_id in scene_keys}):
        manifest = build_scene_review_workbook(out_root=root, domain=str(domain), scene_id=str(scene_id))
        if int(manifest.get("task_count", 0)) <= 0:
            continue
        scene_key = f"{domain}/{scene_id}"
        scene_manifests[scene_key] = manifest
        print(
            f"[done] scene workbook: {manifest['workbook']} "
            f"(tasks={manifest['task_count']}, rows={manifest['inspection_count']})"
        )
    return scene_manifests


def _safe_query_id_dir_name(query_id: str) -> str:
    """Return filesystem-safe variant directory label."""
    value = str(query_id).strip()
    if not value:
        return "default"
    return re.sub(r"[^a-zA-Z0-9._-]+", "_", value)


def _build_inspection_rows(
    *,
    task_id: str,
    out_root: Path,
    task_dir: Path,
    seed_rows_by_query_id: Mapping[str, Sequence[Mapping[str, Any]]],
    max_attempts_per_instance: int,
) -> Dict[str, Any]:
    """Generate inspection artifacts (images/json/workbook rows) for one task."""
    rows_by_query_id: Dict[str, List[Dict[str, Any]]] = {}
    task = create_task(str(task_id))
    taxonomy = resolve_task_taxonomy(
        str(task_id),
        source_domain=str(getattr(task, "domain", "")),
        source_task_group=str(getattr(task, "task_group", "")),
    )
    for artifact_subdir in ("images", "data"):
        shutil.rmtree(task_dir / artifact_subdir, ignore_errors=True)

    for query_id in sorted(seed_rows_by_query_id.keys()):
        query_id_dir = _safe_query_id_dir_name(str(query_id))
        image_dir = task_dir / "images" / query_id_dir
        data_dir = task_dir / "data" / query_id_dir
        image_dir.mkdir(parents=True, exist_ok=True)
        data_dir.mkdir(parents=True, exist_ok=True)

        seed_rows = list(seed_rows_by_query_id.get(str(query_id), []))
        for index, seed_row in enumerate(seed_rows):
            instance_seed = int(seed_row.get("instance_seed", 0))
            generation_params = dict(seed_row.get("generation_params", {}) or {})
            generation_param_candidates: List[Dict[str, Any]] = []
            if str(query_id).strip():
                forced_task_params = dict(generation_params)
                forced_task_params["query_id"] = str(query_id)
                forced_task_params["query_id"] = str(query_id)
                generation_param_candidates.append(forced_task_params)

                query_only_params = dict(generation_params)
                query_only_params.pop("query_id", None)
                query_only_params["query_id"] = str(query_id)
                query_only_params["query_id"] = str(query_id)
                generation_param_candidates.append(query_only_params)

                replay_params = dict(generation_params)
                replay_params.setdefault("query_id", str(query_id))
                replay_params.setdefault("query_id", str(query_id))
                generation_param_candidates.append(replay_params)
            else:
                # For single-sheet public tasks, trace params often include
                # resolved scene facts plus diagnostic query bookkeeping. Those
                # values are useful in review JSON but are not a stable replay
                # contract for narrowed wrapper tasks, so regenerate from the
                # sampled seed with task defaults.
                generation_param_candidates.append({})

            deduped_generation_param_candidates: List[Dict[str, Any]] = []
            seen_generation_param_candidates: set[str] = set()
            for candidate_params in generation_param_candidates:
                candidate_key = repr(sorted(candidate_params.items()))
                if candidate_key in seen_generation_param_candidates:
                    continue
                seen_generation_param_candidates.add(candidate_key)
                deduped_generation_param_candidates.append(dict(candidate_params))

            output = None
            final_seed = int(instance_seed)
            last_error: Exception | None = None
            candidate_seeds = [int(instance_seed)]
            candidate_seeds.extend(
                int(hash64(int(instance_seed), f"{str(task_id)}|{str(query_id)}|inspection", retry_index))
                for retry_index in range(1, 33)
            )
            for candidate_seed in candidate_seeds:
                for candidate_params in deduped_generation_param_candidates:
                    try:
                        candidate_output = task.generate(
                            int(candidate_seed),
                            params=dict(candidate_params),
                            max_attempts=int(max_attempts_per_instance),
                        )
                    except Exception as exc:
                        last_error = exc
                        continue
                    if str(query_id).strip() and resolve_review_query_id(candidate_output) != str(query_id):
                        continue
                    output = candidate_output
                    final_seed = int(candidate_seed)
                    break
                if output is not None:
                    break
            if output is None:
                raise RuntimeError(
                    f"{task_id} failed to build inspection row for variant {query_id!r} "
                    f"after {len(candidate_seeds)} deterministic seed attempts"
                ) from last_error

            image_path = image_dir / f"{index:04d}.png"
            output.image.save(image_path, format="PNG")
            rel_image_path = image_path.relative_to(out_root).as_posix()

            prompt_variants = dict(getattr(output, "prompt_variants", {}) or {})
            prompt_answer = str(prompt_variants.get("answer_only", output.prompt))
            prompt_answer_and_evidence = str(prompt_variants.get("answer_and_evidence", output.prompt))

            sanitized_trace_payload = sanitize_trace_payload_for_public_evidence(
                output.trace_payload if isinstance(output.trace_payload, Mapping) else {},
                evidence_gt=output.evidence_gt,
            )
            output_variant = str(getattr(output, "query_id", "") or "")
            review_query_id = resolve_review_query_id(output)
            query_id = str(
                getattr(output, "query_id", "")
                or resolve_task_query_id(query_id=output_variant, trace_payload=sanitized_trace_payload)
            )
            sanitized_trace_payload = inject_taxonomy_metadata(
                sanitized_trace_payload,
                task_id=str(task_id),
                taxonomy=taxonomy,
                query_id=query_id,
                registered_domain=str(getattr(task, "domain", "")),
                registered_task_group=str(getattr(task, "task_group", "")),
            )
            data_payload = {
                "task": str(task_id),
                "domain": str(taxonomy.domain),
                "scene_id": str(taxonomy.scene_id),
                "query_id": query_id,
                "instance_seed": int(final_seed),
                "prompt": prompt_answer_and_evidence,
                "prompt_variants": prompt_variants,
                "answer_gt": output.answer_gt.to_dict(),
                "evidence_gt": output.evidence_gt.to_dict(),
                "image": {
                    "path": str(rel_image_path),
                    "format": "png",
                },
                "trace_payload": sanitized_trace_payload,
                "versions": dict(output.task_versions),
            }
            data_path = data_dir / f"{index:04d}.json"
            write_json_file(data_path, data_payload)
            rel_data_path = data_path.relative_to(out_root).as_posix()

            overlay_evidence_type, overlay_evidence_value = resolve_overlay_evidence(
                evidence_type=str(output.evidence_gt.type),
                evidence_value=output.evidence_gt.value,
                trace_payload=sanitized_trace_payload,
            )
            canonical_answer = {
                "evidence": output.evidence_gt.value,
                "answer": output.answer_gt.value,
            }
            answer_only_ground_truth = {
                "answer": output.answer_gt.value,
            }
            variant_rows = rows_by_query_id.setdefault(str(review_query_id), [])
            variant_rows.append(
                {
                    "task": str(task_id),
                    "query_id": str(output_variant),
                    "scene_id": str(taxonomy.scene_id),
                    "query_id": query_id,
                    "prompt": prompt_answer_and_evidence,
                    "prompt_answer": prompt_answer,
                    "prompt_answer_only": prompt_answer,
                    "prompt_answer_and_evidence": prompt_answer_and_evidence,
                    "ground_truth_answer": answer_only_ground_truth,
                    "ground_truth_answer_and_evidence": canonical_answer,
                    "answer": canonical_answer,
                    "answer_evidence": output.evidence_gt.value,
                    "answer_type": str(output.answer_gt.type),
                    "evidence_type": str(output.evidence_gt.type),
                    "instance_seed": int(final_seed),
                    "image_path": str(rel_image_path),
                    "data_path": str(rel_data_path),
                    "overlay_evidence_type": str(overlay_evidence_type),
                    "overlay_evidence_value": overlay_evidence_value,
                }
            )

    inspection_total = 0
    for query_id in sorted(rows_by_query_id.keys()):
        rows_by_query_id[str(query_id)] = sorted(
            list(rows_by_query_id.get(str(query_id), [])),
            key=lambda item: int(item.get("instance_seed", 0)),
        )
        inspection_total += int(len(rows_by_query_id[str(query_id)]))

    workbook_path = task_dir / f"{task_id}.xlsx"
    workbook_sheets = _write_inspection_excel(rows_by_query_id, workbook_path, out_root=out_root)

    manifest = {
        "task_id": str(task_id),
        "inspection_count": int(inspection_total),
        "variants": {
            str(variant): int(len(seed_rows_by_query_id.get(str(variant), [])))
            for variant in sorted(seed_rows_by_query_id.keys())
        },
        "workbook": str(workbook_path.relative_to(out_root).as_posix()),
        "workbook_sheets": dict(workbook_sheets),
    }
    write_json_file(task_dir / "manifest.json", manifest)
    return manifest


def _group_random_inspection_rows(rows: Sequence[Mapping[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
    """Group one public-task random sample for workbook sheets without changing total sample count."""

    grouped: Dict[str, List[Dict[str, Any]]] = {}
    for row in rows:
        review_key = str(
            row.get("review_query_id")
            or row.get("query_id")
            or row.get("query_id")
            or ""
        )
        grouped.setdefault(str(review_key), []).append(dict(row))
    return {str(key): list(value) for key, value in sorted(grouped.items(), key=lambda item: item[0])}


def main() -> int:
    """Entry point for task-review workflow execution."""
    args = _parse_cli()
    if int(args.random_count) <= 0:
        raise ValueError("--random-count must be > 0")
    if int(args.count_per_query_id) <= 0:
        raise ValueError("--count-per-variant must be > 0")
    if int(args.inspection_count_per_query_id) <= 0:
        raise ValueError("--inspection-count-per-variant must be > 0")
    if int(args.max_attempts_per_instance) <= 0:
        raise ValueError("--max-attempts-per-instance must be > 0")
    if int(args.max_total_samples_per_task) <= 0:
        raise ValueError("--max-total-samples-per-task must be > 0")
    if int(args.workers) <= 0:
        raise ValueError("--workers must be > 0")

    task_ids = _resolve_task_ids(str(args.tasks))
    out_root = Path(str(args.out_root)).resolve()
    out_root.mkdir(parents=True, exist_ok=True)

    summary: Dict[str, Any] = {
        "config": {
            "mode": str(args.mode),
            "seed": int(args.seed),
            "tasks": list(task_ids),
            "random_count": int(args.random_count),
            "count_per_query_id": int(args.count_per_query_id),
            "inspection_count_per_query_id": int(args.inspection_count_per_query_id),
            "balanced_inspection_by_query": bool(args.balanced_inspection_by_query),
            "max_attempts_per_instance": int(args.max_attempts_per_instance),
            "max_total_samples_per_task": int(args.max_total_samples_per_task),
            "workers": int(args.workers),
        },
        "tasks": [],
    }

    failed_distribution_tasks: List[str] = []
    touched_scene_keys: set[tuple[str, str]] = set()

    for task_id in task_ids:
        task = create_task(str(task_id))
        taxonomy = resolve_task_taxonomy(
            str(task_id),
            source_domain=str(getattr(task, "domain", "")),
            source_task_group=str(getattr(task, "task_group", "")),
        )
        task_dir = _resolve_task_review_dir(out_root=out_root, task_id=str(task_id), task_obj=task)
        task_dir.mkdir(parents=True, exist_ok=True)

        task_summary: Dict[str, Any] = {
            "task_id": str(task_id),
            "domain": str(taxonomy.domain),
            "scene_id": str(taxonomy.scene_id),
            "source_domain": str(task.domain),
            "task_group": str(task.task_group),
            "reports": {},
        }
        touched_scene_keys.add((str(taxonomy.domain), str(taxonomy.scene_id)))

        random_rows: List[Dict[str, Any]] = []
        random_report: Dict[str, Any] | None = None
        distribution_variant_rows: Dict[str, List[Dict[str, Any]]] | None = None

        if str(args.mode) in {"full", "distribution"}:
            random_rows = generate_random_samples(
                task_id=str(task_id),
                count=int(args.random_count),
                seed=int(args.seed),
                max_attempts_per_instance=int(args.max_attempts_per_instance),
                workers=int(args.workers),
                collector=_random_collector,
            )
            random_report = _build_random_review_report(task_id=str(task_id), rows=random_rows)
            random_path = task_dir / "random_review_100.json"
            write_json_file(random_path, random_report)
            task_summary["reports"]["random_review"] = str(random_path.relative_to(out_root).as_posix())

            has_variants = bool(random_report["query_id_distribution"]["has_variants"])
            variant_rows: Dict[str, List[Dict[str, Any]]] | None = None
            collection_meta: Dict[str, Any] | None = None
            if has_variants:
                collected = collect_query_id_samples(
                    task_id=str(task_id),
                    target_count_per_query_id=int(args.count_per_query_id),
                    seed=int(args.seed),
                    max_attempts_per_instance=int(args.max_attempts_per_instance),
                    max_total_samples_per_task=int(args.max_total_samples_per_task),
                    workers=int(args.workers),
                    collector=_answer_collector,
                )
                variant_rows = {
                    str(variant): list(collected.get("samples_by_query_id", {}).get(str(variant), []))
                    for variant in list(collected.get("expected_query_ids", []))
                }
                distribution_variant_rows = dict(variant_rows)
                collection_meta = {
                    key: collected.get(key)
                    for key in (
                        "target_count_per_query_id",
                        "total_generated",
                        "expected_query_ids",
                        "generated_query_id_counts",
                        "collected_query_id_counts",
                        "incomplete_query_ids",
                        "generation_error_counts",
                    )
                }

            distribution_report = _build_distribution_review_report(
                task_id=str(task_id),
                task_group=str(task.task_group),
                domain=str(taxonomy.domain),
                scene_id=str(taxonomy.scene_id),
                random_rows=random_rows,
                random_report=random_report,
                variant_rows=variant_rows,
                query_id_collection_meta=collection_meta,
            )
            dist_path = task_dir / "distribution_review.json"
            write_json_file(dist_path, distribution_report)
            task_summary["reports"]["distribution_review"] = str(dist_path.relative_to(out_root).as_posix())
            task_summary["distribution_pass"] = bool(distribution_report.get("pass", False))
            if not bool(distribution_report.get("pass", False)):
                failed_distribution_tasks.append(str(task_id))

            status = "PASS" if bool(distribution_report.get("pass", False)) else "FAIL"
            print(f"[{status}] {task_id} distribution review")

        if str(args.mode) in {"full", "inspection"}:
            seed_rows_by_query_id: Dict[str, List[Dict[str, Any]]]
            random_rows_for_inspection: List[Dict[str, Any]] | None = None
            if not bool(args.balanced_inspection_by_query):
                random_rows_for_inspection = list(random_rows)
                if not random_rows_for_inspection:
                    random_rows_for_inspection = generate_random_samples(
                        task_id=str(task_id),
                        count=int(args.random_count),
                        seed=int(args.seed),
                        max_attempts_per_instance=int(args.max_attempts_per_instance),
                        workers=int(args.workers),
                        collector=_answer_collector,
                    )
                seed_rows_by_query_id = _group_random_inspection_rows(random_rows_for_inspection)
            elif random_report is None:
                inspection_collected = collect_query_id_samples(
                    task_id=str(task_id),
                    target_count_per_query_id=int(args.inspection_count_per_query_id),
                    seed=int(args.seed) + 17,
                    max_attempts_per_instance=int(args.max_attempts_per_instance),
                    max_total_samples_per_task=int(args.max_total_samples_per_task),
                    workers=int(args.workers),
                    collector=_answer_collector,
                )
                expected_query_ids = list(inspection_collected.get("expected_query_ids", []))
                if not expected_query_ids:
                    expected_query_ids = sorted(inspection_collected.get("samples_by_query_id", {}).keys())
                seed_rows_by_query_id = {
                    str(variant): list(inspection_collected.get("samples_by_query_id", {}).get(str(variant), []))
                    for variant in expected_query_ids
                }
                if not seed_rows_by_query_id:
                    fallback_rows = generate_random_samples(
                        task_id=str(task_id),
                        count=int(args.inspection_count_per_query_id),
                        seed=int(args.seed) + 17,
                        max_attempts_per_instance=int(args.max_attempts_per_instance),
                        workers=int(args.workers),
                        collector=_answer_collector,
                    )
                    seed_rows_by_query_id = {"": list(fallback_rows)}
            elif bool(random_report["query_id_distribution"]["has_variants"]):
                if distribution_variant_rows is not None and str(args.mode) == "full":
                    seed_rows_by_query_id = {
                        str(variant): list(rows[: int(args.inspection_count_per_query_id)])
                        for variant, rows in distribution_variant_rows.items()
                    }
                else:
                    inspection_collected = collect_query_id_samples(
                        task_id=str(task_id),
                        target_count_per_query_id=int(args.inspection_count_per_query_id),
                        seed=int(args.seed) + 17,
                        max_attempts_per_instance=int(args.max_attempts_per_instance),
                        max_total_samples_per_task=int(args.max_total_samples_per_task),
                        workers=int(args.workers),
                        collector=_answer_collector,
                    )
                    seed_rows_by_query_id = {
                        str(variant): list(inspection_collected.get("samples_by_query_id", {}).get(str(variant), []))
                        for variant in list(inspection_collected.get("expected_query_ids", []))
                    }
            else:
                if random_rows_for_inspection is None:
                    random_rows_for_inspection = generate_random_samples(
                        task_id=str(task_id),
                        count=int(args.inspection_count_per_query_id),
                        seed=int(args.seed) + 17,
                        max_attempts_per_instance=int(args.max_attempts_per_instance),
                        workers=int(args.workers),
                        collector=_answer_collector,
                    )
                seed_rows_by_query_id = {
                    "": list(random_rows_for_inspection[: int(args.inspection_count_per_query_id)]),
                }

            inspection_manifest = _build_inspection_rows(
                task_id=str(task_id),
                out_root=out_root,
                task_dir=task_dir,
                seed_rows_by_query_id=seed_rows_by_query_id,
                max_attempts_per_instance=int(args.max_attempts_per_instance),
            )
            task_summary["reports"]["inspection_manifest"] = str(
                (task_dir / "manifest.json").relative_to(out_root).as_posix()
            )
            task_summary["inspection_count"] = int(inspection_manifest.get("inspection_count", 0))
            task_summary["inspection_workbook"] = str(inspection_manifest.get("workbook", ""))
            task_summary["inspection_workbook_sheets"] = dict(inspection_manifest.get("workbook_sheets", {}))
            workbook_hint = str(task_summary["inspection_workbook"]) or "none"
            sheet_count = len(task_summary["inspection_workbook_sheets"])
            print(
                f"[done] {task_id} inspection workbook: {workbook_hint} (sheets={sheet_count})"
            )

        summary["tasks"].append(task_summary)

    summary["summary"] = {
        "total_tasks": int(len(task_ids)),
        "distribution_tasks_failed": int(len(failed_distribution_tasks)),
        "failed_distribution_task_ids": sorted(failed_distribution_tasks),
    }
    if str(args.mode) in {"full", "inspection"} and not bool(args.skip_scene_workbooks):
        summary["scene_workbooks"] = build_scene_review_workbooks(
            out_root=out_root,
            scene_keys=sorted(touched_scene_keys),
        )
    summary_path = out_root / "review_summary.json"
    write_json_file(summary_path, summary)
    print(f"[done] wrote review summary: {summary_path}")

    if failed_distribution_tasks and str(args.mode) in {"full", "distribution"} and not bool(args.allow_fail):
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
