#!/usr/bin/env python3
"""Summarize TRACE answer+annotation rollout diagnostics."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import gzip
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _json_default(value: Any) -> Any:
    if hasattr(value, "tolist"):
        return value.tolist()
    if hasattr(value, "item"):
        return value.item()
    return str(value)


def _as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return float(value) != 0.0
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "y"}
    return bool(value)


def _as_float(value: Any) -> float:
    if value is None:
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _read_jsonl(path: Path):
    opener = gzip.open if str(path).endswith(".gz") else open
    with opener(path, "rt", encoding="utf-8", errors="ignore") as handle:
        for line in handle:
            stripped = line.strip()
            if not stripped:
                continue
            yield json.loads(stripped)


def _new_aggregate() -> dict[str, Any]:
    return {
        "rollout_count": 0,
        "prompt_rollouts": defaultdict(int),
        "prompt_positive": defaultdict(int),
        "sum_strict_task_reward": 0.0,
        "sum_strict_answer_reward": 0.0,
        "sum_strict_annotation_reward": 0.0,
        "sum_strict_format_reward": 0.0,
        "sum_fallback_task_reward": 0.0,
        "sum_fallback_answer_reward": 0.0,
        "sum_fallback_annotation_reward": 0.0,
        "sum_fallback_format_reward": 0.0,
        "strict_json_found": 0,
        "strict_format_structure_ok": 0,
        "strict_format_json_ok": 0,
        "strict_format_schema_ok": 0,
        "strict_answer_parse_ok": 0,
        "strict_annotation_parse_ok": 0,
        "fallback_positive": 0,
        "key_order_ok": 0,
        "single_answer_tag": 0,
        "hit_response_cap": 0,
        "sum_generated_tokens": 0.0,
        "max_generated_tokens": 0,
        "failure_categories": Counter(),
    }


def _add_row(aggregate: dict[str, Any], row: dict[str, Any]) -> None:
    rollout_count = int(aggregate["rollout_count"]) + 1
    aggregate["rollout_count"] = rollout_count
    prompt_key = str(row.get("dataset_index"))
    aggregate["prompt_rollouts"][prompt_key] += 1
    if _as_bool(row.get("strict_positive")):
        aggregate["prompt_positive"][prompt_key] += 1
    aggregate["sum_strict_task_reward"] += _as_float(row.get("strict_task_reward"))
    aggregate["sum_strict_answer_reward"] += _as_float(row.get("strict_answer_reward"))
    aggregate["sum_strict_annotation_reward"] += _as_float(row.get("strict_annotation_reward"))
    aggregate["sum_strict_format_reward"] += _as_float(row.get("strict_format_reward"))
    aggregate["sum_fallback_task_reward"] += _as_float(row.get("fallback_task_reward"))
    aggregate["sum_fallback_answer_reward"] += _as_float(row.get("fallback_answer_reward"))
    aggregate["sum_fallback_annotation_reward"] += _as_float(row.get("fallback_annotation_reward"))
    aggregate["sum_fallback_format_reward"] += _as_float(row.get("fallback_format_reward"))
    for key in (
        "strict_json_found",
        "strict_format_structure_ok",
        "strict_format_json_ok",
        "strict_format_schema_ok",
        "strict_answer_parse_ok",
        "strict_annotation_parse_ok",
        "fallback_positive",
        "key_order_ok",
        "hit_response_cap",
    ):
        aggregate[key] += 1 if _as_bool(row.get(key)) else 0
    aggregate["single_answer_tag"] += 1 if int(row.get("answer_tag_count") or 0) == 1 else 0
    generated_tokens = int(row.get("generated_tokens") or 0)
    aggregate["sum_generated_tokens"] += float(generated_tokens)
    aggregate["max_generated_tokens"] = max(int(aggregate["max_generated_tokens"]), generated_tokens)
    for category in row.get("failure_categories") or []:
        aggregate["failure_categories"][str(category)] += 1


def _rate(numerator: float, denominator: float) -> float:
    return float(numerator / denominator) if denominator else 0.0


def _finalize(name: str, aggregate: dict[str, Any]) -> dict[str, Any]:
    rollout_count = int(aggregate["rollout_count"])
    prompt_rollouts: dict[str, int] = dict(aggregate["prompt_rollouts"])
    prompt_positive: dict[str, int] = dict(aggregate["prompt_positive"])
    prompt_count = len(prompt_rollouts)
    easy_count = 0
    hard_count = 0
    solve_sum = 0.0
    for prompt_key, prompt_rollout_count in prompt_rollouts.items():
        positives = int(prompt_positive.get(prompt_key, 0))
        if positives == 0:
            hard_count += 1
        if prompt_rollout_count > 0 and positives == prompt_rollout_count:
            easy_count += 1
        solve_sum += _rate(positives, prompt_rollout_count)
    band_count = max(0, prompt_count - easy_count - hard_count)
    failure_categories: Counter[str] = aggregate["failure_categories"]
    top_failures = ", ".join(
        f"{category}:{count}" for category, count in failure_categories.most_common(5)
    )
    return {
        "name": str(name),
        "prompt_count": int(prompt_count),
        "rollout_count": int(rollout_count),
        "mean_solve_rate": _rate(solve_sum, prompt_count),
        "easy_rate": _rate(easy_count, prompt_count),
        "hard_rate": _rate(hard_count, prompt_count),
        "band_rate": _rate(band_count, prompt_count),
        "easy_count": int(easy_count),
        "hard_count": int(hard_count),
        "band_count": int(band_count),
        "strict_task_reward_mean": _rate(aggregate["sum_strict_task_reward"], rollout_count),
        "strict_answer_reward_mean": _rate(aggregate["sum_strict_answer_reward"], rollout_count),
        "strict_annotation_reward_mean": _rate(aggregate["sum_strict_annotation_reward"], rollout_count),
        "strict_format_reward_mean": _rate(aggregate["sum_strict_format_reward"], rollout_count),
        "fallback_task_reward_mean": _rate(aggregate["sum_fallback_task_reward"], rollout_count),
        "fallback_answer_reward_mean": _rate(aggregate["sum_fallback_answer_reward"], rollout_count),
        "fallback_annotation_reward_mean": _rate(aggregate["sum_fallback_annotation_reward"], rollout_count),
        "fallback_format_reward_mean": _rate(aggregate["sum_fallback_format_reward"], rollout_count),
        "fallback_positive_rate": _rate(aggregate["fallback_positive"], rollout_count),
        "single_answer_tag_rate": _rate(aggregate["single_answer_tag"], rollout_count),
        "strict_json_found_rate": _rate(aggregate["strict_json_found"], rollout_count),
        "strict_format_structure_ok_rate": _rate(aggregate["strict_format_structure_ok"], rollout_count),
        "strict_format_json_ok_rate": _rate(aggregate["strict_format_json_ok"], rollout_count),
        "strict_format_schema_ok_rate": _rate(aggregate["strict_format_schema_ok"], rollout_count),
        "key_order_ok_rate": _rate(aggregate["key_order_ok"], rollout_count),
        "strict_answer_parse_ok_rate": _rate(aggregate["strict_answer_parse_ok"], rollout_count),
        "strict_annotation_parse_ok_rate": _rate(aggregate["strict_annotation_parse_ok"], rollout_count),
        "hit_response_cap_rate": _rate(aggregate["hit_response_cap"], rollout_count),
        "mean_generated_tokens": _rate(aggregate["sum_generated_tokens"], rollout_count),
        "max_generated_tokens": int(aggregate["max_generated_tokens"]),
        "top_failure_categories": top_failures,
        "failure_categories": dict(sorted(failure_categories.items())),
    }


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=_json_default), encoding="utf-8")


def _write_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]], columns: list[str]) -> None:
    sheet = workbook.create_sheet(title=title)
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row.get(column) for column in columns])
    for column_index, column in enumerate(columns, 1):
        max_width = len(column)
        for cell in sheet.iter_cols(min_col=column_index, max_col=column_index, min_row=2, values_only=True):
            for value in cell:
                max_width = max(max_width, len(str(value)) if value is not None else 0)
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max_width + 2, 60)
    sheet.freeze_panes = "A2"


def _summary_columns() -> list[str]:
    return [
        "name",
        "prompt_count",
        "rollout_count",
        "mean_solve_rate",
        "easy_rate",
        "hard_rate",
        "band_rate",
        "strict_task_reward_mean",
        "strict_answer_reward_mean",
        "strict_annotation_reward_mean",
        "strict_format_reward_mean",
        "fallback_task_reward_mean",
        "fallback_positive_rate",
        "single_answer_tag_rate",
        "strict_json_found_rate",
        "strict_format_json_ok_rate",
        "strict_format_schema_ok_rate",
        "key_order_ok_rate",
        "strict_answer_parse_ok_rate",
        "strict_annotation_parse_ok_rate",
        "hit_response_cap_rate",
        "mean_generated_tokens",
        "max_generated_tokens",
        "top_failure_categories",
    ]


def main() -> int:
    parser = argparse.ArgumentParser(description="Summarize TRACE annotation diagnostic rollouts.")
    parser.add_argument("--per-rollout", type=Path, required=True, help="per_rollout.jsonl or per_rollout.jsonl.gz")
    parser.add_argument("--output-dir", type=Path, required=True, help="Directory for JSON summaries and workbook.")
    parser.add_argument("--label", default="trace_annotation_eval", help="Output filename label.")
    args = parser.parse_args()

    aggregates = {
        "overall": defaultdict(_new_aggregate),
        "domain": defaultdict(_new_aggregate),
        "task": defaultdict(_new_aggregate),
        "query_id": defaultdict(_new_aggregate),
        "annotation_type": defaultdict(_new_aggregate),
    }
    failure_totals: Counter[str] = Counter()

    for row in _read_jsonl(args.per_rollout):
        domain = str(row.get("domain") or "")
        task = str(row.get("task") or "")
        annotation_type = str(row.get("annotation_type") or "")
        query_id = str(row.get("query_id") or "unknown")
        group_keys = {
            "overall": "overall",
            "domain": domain,
            "task": task,
            "query_id": f"{task}::{query_id}",
            "annotation_type": annotation_type,
        }
        for group_name, group_key in group_keys.items():
            _add_row(aggregates[group_name][group_key], row)
        for category in row.get("failure_categories") or []:
            failure_totals[str(category)] += 1

    finalized = {
        group_name: [
            _finalize(group_key, aggregate)
            for group_key, aggregate in sorted(group_aggregates.items())
        ]
        for group_name, group_aggregates in aggregates.items()
    }

    args.output_dir.mkdir(parents=True, exist_ok=True)
    for group_name, rows in finalized.items():
        _write_json(args.output_dir / f"{args.label}_{group_name}_summary.json", rows)
    _write_json(
        args.output_dir / f"{args.label}_failure_category_summary.json",
        [
            {"failure_category": category, "count": count}
            for category, count in failure_totals.most_common()
        ],
    )

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    columns = _summary_columns()
    _write_sheet(workbook, "overall", finalized["overall"], columns)
    _write_sheet(workbook, "by_task", finalized["task"], columns)
    _write_sheet(workbook, "by_domain", finalized["domain"], columns)
    _write_sheet(workbook, "by_query_id", finalized["query_id"], columns)
    _write_sheet(workbook, "by_annotation_type", finalized["annotation_type"], columns)
    _write_sheet(
        workbook,
        "failure_categories",
        [{"name": category, "count": count} for category, count in failure_totals.most_common()],
        ["name", "count"],
    )
    workbook_path = args.output_dir / f"{args.label}.xlsx"
    workbook.save(workbook_path)
    print(f"[done] wrote {workbook_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
