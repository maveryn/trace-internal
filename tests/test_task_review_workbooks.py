"""Regression tests for shared task-review workbook helpers."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from trace.core.task_review_workbooks import write_inspection_excel, write_scene_inspection_excel


def _inspection_row(*, query_id: str, image_path: str = "missing.png") -> dict[str, object]:
    return {
        "task": "task_dummy__scene__foo",
        "query_id": str(query_id),
        "scene_id": "scene",
        "prompt_answer": "Return the value.",
        "ground_truth_answer": {"answer": 3},
        "prompt_answer_and_annotation": "Return the value and annotation.",
        "ground_truth_answer_and_annotation": {"annotation": [[1, 2, 3, 4]], "answer": 3},
        "answer_type": "integer",
        "annotation_type": "bbox_set",
        "answer_annotation": [[1, 2, 3, 4]],
        "instance_seed": 11,
        "image_path": str(image_path),
        "data_path": "data/0000.json",
    }


def test_write_inspection_excel_dedupes_query_id_sheet_titles(tmp_path: Path) -> None:
    image_path = tmp_path / "sample.png"
    Image.new("RGB", (32, 32), color=(255, 255, 255)).save(image_path)

    workbook_path = tmp_path / "task.xlsx"
    sheet_map = write_inspection_excel(
        {
            "a/b": [_inspection_row(query_id="a/b", image_path="sample.png")],
            "a?b": [_inspection_row(query_id="a?b")],
        },
        workbook_path,
        out_root=tmp_path,
    )

    assert sheet_map == {"a/b": "a_b", "a?b": "a_b_2"}
    workbook = load_workbook(workbook_path)
    assert workbook.sheetnames == ["a_b", "a_b_2"]
    assert workbook["a_b"]["C2"].value == "task_dummy__scene__foo"
    assert workbook["a_b"]["H2"].value == '{"answer": 3}'


def test_write_scene_inspection_excel_adds_model_stats_sheet(tmp_path: Path) -> None:
    workbook_path = tmp_path / "scene_review.xlsx"

    task_sheets, model_stats_sheet, model_stats_count = write_scene_inspection_excel(
        {"task_dummy__scene__foo": [_inspection_row(query_id="default")]},
        workbook_path,
        out_root=tmp_path,
        domain="dummy",
        model_stats_rows=[
            {
                "task": "task_dummy__scene__foo",
                "combined_status": "accepted",
                "model": "qwen25vl7b",
            }
        ],
    )

    assert model_stats_sheet == "model_stats"
    assert model_stats_count == 1
    assert task_sheets == {"task_dummy__scene__foo": "_scene__foo"}
    workbook = load_workbook(workbook_path)
    assert workbook.sheetnames[0] == "model_stats"
    assert workbook["model_stats"]["A2"].value == "task_dummy__scene__foo"
